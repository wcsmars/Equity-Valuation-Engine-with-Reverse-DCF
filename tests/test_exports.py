"""Offline checks for the demo run, every export format, and the reverse DCF.

All tests use the synthetic company from ``equity_valuation.data.synthetic``,
so none of them touch the network. Run with:  python -m unittest tests.test_exports
"""

from __future__ import annotations

import contextlib
import dataclasses
import io
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from equity_valuation import value_company
from equity_valuation.data.synthetic import DEMO_PEERS, DEMO_TICKER, SyntheticProvider


def _demo_report():
    return value_company(DEMO_TICKER, provider=SyntheticProvider(), peers=DEMO_PEERS)


class DemoResultTests(unittest.TestCase):
    def test_summary_matches_readme_results(self):
        # The README "Results" table quotes these figures; keep them in step.
        s = _demo_report().summary
        self.assertAlmostEqual(s["current_price"], 40.84, delta=0.005)
        expected = {"DCF": 33.38, "Comps (median)": 46.96, "DDM": 10.99, "FCFE": 31.41}
        self.assertEqual(set(s["methods"]), set(expected))
        for method, price in expected.items():
            self.assertAlmostEqual(s["methods"][method], price, delta=0.005, msg=method)
        self.assertAlmostEqual(s["blended_target"], 32.40, delta=0.005)
        self.assertAlmostEqual(s["blended_upside"], -0.207, delta=0.0005)
        self.assertEqual(s["recommendation"], "Overvalued")

    def test_cli_demo_writes_excel_and_html(self):
        from equity_valuation.cli import main

        with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(main(["--demo", "--quiet", "--out", tmp]), 0)
            self.assertTrue(os.path.isfile(os.path.join(tmp, "SYNT_valuation.xlsx")))
            self.assertTrue(os.path.isfile(os.path.join(tmp, "SYNT_valuation.html")))

    def test_cli_demo_rejects_a_live_ticker(self):
        from equity_valuation.cli import main

        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
            main(["AAPL", "--demo"])


class EngineExportTests(unittest.TestCase):
    def test_excel_keeps_formulas_and_html_renders(self):
        from openpyxl import load_workbook

        from equity_valuation.report.excel import write_excel
        from equity_valuation.report.html import write_html

        report = _demo_report()
        with tempfile.TemporaryDirectory() as tmp:
            xlsx = write_excel(report, os.path.join(tmp, "SYNT_valuation.xlsx"))
            html = write_html(report, os.path.join(tmp, "SYNT_valuation.html"))

            wb = load_workbook(xlsx)
            for sheet in ("Summary", "DCF", "Comps", "DDM_FCFE", "Sensitivity"):
                self.assertIn(sheet, wb.sheetnames)
            formulas = [
                c for row in wb["DCF"].iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")
            ]
            self.assertTrue(formulas, "DCF sheet should contain live formulas")

            text = Path(html).read_text(encoding="utf-8")
            self.assertIn("Synthetic Corp", text)
            self.assertIn("<html", text.lower())


class OfficeExportTests(unittest.TestCase):
    """The dashboard's Word memo and PowerPoint briefing, without a live provider."""

    def test_memo_and_deck(self):
        from docx import Document
        from pptx import Presentation

        from backend import exports, valuation_service

        payload = {"peers": ",".join(DEMO_PEERS)}
        with tempfile.TemporaryDirectory() as tmp, patch.object(
            valuation_service, "_get_provider",
            lambda ticker, refresh=False: SyntheticProvider(),
        ), patch.object(exports, "_ensure_out", lambda: Path(tmp)):
            memo = exports.export_memo(DEMO_TICKER, payload)
            deck = exports.export_deck(DEMO_TICKER, payload)

            self.assertEqual(Path(memo).parent, Path(tmp))
            text = "\n".join(p.text for p in Document(memo).paragraphs)
            self.assertIn("Synthetic Corp", text)
            self.assertIn("Overvalued", text)
            self.assertGreaterEqual(len(Presentation(deck).slides), 1)

    def test_zero_method_value_is_minus_100_percent_in_both_exports(self):
        from docx import Document
        from pptx import Presentation
        from backend import exports

        report = _demo_report()
        report.summary["methods"]["DCF"] = 0.0
        report.dcf.implied_price = 0.0
        report.dcf.assumptions["valuation_available"] = True
        with tempfile.TemporaryDirectory() as tmp, \
                patch.object(exports, "run_valuation_report", return_value=(report, {})), \
                patch.object(exports, "_ensure_out", lambda: Path(tmp)):
            memo = Document(exports.export_memo(DEMO_TICKER, {}))
            self.assertEqual(memo.tables[0].rows[1].cells[2].text, "-100.0%")
            deck = Presentation(exports.export_deck(DEMO_TICKER, {}))
            table = next(s.table for s in deck.slides[1].shapes if s.has_table)
            self.assertEqual(table.cell(1, 2).text, "-100.0%")

    def test_unavailable_method_is_not_minus_100_percent_in_office_exports(self):
        from docx import Document
        from pptx import Presentation
        from backend import exports

        report = _demo_report()
        report.summary["methods"]["DCF"] = 0.0
        report.dcf.implied_price = 0.0
        report.dcf.assumptions["valuation_available"] = False
        with tempfile.TemporaryDirectory() as tmp, \
                patch.object(exports, "run_valuation_report", return_value=(report, {})), \
                patch.object(exports, "_ensure_out", lambda: Path(tmp)):
            memo = Document(exports.export_memo(DEMO_TICKER, {}))
            self.assertEqual(memo.tables[0].rows[1].cells[1].text, "n/a")
            self.assertEqual(memo.tables[0].rows[1].cells[2].text, "n/a")
            deck = Presentation(exports.export_deck(DEMO_TICKER, {}))
            table = next(s.table for s in deck.slides[1].shapes if s.has_table)
            self.assertEqual(table.cell(1, 1).text, "n/a")
            self.assertEqual(table.cell(1, 2).text, "n/a")

    def test_deck_paginates_peers_notes_and_preserves_sources(self):
        from pptx import Presentation
        from backend import exports

        report = _demo_report()
        report.comps.peers = [dataclasses.replace(report.comps.peers[0],
                              ticker=f"PEER{i}", currency="HKD") for i in range(35)]
        note = {
            "executive_summary": "Summary material.",
            "valuation_view": "Valuation discussion.",
            "thesis": [f"Thesis item {i}: " + "Revenue assumptions need evidence. " * 30
                       for i in range(9)],
            "risks": ["A material risk."],
            "key_drivers": ["Demand."],
            "catalysts": [f"Catalyst {i}" for i in range(6)],
            "red_flags": ["Verify accounting."],
            "what_would_change_my_mind": ["Failure to meet the operating plan."],
            "citations": [{"source": "Annual report", "note": "Page 42"}],
        }
        with tempfile.TemporaryDirectory() as tmp, \
                patch.object(exports, "run_valuation_report", return_value=(report, {})), \
                patch.object(exports, "_ensure_out", lambda: Path(tmp)):
            deck = Presentation(exports.export_deck(DEMO_TICKER, {}, note))
        texts, peer_rows = [], []
        for slide in deck.slides:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.top + shape.height, deck.slide_height)
                self.assertLessEqual(shape.left + shape.width, deck.slide_width)
                if shape.has_text_frame:
                    texts.append(shape.text)
                if shape.has_table and shape.table.cell(0, 0).text == "Ticker":
                    self.assertLessEqual(len(shape.table.rows), 13)
                    peer_rows += [[c.text for c in row.cells] for row in list(shape.table.rows)[1:]]
        joined = "\n".join(texts)
        for i in range(9):
            self.assertIn(f"Thesis item {i}:", joined)
        self.assertIn("Catalyst 5", joined)
        self.assertIn("Annual report: Page 42", joined)
        self.assertIn("Valuation discussion.", joined)
        self.assertIn("Failure to meet the operating plan.", joined)
        self.assertEqual(len(peer_rows), 36)
        self.assertEqual(peer_rows[-1][:2], ["PEER34", "HKD"])

    def test_missing_price_does_not_break_football_field(self):
        from pptx import Presentation
        from backend import exports

        report = _demo_report()
        report.summary["current_price"] = None
        with tempfile.TemporaryDirectory() as tmp, \
                patch.object(exports, "run_valuation_report", return_value=(report, {})), \
                patch.object(exports, "_ensure_out", lambda: Path(tmp)):
            deck = Presentation(exports.export_deck(DEMO_TICKER, {}))
        field = deck.slides[2]
        self.assertIn("Football field", field.shapes[0].text)
        self.assertFalse(any(s.has_text_frame and s.text.startswith("price ") for s in field.shapes))


def _office(report, note=None):
    """(memo Document, deck Presentation) built from ``report`` offline."""
    from docx import Document
    from pptx import Presentation
    from backend import exports

    with tempfile.TemporaryDirectory() as tmp, \
            patch.object(exports, "run_valuation_report", return_value=(report, {})), \
            patch.object(exports, "_ensure_out", lambda: Path(tmp)):
        memo = Document(exports.export_memo(DEMO_TICKER, {}, note))
        deck = Presentation(exports.export_deck(DEMO_TICKER, {}, note))
    return memo, deck


def _slide(deck, title):
    return next(s for s in deck.slides
                if s.shapes and s.shapes[0].has_text_frame and s.shapes[0].text == title)


class DeckLayoutTests(unittest.TestCase):
    """How the briefing deck lays its text out (renderers draw only what the
    slide XML says: a line feed inside a run is not a line break)."""

    NOTE = {
        "thesis": ["Thesis point 0: " + "long-term demand " * 12, "Thesis point 1: short."],
        "risks": ["Risk A short.", "Risk B short.", "Risk C short."],
        "citations": [{"source": "Annual report", "note": "Page 42"}],
    }

    def test_note_items_and_model_notes_are_separate_paragraphs(self):
        from pptx.util import Pt

        report = _demo_report()
        report.warnings.append("Second model note.")
        _memo, deck = _office(report, self.NOTE)
        for slide in deck.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        for r in p.runs:
                            self.assertNotIn("\n", r.text)
        body = _slide(deck, "Risks").shapes[1].text_frame
        self.assertEqual([p.text for p in body.paragraphs],
                         ["Risk A short.", "", "Risk B short.", "", "Risk C short."])
        # A blank separator line is as tall as the text around it.
        self.assertEqual(body.paragraphs[1]._p.endParaRPr.sz, Pt(18).centipoints)
        thesis = [p.text for p in _slide(deck, "Thesis").shapes[1].text_frame.paragraphs]
        # One paragraph per item for the slide to wrap, not a hard break at 95 characters.
        self.assertEqual(thesis, [" ".join(self.NOTE["thesis"][0].split()), "",
                                  "Thesis point 1: short."])
        notes = [p.text for p in _slide(deck, "Model notes").shapes[1].text_frame.paragraphs]
        self.assertEqual(notes[-1], "Second model note.")
        self.assertEqual(notes[-2], "")

    def test_text_pages_split_long_items_without_changing_their_text(self):
        from backend.exports import _text_pages

        url = "https://example.com/" + "a-very-long-path-segment/" * 8 + "page.htm"
        long_item = " ".join(f"word{i} long-term" for i in range(200))
        pages = _text_pages([f"Source: {url}", "Short   item.", long_item])
        self.assertGreater(len(pages), 2)
        paragraphs = [p for page in pages for p in page.split("\n")]
        # Whole paragraphs, whitespace runs collapsed, never cut inside a URL
        # or a hyphenated word (the old 95-character wrap became hard breaks).
        self.assertEqual(paragraphs[:4], [f"Source: {url}", "", "Short item.", ""])
        # The long item runs over several pages; its pieces are consecutive
        # slices of it, split only between two wrapped lines.
        pieces = paragraphs[4:]
        self.assertGreater(len(pieces), 1)
        rebuilt = pieces[0]
        for piece in pieces[1:]:
            self.assertTrue(long_item.startswith(rebuilt + " " + piece)
                            or long_item.startswith(rebuilt + piece), piece[:30])
            rebuilt += (" " if long_item.startswith(rebuilt + " " + piece) else "") + piece
        self.assertEqual(rebuilt, long_item)
        for page in pages:  # each page stays within its wrapped-line budget
            self.assertLessEqual(sum(max(1, -(-len(p) // 95)) for p in page.split("\n")), 15)

    def test_control_characters_do_not_sink_the_memo_or_the_deck(self):
        report = _demo_report()
        report.summary["name"] = "Synthetic\x07 Corp"
        report.warnings.append("Comps: peer fetch failed: bad byte \x1b[0m in response")
        note = {"stance": "Neutral\x0c", "thesis": ["Line one\x0bstill one.",
                                                    "Group\x1drecord\x1eend."],
                "citations": [{"source": "Filing\x00", "note": "p. 4"}]}
        memo, deck = _office(report, note)
        text = "\n".join(p.text for p in memo.paragraphs)
        self.assertIn("bad byte [0m in response", text)
        self.assertIn("Analytical stance: Neutral", text)
        # A vertical tab (Word's manual line break), form feed or separator is
        # a line break in the memo, as in the deck: the words stay apart.
        self.assertIn("Line one\nstill one.", text)
        self.assertIn("Group\nrecord\nend.", text)
        self.assertNotIn("onestill", text)
        self.assertIn("Filing: p. 4", text)
        self.assertEqual(memo.core_properties.title, "Synthetic Corp (SYNT) — Research Memo")
        shapes = "\n".join(s.text for sl in deck.slides for s in sl.shapes if s.has_text_frame)
        self.assertIn("bad byte [0m in response", shapes)
        self.assertIn("Line one\nstill one.", shapes)
        self.assertIn("Group\nrecord\nend.", shapes)
        self.assertNotIn("_x00", shapes)

    def test_a_long_company_name_does_not_run_into_the_subtitle(self):
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE

        for name, size in (("INTERNATIONAL BUSINESS MACHINES CORP", 40),
                           ("Taiwan Semiconductor Manufacturing Company Limited Holdings "
                            "and Subsidiaries", 30)):
            report = _demo_report()
            report.summary["name"] = name
            _memo, deck = _office(report)
            title, subtitle = deck.slides[0].shapes[0], deck.slides[0].shapes[1]
            with self.subTest(name):
                self.assertTrue(title.text.startswith(name))
                self.assertTrue(subtitle.text.startswith("Valuation briefing"))
                # Two 40pt lines fit the box, which ends above the subtitle and
                # grows the text upward from its bottom edge.
                self.assertLessEqual(title.top + title.height, subtitle.top)
                self.assertGreaterEqual(title.height, 2 * 40 * 1.2 * 12700)
                self.assertEqual(title.text_frame.vertical_anchor, MSO_ANCHOR.BOTTOM)
                self.assertEqual(title.text_frame.auto_size, MSO_AUTO_SIZE.NONE)
                self.assertEqual(title.text_frame.paragraphs[0].runs[0].font.size.pt, size)

    def test_football_field_labels_each_range_and_the_exact_price(self):
        from backend.exports import _money

        report = _demo_report()
        _memo, deck = _office(report)
        texts = [s.text for s in _slide(deck, "Football field — value ranges by method").shapes
                 if s.has_text_frame]
        self.assertIn(f"price {_money(report.current_price)}", texts)
        self.assertIn("price $40.84", texts)
        for r in report.football_field:
            self.assertIn(f"{_money(r.low)} – {_money(r.high)}", texts, r.method)

    def test_money_uses_the_currency_prefix_and_signs_first(self):
        from backend.exports import _money

        self.assertEqual(_money(40.84, "CAD"), "C$40.84")
        self.assertEqual(_money(40.84, "SEK"), "SEK 40.84")
        self.assertEqual(_money(-6.52), "-$6.52")
        self.assertEqual(_money(-0.004), "$0.00")
        self.assertEqual(_money(40.84, None), "40.84")
        # Provider text that is not a three-letter code gets no prefix.
        for code in ('U"S', "<B>X</B>", "US", "USDX", "\x1b"):
            self.assertEqual(_money(40.84, code), "40.84", repr(code))
        report = _demo_report()
        report.summary["currency"] = "SEK"
        memo, deck = _office(report)
        self.assertIn("price SEK 40.84", memo.paragraphs[1].text)
        self.assertIn("Price SEK 40.84", deck.slides[0].shapes[2].text)


class SerializationTests(unittest.TestCase):
    """The dashboard payload and the AI context for an adjusted beta and for a
    summary with no blended target or a withheld upside."""

    def _adjusted_report(self):
        from equity_valuation.data.synthetic import make_company

        class _Adjusted(SyntheticProvider):
            def get_company_data(self, ticker):
                company = make_company()
                company.market.raw_beta = 2.217
                company.market.beta = 0.67 * 2.217 + 0.33
                return company

        return value_company(DEMO_TICKER, provider=_Adjusted(), peers=DEMO_PEERS)

    def test_raw_beta_reaches_the_payload_and_the_ai_context(self):
        from backend.serialization import build_ai_context, report_to_dict

        d = report_to_dict(self._adjusted_report())
        self.assertEqual(d["company"]["market"]["raw_beta"], 2.217)
        self.assertAlmostEqual(d["company"]["market"]["beta"], 1.81539, places=9)
        self.assertEqual(d["dcf"]["wacc"]["detail"]["beta_raw"], 2.217)
        ctx = build_ai_context(d)
        self.assertIn("beta 1.815 (Blume-adjusted from raw 2.217)", ctx)
        self.assertEqual(ctx.count("Blume-adjusted from raw 2.217"), 2)  # PRICE and DCF

        # The synthetic beta is not adjusted: raw_beta is null, no adjustment is claimed.
        d = report_to_dict(_demo_report())
        self.assertIsNone(d["company"]["market"]["raw_beta"])
        ctx = build_ai_context(d)
        self.assertIn("beta 1.100 |", ctx)
        self.assertNotIn("Blume", ctx)

    @staticmethod
    def _kind_payload(kind=None, industry=None):
        """The dashboard payload of the engine's run for the synthetic company
        flagged as that kind (or presented as a bank), without peers."""
        from backend.serialization import report_to_dict
        from equity_valuation.data.synthetic import make_company

        class _Kind(SyntheticProvider):
            def get_company_data(self, ticker):
                company = make_company()
                if kind:
                    company.financials._financial_kind = kind
                if industry:
                    company.market.sector = "Financial Services"
                    company.market.industry = industry
                return company

        return report_to_dict(value_company(DEMO_TICKER, provider=_Kind(), run_comps=False))

    @staticmethod
    def _line(ctx, prefix):
        return next(ln for ln in ctx.splitlines() if ln.startswith(prefix))

    def test_ai_context_says_when_there_is_no_target_or_no_upside(self):
        from backend.serialization import build_ai_context, report_to_dict

        ctx = build_ai_context(report_to_dict(_demo_report()))
        self.assertIn("VERDICT: Overvalued | blended target $32.40 (-20.7% vs price)", ctx)
        # The demo's DCF is in the blend: a signed upside, no marker.
        dcf = self._line(ctx, "DCF: ")
        self.assertTrue(dcf.endswith("implied $33.38 (-18.3% vs price)"), dcf)
        self.assertNotIn("reference only", dcf)

        # A lessor without peers (as AER): no target, every method reference only.
        d = self._kind_payload("lessor")
        excluded = d["summary"]["excluded_from_blend"]
        self.assertEqual(set(excluded), {"DCF", "DDM", "FCFE"})
        ctx = build_ai_context(d)
        self.assertIn("VERDICT: N/A | blended target: none (no target; supply peers to add "
                      "trading comps)", ctx)
        for name in ("DCF", "DDM", "FCFE"):
            self.assertIn(f"{name} ${d['summary']['methods'][name]:,.2f} (reference only, "
                          f"not in blend: {excluded[name]})", ctx)
        # The DCF line is marked too, so its upside is not quoted as the view.
        dcf = self._line(ctx, "DCF: ")
        self.assertTrue(dcf.endswith(f"(-18.3% vs price) (reference only, not in blend: "
                                     f"{excluded['DCF']})"), dcf)

        # A bank without peers (as BAC): the DDM-only target stays, the verdict
        # is withheld; the DDM is in the blend, the DCF is not.
        d = self._kind_payload(industry="Banks - Diversified")
        ctx = build_ai_context(d)
        self.assertEqual(d["summary"]["verdict_withheld"], "the target rests on the DDM alone")
        self.assertIn("VERDICT: N/A | blended target $10.99 (upside withheld: the target rests "
                      "on the DDM alone)", ctx)
        self.assertIn("DDM $10.99;", ctx)
        self.assertIn("(reference only, not in blend: not meaningful for a financial "
                      "institution)", self._line(ctx, "DCF: "))

        # The engine's reason is quoted (a corporate action), "n/a" without one.
        d["summary"].update(verdict_withheld="a corporate action after FY2025 (likely a "
                                             "spin-off)")
        self.assertIn("VERDICT: N/A | blended target $10.99 (upside withheld: a corporate "
                      "action after FY2025 (likely a spin-off))", build_ai_context(d))
        d["summary"].update(verdict_withheld=None)
        self.assertIn("VERDICT: N/A | blended target $10.99 (upside withheld: n/a)",
                      build_ai_context(d))
        # Hand-built (the engine needs a price): no upside because there is no
        # current price is not a withheld verdict.
        d["summary"].update(current_price=None)
        self.assertIn("VERDICT: N/A | blended target $10.99 (upside n/a: no current price)",
                      build_ai_context(d))
        # A positive upside carries its sign.
        d = report_to_dict(_demo_report())
        d["summary"].update(blended_target=49.0, blended_upside=0.2)
        self.assertIn("blended target $49.00 (+20.0% vs price)", build_ai_context(d))
        # An ordinary company whose methods gave no valuation: no peers asked for.
        d["summary"].update(methods={"DCF": 0.0}, blended_target=None, blended_upside=None,
                            recommendation="N/A", financial_kind=None,
                            financial_institution=None)
        d["comps"] = None
        self.assertIn("VERDICT: N/A | blended target: n/a", build_ai_context(d))


def _office_texts(report):
    """(memo paragraphs, every deck text) of the memo and deck built offline
    from ``report``."""
    memo, deck = _office(report)
    paragraphs = [p.text for p in memo.paragraphs]
    texts = [shape.text for slide in deck.slides for shape in slide.shapes
             if shape.has_text_frame]
    return paragraphs, texts, deck


DEMO_REVERSE = ("Reverse DCF (other inputs held) — The price implies 17.5% year-1 growth "
                "(model 8.0%), a 30.8% target EBIT margin (model 25.0%) or an 8.2% discount "
                "rate (model 9.5%).")


class ReverseDcfOfficeTests(unittest.TestCase):
    """The memo and the deck state what the price implies (the engine's
    reverse DCF) and, when the DCF flagged it, the terminal reinvestment check;
    the demo's pinned figures are unchanged."""

    def test_memo_and_deck_state_what_the_price_implies(self):
        report = _demo_report()
        paragraphs, texts, deck = _office_texts(report)
        self.assertIn(DEMO_REVERSE, paragraphs)
        # One bullet in the memo, among the key assumptions; no terminal check
        # for the demo (its 37% implied RONIC raises no note).
        self.assertEqual(sum("Reverse DCF" in p for p in paragraphs), 1)
        self.assertFalse([p for p in paragraphs if p.startswith("Terminal check")])
        # The deck says it on the valuation summary slide, under the table.
        summary = _slide(deck, "Valuation summary")
        boxes = [s for s in summary.shapes if s.has_text_frame and "Reverse DCF" in s.text]
        self.assertEqual([b.text for b in boxes], [DEMO_REVERSE])
        table = next(s for s in summary.shapes if s.has_table)
        self.assertGreater(boxes[0].top, table.top + table.height)
        self.assertLessEqual(boxes[0].top + boxes[0].height, deck.slide_height)
        self.assertEqual(sum("Reverse DCF" in t for t in texts), 1)

    def test_a_flagged_issuer_gets_the_note_and_no_implied_figures(self):
        from tests.test_reports_cli import _bank_report, _lessor_report

        for build, whom in ((_bank_report, "a financial institution"),
                            (_lessor_report, "a debt-funded lessor")):
            with self.subTest(whom):
                paragraphs, texts, _deck = _office_texts(build())
                line = (f"Reverse DCF — Not meaningful for {whom}: the FCFF DCF is "
                        "reference-only.")
                self.assertIn(line, paragraphs)
                self.assertIn(line, texts)
                for text in paragraphs + texts:
                    self.assertNotIn("price implies", text)

    def test_an_unsolved_driver_is_left_out(self):
        from backend.exports import _reverse_dcf_text

        reverse = dict(_demo_report().reverse_dcf)
        reverse["margin"] = dict(reverse["margin"], converged=False,
                                 implied_target_ebit_margin=None, note="No margin fits.")
        self.assertEqual(_reverse_dcf_text(reverse),
                         "Reverse DCF (other inputs held) — The price implies 17.5% year-1 "
                         "growth (model 8.0%) or an 8.2% discount rate (model 9.5%).")
        self.assertEqual(_reverse_dcf_text(None), "Reverse DCF — Not available for this run.")

    def test_none_solved_gives_each_drivers_own_reason(self):
        from backend.exports import _reverse_dcf_text

        # A flat curve (every growth rate reproduces the price), a margin with
        # no solution and a discount rate that could not be computed: each
        # driver keeps its own reason, never 'no value reproduces the price'.
        flat = ("The model matches the market price throughout the tested year-1 growth "
                "range; year-1 growth cannot be uniquely inferred.")
        outside = ("Market price is outside the solvable target EBIT margin range (5.0% to "
                   "60.0%) with the current assumptions.")
        reverse = dict(_demo_report().reverse_dcf, converged=False, implied_growth_y1=None,
                       note=flat, wacc=None)
        reverse["margin"] = dict(reverse["margin"], converged=False,
                                 implied_target_ebit_margin=None, note=outside)
        self.assertEqual(_reverse_dcf_text(reverse),
                         "Reverse DCF (other inputs held) — No driver was solved. Year-1 "
                         f"revenue growth: {flat} Target EBIT margin: {outside} Discount rate "
                         "(WACC): could not be computed.")
        # The growth leg the engine could not re-run, the others unsolved too.
        reverse.update(note="The DCF could not be re-run across the growth range.",
                       wacc={"converged": False, "implied_wacc": None, "current_wacc": 0.095,
                             "note": "No discount rate to test"})
        text = _reverse_dcf_text(reverse)
        self.assertIn("Year-1 revenue growth: The DCF could not be re-run across the growth "
                      "range.", text)
        self.assertTrue(text.endswith("Discount rate (WACC): No discount rate to test."), text)
        self.assertNotIn("reproduces the price", text)

    def test_the_terminal_check_restates_the_dcf_note(self):
        from equity_valuation.data.synthetic import make_company
        from tests.test_terminal_value import _company

        class _Fixed(SyntheticProvider):
            def __init__(self, company):
                super().__init__()
                self.company = company

            def get_company_data(self, ticker):
                return self.company

        # Capex barely above D&A: growth on almost no reinvestment.
        report = value_company(DEMO_TICKER, provider=_Fixed(_company(capex_share=0.041)),
                               peers=DEMO_PEERS, run_sensitivity=False)
        a = report.dcf.assumptions
        self.assertTrue(any("implied return on new invested capital" in n for n in a["notes"]))
        line = (f"Terminal check — Terminal growth of 2.5% on net reinvestment of "
                f"{a['terminal_reinvestment_rate'] * 100:.1f}% of NOPAT implies a "
                f"{a['implied_terminal_ronic'] * 100:.0f}% return on new invested capital "
                f"(latest ROIC {a['latest_roic'] * 100:.1f}%).")
        paragraphs, texts, _deck = _office_texts(report)
        self.assertIn(line, paragraphs)
        self.assertTrue(any(line in t for t in texts))
        # Capex below D&A: growth without reinvestment.
        report = value_company(DEMO_TICKER, provider=_Fixed(_company(capex_share=0.036)),
                               peers=DEMO_PEERS, run_sensitivity=False)
        rate = report.dcf.assumptions["terminal_reinvestment_rate"]
        self.assertLess(rate, 0)
        paragraphs, _texts, _deck = _office_texts(report)
        self.assertIn(f"Terminal check — Terminal growth of 2.5% with net reinvestment of "
                      f"{rate * 100:.1f}% of NOPAT implies growth without reinvestment.",
                      paragraphs)
        # No DCF, no lines.
        from backend.exports import _valuation_checks

        self.assertEqual(_valuation_checks(value_company(
            DEMO_TICKER, provider=_Fixed(make_company()), run_dcf=False,
            run_sensitivity=False)), [])

    def test_articles_follow_how_the_figure_is_read(self):
        from backend.exports import _article

        for figure, want in (("8.2%", "an"), ("80.0%", "an"), ("11.0%", "an"),
                             ("18.5%", "an"), ("830%", "an"), ("30.8%", "a"), ("1.5%", "a"),
                             ("110.0%", "a"), ("-5.0%", "a"), ("371%", "a"),
                             # eleven / eighteen thousand (million), but one thousand ...
                             ("11500%", "an"), ("11,500%", "an"), ("18000%", "an"),
                             ("11000000%", "an"), ("1100%", "a"), ("1800%", "a"),
                             ("180%", "a"), ("111%", "a"), ("8000%", "an")):
            self.assertEqual(_article(figure), want, figure)

    def test_office_notes_do_not_name_the_command_line_flag(self):
        from tests.test_reports_cli import _kind_report, _lessor_report

        for report in (_lessor_report(), _kind_report("lessor", run_comps=True)):
            self.assertTrue([w for w in report.warnings if "--peers" in w])
            paragraphs, texts, _deck = _office_texts(report)
            joined = "\n".join(paragraphs + texts)
            self.assertNotIn("--peers", joined)
            self.assertIn("there are no comps. Supply peers to get a target.", joined)
        self.assertIn("Comps: no usable peers found (supply peers to add trading comps).",
                      joined)

    def test_dashboard_excel_and_html_do_not_name_the_command_line_flag(self):
        from openpyxl import load_workbook

        from backend import exports
        from equity_valuation.report.excel import write_excel
        from tests.test_reports_cli import _kind_report, _lessor_report

        for report in (_lessor_report(), _kind_report("lessor", run_comps=True)):
            before = list(report.warnings)
            self.assertTrue([w for w in before if "--peers" in w])
            with tempfile.TemporaryDirectory() as tmp, \
                    patch.object(exports, "run_valuation_report", return_value=(report, {})), \
                    patch.object(exports, "_ensure_out", lambda: Path(tmp)):
                wb = load_workbook(exports.export_excel(DEMO_TICKER, {}))
                html = Path(exports.export_html(DEMO_TICKER, {})).read_text("utf-8")
                cli = load_workbook(write_excel(report, os.path.join(tmp, "cli.xlsx")))
            cells = "\n".join(c.value for ws in wb for row in ws.iter_rows() for c in row
                              if isinstance(c.value, str))
            for text in (cells, html):
                self.assertNotIn("--peers", text)
            self.assertEqual(report.warnings, before)  # the snapshot is not changed
            # The command line's own workbook keeps its wording.
            self.assertIn("--peers", "\n".join(
                c.value for ws in cli for row in ws.iter_rows() for c in row
                if isinstance(c.value, str)))
        self.assertIn("Comps: no usable peers found (supply peers to add trading comps).",
                      cells)
        self.assertIn("Supply peers to get a target.", html)

    def test_memo_names_the_value_driver_ronic(self):
        from equity_valuation.schemas import DCFAssumptions

        report = value_company(DEMO_TICKER, provider=SyntheticProvider(), peers=DEMO_PEERS,
                               dcf_assumptions=DCFAssumptions(terminal_method="value_driver",
                                                              terminal_ronic=0.15),
                               run_sensitivity=False)
        paragraphs, _texts, _deck = _office_texts(report)
        horizon = next(p for p in paragraphs if p.startswith("Forecast horizon"))
        self.assertTrue(horizon.endswith(
            "terminal growth 2.5% · terminal method: Value driver (RONIC 15.0%)"), horizon)
        self.assertNotIn("value_driver", "\n".join(paragraphs))
        # A Gordon run keeps the method's name in brackets.
        paragraphs, _texts, _deck = _office_texts(_demo_report())
        horizon = next(p for p in paragraphs if p.startswith("Forecast horizon"))
        self.assertTrue(horizon.endswith("terminal growth 2.5% (gordon)"), horizon)

    def test_demo_exports_keep_the_pinned_figures(self):
        from openpyxl import load_workbook

        from equity_valuation.report.excel import write_excel
        from equity_valuation.report.html import write_html

        report = _demo_report()
        memo, deck = _office(report)
        rows = {r.cells[0].text: (r.cells[1].text, r.cells[2].text) for r in memo.tables[0].rows}
        self.assertEqual(rows["DCF"], ("$33.38", "-18.3%"))
        self.assertEqual(rows["Comps (median)"], ("$46.96", "+15.0%"))
        self.assertEqual(rows["DDM"], ("$10.99", "-73.1%"))
        self.assertEqual(rows["FCFE"], ("$31.41", "-23.1%"))
        self.assertEqual(rows["Blended target"], ("$32.40", "-20.7%"))
        table = next(s.table for s in deck.slides[1].shapes if s.has_table)
        self.assertEqual([table.cell(5, c).text for c in range(3)],
                         ["Blended target", "$32.40", "-20.7%"])
        with tempfile.TemporaryDirectory() as tmp:
            wb = load_workbook(write_excel(report, os.path.join(tmp, "d.xlsx")))
            ws = wb["Summary"]
            values = {}  # the first row of each label: the method table's
            for r in range(1, ws.max_row + 1):
                values.setdefault(ws[f"A{r}"].value, ws[f"B{r}"].value)
            for label, pinned in (("DCF (FCFF)", 33.38), ("Trading comps (median)", 46.96),
                                  ("DDM", 10.99), ("FCFE", 31.41),
                                  ("Blended target (median)", 32.40)):
                self.assertAlmostEqual(values[label], pinned, delta=0.005, msg=label)
            self.assertAlmostEqual(values["Market-implied year-1 growth"], 0.175, delta=0.0005)
            html = Path(write_html(report, os.path.join(tmp, "d.html"))).read_text("utf-8")
            self.assertIn("Blended target</div><div class='card-value neutral'>$32.40<", html)
            self.assertIn("Year-1 revenue growth</th><td>8.0%</td><td>17.5%</td>", html)


class ReverseDCFTests(unittest.TestCase):
    def test_solved_growth_reprices_to_market(self):
        from backend import valuation_service as vs
        from equity_valuation.models.dcf import run_dcf

        macro, dcf_a, ddm_a, _peers, _toggles, _echo = vs.parse_assumptions({})
        report = value_company(
            DEMO_TICKER, provider=SyntheticProvider(), macro=macro,
            dcf_assumptions=dcf_a, ddm_assumptions=ddm_a, peers=DEMO_PEERS,
        )
        result = vs._reverse_dcf(report, dcf_a, macro)
        self.assertTrue(result["converged"])
        g1 = result["implied_growth_y1"]
        # The synthetic firm grew 8% a year; a 20x P/E price needs more than that.
        self.assertGreater(g1, 0.08)

        path = vs._revenue_growth_path(g1, dcf_a.terminal_growth, dcf_a.forecast_years)
        solved = dataclasses.replace(dcf_a, revenue_growth=path)
        price = run_dcf(report.company, macro, solved, report.current_price).implied_price
        self.assertAlmostEqual(price, report.current_price, delta=0.01)


if __name__ == "__main__":
    unittest.main()
