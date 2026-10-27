"""Office artifact generation — the analyst-facing outputs.

  * Excel model  -> engine's write_excel (live formulas, multi-tab)   [.xlsx]
  * HTML report  -> engine's write_html (interactive, offline)        [.html]
  * Research memo-> python-docx; model summary + the AI research note [.docx]
  * Briefing deck-> python-pptx; valuation summary, football field,
                    comps, thesis/risks                               [.pptx]

The memo/deck render fine with NO AI note (model-only); when a generated
research note dict is supplied they include the full cited write-up. Both
state what the market price implies (the engine's reverse DCF: year-1 growth,
target EBIT margin and discount rate, one at a time) and, when the DCF flagged
its Gordon terminal value, the terminal reinvestment check (the implied return
on new invested capital, or growth without net reinvestment).

All four come from the dashboard, so the engine notes they carry name no
command-line flag ('Supply peers with --peers ...' reads 'Supply peers ...');
the Excel and HTML reports the CLI writes keep its wording.
"""

from __future__ import annotations

import dataclasses
import os
import re
import textwrap
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Optional

from equity_valuation.report.excel import (
    currency_prefix,
    ddm_method_label,
    no_target_text,
    reverse_dcf_message,
    reverse_dcf_rows,
    terminal_method_label,
)
from equity_valuation.schemas import DDMAssumptions
from equity_valuation.utils import available_model_price, is_num

from .valuation_service import get_report_snapshot, run_valuation_report

_OUT_DIR = Path(__file__).resolve().parent.parent / "output"

_AUTHOR = "Equity Research Automation"

# The DDM drivers' defaults: the memo quotes the DDM's inputs only when the
# request changed one (``_ddm_text``).
_DDM_DEFAULTS = DDMAssumptions()

# Characters XML 1.0 forbids. python-docx refuses any string holding one (the
# memo export would fail) and python-pptx writes it as a visible '_x001B_'.
_XML_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff\ud800-\udfff]")
# The forbidden characters that str.splitlines() reads as line breaks (vertical
# tab, Word's manual line break; form feed; the file, group and record
# separators). They become a line feed, so the words either side stay apart and
# the memo breaks the line where the deck (which splits on them) does.
_XML_LINE_BREAKS = re.compile(r"[\x0b\x0c\x1c-\x1e]")


def _xml_text(s) -> str:
    """Provider, engine or AI text made safe for a Word or PowerPoint file: a
    forbidden line-break character becomes a line feed, any other is removed."""
    return _XML_ILLEGAL.sub("", _XML_LINE_BREAKS.sub("\n", str(s)))


# Engine notes written for the command line name its --peers flag. Every
# export here comes from the dashboard, so they say 'supply peers' instead, the
# reports' wording (``no_target_text``: 'No target (supply peers)').
_CLI_PEERS_PHRASES = (("pass --peers to supply them", "supply peers to add trading comps"),
                      (" with --peers", ""))


def _dashboard_note(s) -> str:
    """An engine note in the dashboard's wording: the command-line peers
    phrasing reworded ('... Supply peers with --peers to get a target.' ->
    '... Supply peers to get a target.')."""
    text = str(s)
    for cli, plain in _CLI_PEERS_PHRASES:
        text = text.replace(cli, plain)
    return text


def _note_text(s) -> str:
    """An engine note for the memo or deck: ``_dashboard_note``, then made
    XML-safe (``_xml_text``)."""
    return _xml_text(_dashboard_note(s))


def _dashboard_report(report):
    """The report for the Excel and HTML downloads: a copy whose warnings are
    in the dashboard's wording (``_dashboard_note``), so they name no CLI flag.
    The cached snapshot itself is never changed; a report with nothing to
    reword (or not a dataclass) comes back as is."""
    warnings = list(getattr(report, "warnings", None) or [])
    reworded = [_dashboard_note(w) for w in warnings]
    if reworded == warnings or not dataclasses.is_dataclass(report):
        return report
    return dataclasses.replace(report, warnings=reworded)


def _money(x, cur="USD", dec=2) -> str:
    """'-$6.52': the minus sign before the currency prefix (``currency_prefix``:
    the symbol, or the ISO code for an unmapped currency); no sign on a value
    that rounds to zero."""
    if not is_num(x):
        return "n/a"
    x = round(x, dec)
    sign = "-" if x < 0 else ""
    return f"{sign}{currency_prefix(cur)}{abs(x):,.{dec}f}"


def _cap(x) -> str:
    """Compact market cap for the comps table: 2,345B / 3.4B / 300M."""
    if not is_num(x) or x == 0:
        return "n/a"
    if abs(x) >= 1e10:
        return f"{x / 1e9:,.0f}B"
    if abs(x) >= 1e9:
        return f"{x / 1e9:.1f}B"
    if abs(x) >= 1e6:
        return f"{x / 1e6:,.0f}M"
    return f"{x:,.0f}"


def _set_core_properties(props, title: str) -> None:
    """Replace the python-docx/python-pptx template metadata (third-party
    author, 2013 dates, 'generated using ...' comment) with our own."""
    now = datetime.now(timezone.utc).replace(microsecond=0)
    props.title = _xml_text(title)
    props.author = _AUTHOR
    props.last_modified_by = _AUTHOR
    props.created = now
    props.modified = now
    props.comments = ""
    props.subject = ""
    props.keywords = ""
    props.category = ""
    props.revision = 1


def _pct(x, signed=False) -> str:
    if not is_num(x):
        return "n/a"
    s = "+" if (signed and x > 0) else ""
    return f"{s}{x * 100:.1f}%"


def _method_upside(value, price) -> str:
    return _pct(value / price - 1, signed=True) if (
        is_num(value) and is_num(price) and price > 0
    ) else "n/a"


def _text_pages(items, *, width=95, lines_per_page=15):
    """Bound slide text by wrapped lines, retaining every item and paragraph.

    Long individual items split across slides instead of being truncated.
    A conservative width leaves room for proportional fonts and margins. The
    wrap only sizes the pages: a page has one line per paragraph, which the
    slide wraps itself, and a blank line between items. A paragraph that runs
    onto the next page is split between two wrapped lines, keeping its text.
    """
    if isinstance(items, str):
        items = [items]
    lines = []  # (paragraph no., its text, start, end) per wrapped line; None = blank
    for item in items or []:
        for paragraph in str(item).splitlines() or [""]:
            text = " ".join(paragraph.split())
            wrapped = textwrap.wrap(text, width=width)
            pid, pos = len(lines), 0  # pid: a number unique to this paragraph
            for ln in wrapped:
                start = text.index(ln, pos)  # each wrapped line is a slice of text
                pos = start + len(ln)
                lines.append((pid, text, start, pos))
            if not wrapped:
                lines.append(None)
        lines.append(None)
    while lines and lines[-1] is None:
        lines.pop()
    pages = []
    for i in range(0, len(lines), lines_per_page):
        paragraphs = []  # [paragraph no., text, start, end] or None
        for line in lines[i:i + lines_per_page]:
            last = paragraphs[-1] if paragraphs else None
            if line is not None and last is not None and last[0] == line[0]:
                last[3] = line[3]  # the same paragraph: one slide paragraph
            else:
                paragraphs.append(None if line is None else list(line))
        pages.append("\n".join("" if p is None else p[1][p[2]:p[3]] for p in paragraphs))
    return pages


def _article(figure: str) -> str:
    """'a' or 'an' before a figure as it is read aloud: 'an 8.2%', 'an 11.0%',
    'an 18.5%', 'an 80.0%', 'an 11,500%' (eleven thousand), but 'a 30.8%',
    'a 1.5%', 'a 110%' (one hundred) and 'a -5.0%' (minus). A whole part
    starting with 8 reads 'eight...'; one whose leading group of digits is 11
    or 18 (11, 11,000, 18,000,000) reads 'eleven' or 'eighteen'."""
    whole = figure.lstrip("+").split(".", 1)[0].rstrip("%").replace(",", "")
    leading_pair = len(whole) % 3 == 2 and whole[:2] in ("11", "18")
    return "an" if whole.startswith("8") or leading_pair else "a"


# The reason given for a reverse-DCF driver the engine could not compute (its
# leg is None).
_UNCOMPUTED_LEG = "could not be computed."


def _sentence(text: str) -> str:
    """``text`` ending in a full stop."""
    text = str(text).strip()
    return text if text.endswith((".", "!", "?")) else text + "."


# How the reverse-DCF sentence names each driver (``excel.REVERSE_DCF_DRIVERS``
# labels): (noun, whether it takes an article).
_REVERSE_DCF_NOUNS = {"Year-1 revenue growth": ("year-1 growth", False),
                      "Target EBIT margin": ("target EBIT margin", True),
                      "Discount rate (WACC)": ("discount rate", True)}


def _reverse_dcf_text(reverse) -> str:
    """The reverse DCF in one line, built from the engine's
    ``report.reverse_dcf`` as the Excel and HTML reports tabulate it: 'Reverse
    DCF (other inputs held) — The price implies 17.5% year-1 growth (model
    8.0%), a 30.8% target EBIT margin (model 25.0%) or an 8.2% discount rate
    (model 9.5%).' (the drivers whose solve converged); when none converged,
    'No driver was solved.' and each driver's own reason (the solver's note,
    or that it could not be computed); or the engine's single note (e.g. 'Not
    meaningful for a financial institution: the FCFF DCF is reference-only.')."""
    message = reverse_dcf_message(reverse)
    if message is not None:
        return f"Reverse DCF — {message}"
    legs = []
    for row in reverse_dcf_rows(reverse):
        if not row["converged"]:
            continue
        noun, article = _REVERSE_DCF_NOUNS.get(row["label"], (row["label"].lower(), False))
        figure = _pct(row["implied"])
        model = f" (model {_pct(row['model'])})" if is_num(row["model"]) else ""
        legs.append(f"{_article(figure) + ' ' if article else ''}{figure} {noun}{model}")
    if not legs:
        # Nothing solved: each driver's own reason (no value in the range
        # reproduces the price, the model matches it across the whole range, or
        # the leg could not be computed), as the Excel and HTML tables note it.
        reasons = " ".join(f"{row['label']}: {_sentence(row['note'] or _UNCOMPUTED_LEG)}"
                           for row in reverse_dcf_rows(reverse))
        return f"Reverse DCF (other inputs held) — No driver was solved. {reasons}"
    joined = legs[0] if len(legs) == 1 else ", ".join(legs[:-1]) + " or " + legs[-1]
    return f"Reverse DCF (other inputs held) — The price implies {joined}."


# The Gordon terminal-value notes of ``models/dcf.py`` that the terminal check
# restates (growth with no net reinvestment; an implied RONIC well above the
# latest ROIC). The DCF records them in ``assumptions['notes']``.
_NO_REINVESTMENT_NOTE = "implies growth without reinvestment"
_HIGH_RONIC_NOTE = "implied return on new invested capital"


def _terminal_check_text(dcf) -> Optional[str]:
    """When the DCF's terminal reinvestment note fired, one line with its
    figures: 'Terminal check — Terminal growth of 2.5% on net reinvestment of
    0.7% of NOPAT implies a 371% return on new invested capital (latest ROIC
    29.9%).', or that the growth comes without reinvestment. None otherwise
    (the note never fires for a flagged issuer or a non-Gordon terminal value)."""
    a = getattr(dcf, "assumptions", None) or {}
    notes = [str(n) for n in (a.get("notes") or [])]
    g = _pct(a.get("terminal_growth_used", a.get("terminal_growth")))
    rate = _pct(a.get("terminal_reinvestment_rate"))
    if any(_NO_REINVESTMENT_NOTE in n for n in notes):
        return (f"Terminal check — Terminal growth of {g} with net reinvestment of {rate} "
                "of NOPAT implies growth without reinvestment.")
    ronic = a.get("implied_terminal_ronic")
    if is_num(ronic) and any(_HIGH_RONIC_NOTE in n for n in notes):
        figure = f"{ronic * 100:.0f}%"
        return (f"Terminal check — Terminal growth of {g} on net reinvestment of {rate} of "
                f"NOPAT implies {_article(figure)} {figure} return on new invested capital "
                f"(latest ROIC {_pct(a.get('latest_roic'))}).")
    return None


def _ddm_text(report, echo: dict) -> Optional[str]:
    """The memo's dividend discount model line when the request set a DDM
    driver (``ddm_method``, ``ddm_high_growth_rate`` or
    ``ddm_high_growth_years``), quoting what the DDM applied: 'Dividend
    discount model: H-model, high growth 10.0% over 6 years'. None when the
    DDM did not run or kept its defaults, as the other optional inputs are
    listed only when set."""
    ddm = getattr(report, "ddm", None)
    if ddm is None or (echo.get("ddm_method") in (None, _DDM_DEFAULTS.method)
                       and echo.get("ddm_high_growth_rate") is None
                       and echo.get("ddm_high_growth_years") in (
                           None, _DDM_DEFAULTS.high_growth_years)):
        return None
    detail = getattr(ddm, "detail", None) or {}
    method = getattr(ddm, "method", None)
    text = f"Dividend discount model: {ddm_method_label(method)}"
    years = detail.get("high_growth_years")
    if not is_num(years) and is_num(detail.get("H_half")):
        years = 2 * detail["H_half"]  # the H-model keeps half the horizon
    if method in ("two_stage", "h_model") and is_num(detail.get("high_growth")) \
            and is_num(years):
        unit = "year" if years == 1 else "years"
        text += f", high growth {_pct(detail['high_growth'])} over {years:g} {unit}"
    return text


def _valuation_checks(report) -> list:
    """The memo's and deck's lines on what the price implies (the reverse DCF)
    and, when the DCF flagged it, the terminal reinvestment check; [] without
    a DCF."""
    if getattr(report, "dcf", None) is None:
        return []
    lines = [_reverse_dcf_text(getattr(report, "reverse_dcf", None))]
    check = _terminal_check_text(report.dcf)
    if check:
        lines.append(check)
    return [_xml_text(line) for line in lines]


def _blended_target(summary: dict, cur: str, comps=None) -> str:
    """The blended target, or why there is none (never a made-up figure):
    'No target (supply peers)' or 'n/a', by the rule the Excel and HTML
    reports share (``no_target_text``; ``comps`` is the report's CompsResult)."""
    if is_num(summary.get("blended_target")):
        return _money(summary["blended_target"], cur)
    return no_target_text(summary, comps)


def _blended_upside(summary: dict) -> str:
    """The engine's blended upside; 'n/a' without a target, or when the
    verdict is withheld (the engine then returns None)."""
    if not is_num(summary.get("blended_target")):
        return "n/a"
    return _pct(summary.get("blended_upside"), signed=True)


def _headline(summary: dict, cur: str, comps=None) -> str:
    """'blended fair value $32.40 (-20.7%)', '... $20.87 (upside n/a)' or
    'blended fair value: No target (supply peers)'."""
    target = _blended_target(summary, cur, comps)
    if not is_num(summary.get("blended_target")):
        return f"blended fair value: {target}"
    upside = _blended_upside(summary)
    return f"blended fair value {target} ({'upside n/a' if upside == 'n/a' else upside})"


def _method_label(summary: dict, name: str) -> str:
    """A method's table label, marked when the engine left it out of the blended
    target (e.g. a bank's DCF and FCFE, shown for reference only)."""
    excluded = summary.get("excluded_from_blend") or {}
    return f"{name} (not in blend)" if name in excluded else name


def _method_values(report) -> dict:
    methods = dict(report.summary.get("methods") or {})
    for label, name in (("DCF", "dcf"), ("DDM", "ddm"), ("FCFE", "fcfe")):
        if label in methods:
            methods[label] = available_model_price(getattr(report, name, None))
    return methods


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]", "_", name)


def _ensure_out() -> Path:
    os.makedirs(_OUT_DIR, exist_ok=True)
    return _OUT_DIR


def _export_report(ticker: str, payload: dict):
    if "report_id" in payload:
        return get_report_snapshot(ticker, payload["report_id"])
    return run_valuation_report(ticker, payload)


# --------------------------------------------------------------------------- #
#  Engine-native exports
# --------------------------------------------------------------------------- #
def export_excel(ticker: str, payload: dict) -> str:
    from equity_valuation.report.excel import write_excel

    report, _ = _export_report(ticker, payload)
    path = _ensure_out() / f"{_safe(ticker.upper())}_valuation.xlsx"
    return write_excel(_dashboard_report(report), str(path))


def export_html(ticker: str, payload: dict) -> str:
    from equity_valuation.report.html import write_html

    report, _ = _export_report(ticker, payload)
    path = _ensure_out() / f"{_safe(ticker.upper())}_valuation.html"
    return write_html(_dashboard_report(report), str(path))


# --------------------------------------------------------------------------- #
#  Research memo (.docx)
# --------------------------------------------------------------------------- #
def export_memo(ticker: str, payload: dict, note: Optional[dict] = None) -> str:
    from docx import Document
    from docx.shared import Pt, RGBColor

    report, echo = _export_report(ticker, payload)
    s = report.summary
    cur = s.get("currency", "USD")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    # Provider, engine and AI text passes through _xml_text: python-docx refuses
    # a string with a control character (e.g. from exception text in a warning).
    title = _xml_text(note.get("title")) if note and note.get("title") else None
    doc.add_heading(
        title or _xml_text(f"{s.get('name')} ({s.get('ticker')}) — Research Memo"), level=0
    )
    meta = doc.add_paragraph()
    run = meta.add_run(_xml_text(
        f"Prepared {date.today().isoformat()} · price {_money(s.get('current_price'), cur)} · "
        f"{_headline(s, cur, report.comps)} · model verdict: {s.get('recommendation') or 'n/a'}"
    ))
    run.font.color.rgb = RGBColor(0x60, 0x70, 0x8A)
    if report.valued_at:
        doc.add_paragraph(_xml_text(f"Valuation calculated: {report.valued_at}"))

    if note and note.get("stance"):
        doc.add_paragraph(_xml_text(f"Analytical stance: {note['stance']}")).runs[0].bold = True

    # --- valuation summary table -------------------------------------------- #
    doc.add_heading("Valuation summary", level=1)
    methods = _method_values(report)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Method", "Implied value", "Vs. price"
    price = s.get("current_price")
    for name, val in methods.items():
        row = table.add_row().cells
        row[0].text = _method_label(s, name)
        row[1].text = _money(val, cur)
        row[2].text = _method_upside(val, price)
    row = table.add_row().cells
    row[0].text = "Blended target"
    row[1].text = _blended_target(s, cur, report.comps)
    row[2].text = _blended_upside(s)
    excluded = s.get("excluded_from_blend") or {}
    if excluded:
        doc.add_paragraph(_xml_text(
            "Not in the blended target: "
            + "; ".join(f"{name} ({why})" for name, why in excluded.items()) + "."
        ))

    # --- key assumptions ------------------------------------------------------ #
    doc.add_heading("Key model assumptions", level=1)
    wacc = report.dcf.wacc.wacc if report.dcf else None
    dcf_used = (report.dcf.assumptions or {}) if report.dcf else {}
    # The method the DCF actually applied (it falls back to Gordon when no
    # exit multiple is given).
    method = dcf_used.get("terminal_method") or echo.get("terminal_method")
    # The growth the DCF actually used; it clamps the input below WACC.
    g_used = dcf_used.get("terminal_growth_used", echo.get("terminal_growth"))
    # As the Excel and HTML reports name it: a value-driver run with its RONIC.
    method_text = terminal_method_label({**dcf_used, "terminal_method": method})
    horizon = (
        f"Forecast horizon: {echo.get('forecast_years')} years · terminal growth "
        f"{_pct(g_used)} "
        + (f"· terminal method: {method_text}" if method == "value_driver"
           else f"({method_text})")
    )
    exit_mult = dcf_used.get("exit_ev_ebitda", echo.get("exit_ev_ebitda"))
    if method == "exit_multiple" and is_num(exit_mult):
        # The multiple alone sets the DCF terminal value; the growth rate still
        # ends the revenue-growth fade and drives DDM and FCFE.
        horizon = (
            f"Forecast horizon: {echo.get('forecast_years')} years · terminal value at "
            f"exit EV/EBITDA {exit_mult:.1f}x · terminal growth {_pct(g_used)} "
            "(revenue fade, DDM and FCFE; not the DCF terminal value)"
        )
    bullets = [
        f"WACC: {_pct(wacc)} · risk-free {_pct(echo.get('rf'))} · ERP {_pct(echo.get('erp'))}",
        horizon,
    ]
    if echo.get("revenue_growth_y1") is not None:
        bullets.append(f"Year-1 revenue growth: {_pct(echo.get('revenue_growth_y1'))} (fading to terminal)")
    if echo.get("target_ebit_margin") is not None:
        bullets.append(f"Terminal EBIT margin: {_pct(echo.get('target_ebit_margin'))}")
    ddm_text = _ddm_text(report, echo)
    if ddm_text:
        bullets.append(ddm_text)
    # What the price implies (reverse DCF) and the terminal reinvestment check.
    bullets.extend(_valuation_checks(report))
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")

    # --- AI research note sections -------------------------------------------- #
    def _section(heading: str, items, bullet=True):
        if not items:
            return
        doc.add_heading(heading, level=1)
        if isinstance(items, str):
            doc.add_paragraph(_xml_text(items))
        else:
            for it in items:
                doc.add_paragraph(_xml_text(it), style="List Bullet" if bullet else None)

    if note:
        _section("Executive summary", note.get("executive_summary"))
        _section("Thesis", note.get("thesis"))
        _section("Valuation view", note.get("valuation_view"))
        _section("Key drivers", note.get("key_drivers"))
        _section("Risks", note.get("risks"))
        _section("Red flags / diligence items", note.get("red_flags"))
        _section("Catalysts", note.get("catalysts"))
        _section("What would change my mind", note.get("what_would_change_my_mind"))
        cites = note.get("citations") or []
        if cites:
            doc.add_heading("Sources", level=1)
            for c in cites:
                doc.add_paragraph(
                    _xml_text(f"{c.get('source')}: {c.get('note')}"), style="List Bullet"
                )

    # --- footnotes -------------------------------------------------------------- #
    if report.warnings:
        doc.add_heading("Model notes", level=1)
        for w in report.warnings:
            doc.add_paragraph(_note_text(w), style="List Bullet")
    p = doc.add_paragraph()
    r = p.add_run(
        "Generated by Equity Research Automation. Model estimates driven by user "
        "assumptions — research support, not investment advice."
    )
    r.italic = True
    r.font.size = Pt(8.5)

    _set_core_properties(
        doc.core_properties,
        title or f"{s.get('name')} ({s.get('ticker')}) — Research Memo",
    )
    path = _ensure_out() / f"{_safe(ticker.upper())}_research_memo.docx"
    doc.save(str(path))
    return str(path)


# --------------------------------------------------------------------------- #
#  Briefing deck (.pptx)
# --------------------------------------------------------------------------- #
def export_deck(ticker: str, payload: dict, note: Optional[dict] = None) -> str:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    report, _echo = _export_report(ticker, payload)
    s = report.summary
    cur = s.get("currency", "USD")
    price = s.get("current_price")

    INK = RGBColor(0x12, 0x18, 0x26)
    DIM = RGBColor(0x5C, 0x6B, 0x84)
    BLUE = RGBColor(0x2F, 0x6B, 0xD8)
    GREEN = RGBColor(0x1F, 0x9D, 0x6F)
    ROSE = RGBColor(0xD6, 0x45, 0x61)
    AMBER = RGBColor(0xC8, 0x8A, 0x1A)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def add_slide():
        return prs.slides.add_slide(blank)

    def text(slide, l, t, w, h, s_, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        # One paragraph per line: a line feed inside a run is not a line break
        # in a slide's XML, so renderers would run the items together.
        for i, line in enumerate(_xml_text(s_).splitlines() or [""]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            if not line:  # a blank separator line, the height of the text
                p._p.get_or_add_endParaRPr().sz = int(size * 100)
                continue
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
        return box

    # --- slide 1: title --------------------------------------------------------- #
    sl = add_slide()
    # Room for two title lines above the subtitle, the text anchored to the
    # bottom: a long name (e.g. 'INTERNATIONAL BUSINESS MACHINES CORP') wraps
    # upward instead of over the subtitle. An even longer one is set smaller.
    title = f"{s.get('name')}  ({s.get('ticker')})"
    title_box = text(sl, 0.7, 1.6, 11.9, 1.6, title, 40 if len(title) <= 64 else 30, True)
    title_box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    title_box.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    text(sl, 0.7, 3.3, 11.9, 0.6,
         f"Valuation briefing · {date.today().isoformat()}"
         + (f" · Calculated {report.valued_at}" if report.valued_at else ""), 18, False, DIM)
    verdict = s.get("recommendation") or "n/a"
    # No colour when no verdict is given (no target, or a withheld verdict).
    vcolor = {"Undervalued": GREEN, "Overvalued": ROSE, "Fairly valued": AMBER}.get(verdict, DIM)
    text(sl, 0.7, 4.2, 11.9, 0.6,
         f"Price {_money(price, cur)} · {_headline(s, cur, report.comps)} · {verdict}",
         18, True, vcolor)

    # --- slide 2: method summary ------------------------------------------------- #
    sl = add_slide()
    text(sl, 0.7, 0.4, 12, 0.6, "Valuation summary", 24, True)
    methods = list(_method_values(report).items())
    rows = len(methods) + 2
    tbl = sl.shapes.add_table(rows, 3, Inches(0.7), Inches(1.2), Inches(8.0),
                              Inches(0.4 * rows)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text, tbl.cell(0, 2).text = (
        "Method", "Implied value", "Vs. price")
    for i, (name, val) in enumerate(methods, start=1):
        tbl.cell(i, 0).text = _method_label(s, name)
        tbl.cell(i, 1).text = _money(val, cur)
        tbl.cell(i, 2).text = _method_upside(val, price)
    tbl.cell(rows - 1, 0).text = "Blended target"
    tbl.cell(rows - 1, 1).text = _blended_target(s, cur, report.comps)
    tbl.cell(rows - 1, 2).text = _blended_upside(s)
    # Under the table: what the price implies (reverse DCF) and, when the DCF
    # flagged it, the terminal reinvestment check.
    checks = _valuation_checks(report)
    if checks:
        top = 1.2 + 0.4 * rows + 0.35
        text(sl, 0.7, top, 11.9, 7.5 - top - 0.3, "\n\n".join(checks), 14)

    # --- slide 3: football field --------------------------------------------------- #
    rows_ff = report.football_field or []
    if rows_ff:
        sl = add_slide()
        text(sl, 0.7, 0.4, 12, 0.6, "Football field — value ranges by method", 24, True)
        price_points = [price] if is_num(price) and price > 0 else []
        lows = [r.low for r in rows_ff] + price_points
        highs = [r.high for r in rows_ff] + price_points
        vmin, vmax = min(lows), max(highs)
        pad = (vmax - vmin) * 0.06 or 1.0
        vmin -= pad
        vmax += pad
        span = vmax - vmin
        chart_l, chart_w = 2.8, 9.4
        y = 1.4
        for r in rows_ff:
            text(sl, 0.6, y - 0.06, 2.1, 0.4, r.method, 12, False, DIM)
            # The bar's range, so its values can be read off the slide.
            text(sl, 0.6, y + 0.2, 2.1, 0.3,
                 f"{_money(r.low, cur)} – {_money(r.high, cur)}", 9, False, DIM)
            x0 = chart_l + (r.low - vmin) / span * chart_w
            x1 = chart_l + (r.high - vmin) / span * chart_w
            bar = sl.shapes.add_shape(1, Inches(x0), Inches(y),
                                      Inches(max(x1 - x0, 0.05)), Inches(0.28))
            bar.fill.solid()
            bar.fill.fore_color.rgb = BLUE
            bar.line.fill.background()
            xb = chart_l + (r.base - vmin) / span * chart_w
            tick = sl.shapes.add_shape(1, Inches(xb), Inches(y - 0.05),
                                       Emu(int(914400 * 0.03)), Inches(0.38))
            tick.fill.solid()
            tick.fill.fore_color.rgb = INK
            tick.line.fill.background()
            y += 0.6
        if price_points:
            xp = chart_l + (price - vmin) / span * chart_w
            pl = sl.shapes.add_shape(1, Inches(xp), Inches(1.2),
                                     Emu(int(914400 * 0.025)), Inches(y - 1.2))
            pl.fill.solid()
            pl.fill.fore_color.rgb = ROSE
            pl.line.fill.background()
            text(sl, min(xp - 0.7, 10.4), y + 0.05, 2.2, 0.4,
                 f"price {_money(price, cur)}", 11, True, ROSE)

    # --- slide 4: comps ----------------------------------------------------------- #
    comps = report.comps
    if comps and comps.peers:
        rows_c = [comps.target] + list(comps.peers)
        for start in range(0, len(rows_c), 12):
            page = rows_c[start:start + 12]
            sl = add_slide()
            title = "Trading comps" + (" (continued)" if start else "")
            text(sl, 0.7, 0.4, 12, 0.6, title, 24, True)
            tbl = sl.shapes.add_table(
                len(page) + 1, 7, Inches(0.7), Inches(1.2), Inches(11.9),
                Inches(0.4 * (len(page) + 1))).table
            for j, h in enumerate(["Ticker", "Currency", "Mkt cap", "EV/EBITDA", "EV/Sales", "P/E", "P/B"]):
                tbl.cell(0, j).text = h
            for i, rrow in enumerate(page, start=1):
                tbl.cell(i, 0).text = rrow.ticker + (" (target)" if start == 0 and i == 1 else "")
                tbl.cell(i, 1).text = rrow.currency or "n/a"
                tbl.cell(i, 2).text = _cap(rrow.market_cap)
                for j, attr in enumerate(["ev_ebitda", "ev_sales", "pe", "pb"], start=3):
                    v = getattr(rrow, attr)
                    tbl.cell(i, j).text = f"{v:.1f}x" if is_num(v) else "n/a"
            for row in tbl.rows:
                for cell in row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(13)

    # --- research note: retain every section and source, across as many
    # slides as needed. The old fixed boxes overlapped and silently cut items.
    if note:
        sections = [("Executive summary", note.get("executive_summary")),
                    ("Thesis", note.get("thesis")),
                    ("Valuation view", note.get("valuation_view")),
                    ("Key drivers", note.get("key_drivers")),
                    ("Risks", note.get("risks")),
                    ("Red flags / diligence items", note.get("red_flags")),
                    ("Catalysts", note.get("catalysts")),
                    ("What would change my mind", note.get("what_would_change_my_mind")),
                    ("Sources", [f"{c.get('source')}: {c.get('note')}"
                                 for c in note.get("citations") or []])]
        for heading, items in sections:
            for i, page in enumerate(_text_pages(items)):
                sl = add_slide()
                text(sl, 0.7, 0.4, 12, 0.6,
                     heading + (" (continued)" if i else ""), 24, True)
                text(sl, 0.7, 1.3, 11.9, 5.6, page, 18)

    for i, page in enumerate(_text_pages([_note_text(w) for w in report.warnings or []])):
        sl = add_slide()
        text(sl, 0.7, 0.4, 12, 0.6,
             "Model notes" + (" (continued)" if i else ""), 24, True)
        text(sl, 0.7, 1.3, 11.9, 5.6, page, 18)

    # --- footer slide ------------------------------------------------------------ #
    sl = add_slide()
    text(sl, 0.7, 3.0, 11.9, 1.2,
         "Generated by Equity Research Automation — model estimates driven by user "
         "assumptions. Research support, not investment advice.", 14, False, DIM)

    _set_core_properties(
        prs.core_properties,
        f"{s.get('name')} ({s.get('ticker')}) — Valuation Briefing",
    )
    path = _ensure_out() / f"{_safe(ticker.upper())}_briefing.pptx"
    prs.save(str(path))
    return str(path)
