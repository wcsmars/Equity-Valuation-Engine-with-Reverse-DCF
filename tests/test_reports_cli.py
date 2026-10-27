"""Regression tests for the Excel/HTML reports and the command-line interface.

The Excel checks evaluate the workbook's own live formulas (with the small
evaluator below, which covers the arithmetic, SUM and IF the exporter writes)
and compare them with the model's numbers, so they hold whatever the demo
figures are. Everything runs offline on the synthetic company. Run with:
python -m unittest tests.test_reports_cli
"""

from __future__ import annotations

import contextlib
import dataclasses
import datetime
import io
import json
import os
import re
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from equity_valuation import value_company
from equity_valuation.data.synthetic import (
    DEMO_PEERS,
    DEMO_TICKER,
    SyntheticProvider,
    make_company,
)
from equity_valuation.report.excel import write_excel
from equity_valuation.report.html import _comps_table, write_html
from equity_valuation.schemas import CompRow, DCFAssumptions, SensitivityResult

MINORITY = 30e9
PREFERRED = 10e9


# --------------------------------------------------------------------------- #
#  Fixtures
# --------------------------------------------------------------------------- #
class _ClaimsProvider(SyntheticProvider):
    """The synthetic company with minority interest and preferred equity."""

    def get_company_data(self, ticker):
        company = make_company()
        company.balance_sheet.minority_interest = MINORITY
        company.balance_sheet.preferred_equity = PREFERRED
        return company


class _BankProvider(SyntheticProvider):
    """The synthetic company presented as a bank, so the engine leaves its DCF
    and FCFE out of the blended target."""

    def get_company_data(self, ticker):
        company = make_company()
        company.market.sector = "Financial Services"
        company.market.industry = "Banks - Diversified"
        return company


class _OutlierPeersProvider(SyntheticProvider):
    """Four peers, one with an outlying EV/EBITDA that the comps model trims."""

    def get_peer_comp_rows(self, tickers):
        return [
            CompRow(ticker=f"P{i}", name=f"P{i}", market_cap=5e10, enterprise_value=5.2e10,
                    ev_ebitda=m, ev_sales=4.0, pe=20.0, pb=5.0, peg=1.8)
            for i, m in enumerate((18.0, 22.0, 20.0, 90.0))
        ]


def _report(provider=None, **kwargs):
    kwargs.setdefault("peers", DEMO_PEERS)
    return value_company(DEMO_TICKER, provider=provider or SyntheticProvider(), **kwargs)


def _tmpdir(test: unittest.TestCase) -> str:
    tmp = tempfile.TemporaryDirectory()
    test.addCleanup(tmp.cleanup)
    return tmp.name


# --------------------------------------------------------------------------- #
#  A tiny evaluator for the exporter's formulas
# --------------------------------------------------------------------------- #
class _XlError(str):
    """An Excel error value such as #DIV/0!; propagates through arithmetic."""


_DIV0 = _XlError("#DIV/0!")
_VALUE = _XlError("#VALUE!")
_REF_RE = re.compile(r"(?:(\w+)!)?\$?([A-Z]{1,3})\$?(\d+)")
_TOKEN_RE = re.compile(
    r'\s*(?:(?P<str>"(?:[^"]|"")*")'
    r"|(?P<ref>(?:\w+!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)(?![\w(])"
    r"|(?P<func>[A-Z]+)\("
    r"|(?P<num>\d+(?:\.\d*)?(?:[eE][+-]?\d+)?)"
    r"|(?P<op><>|<=|>=|[-+*/=<>(),]))"
)


def _to_num(v):
    if isinstance(v, _XlError):
        return v
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    return _VALUE  # text in arithmetic


def _arith(op, a, b):
    a, b = _to_num(a), _to_num(b)
    for x in (a, b):
        if isinstance(x, _XlError):
            return x
    if op == "+":
        return a + b
    if op == "-":
        return a - b
    if op == "*":
        return a * b
    return _DIV0 if b == 0 else a / b


def _compare(op, a, b):
    for x in (a, b):
        if isinstance(x, _XlError):
            return x
    if a is None:
        a = 0.0 if not isinstance(b, str) else ""
    if b is None:
        b = 0.0 if not isinstance(a, str) else ""
    return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b,
            "<=": a <= b, ">=": a >= b}[op]


class XlEval:
    """Evaluate a workbook's formulas (numbers, refs, ranges, + - * /, SUM, IF)."""

    def __init__(self, wb):
        self.wb = wb
        self._cache: dict = {}

    def value(self, sheet: str, coord: str):
        key = (sheet, coord)
        if key not in self._cache:
            raw = self.wb[sheet][coord].value
            if isinstance(raw, str) and raw.startswith("="):
                raw = self._formula(sheet, raw[1:])
            self._cache[key] = raw
        return self._cache[key]

    def formulas(self):
        """Yield (sheet, coord, evaluated value) for every formula cell."""
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        yield ws.title, c.coordinate, self.value(ws.title, c.coordinate)

    # -- recursive-descent parser producing lazy thunks ---------------------- #
    def _formula(self, sheet, text):
        toks, pos = [], 0
        text = text.strip()
        while pos < len(text):
            m = _TOKEN_RE.match(text, pos)
            if not m or m.end() == pos:
                raise ValueError(f"cannot parse formula {text!r} at {pos}")
            kind = m.lastgroup
            toks.append((kind, m.group(kind)))
            pos = m.end()
        self._toks, self._i, self._sheet = toks, 0, sheet
        node = self._comparison()
        if self._i != len(toks):
            raise ValueError(f"trailing tokens in {text!r}")
        return node()

    def _peek(self):
        return self._toks[self._i] if self._i < len(self._toks) else (None, None)

    def _take(self, value=None):
        tok = self._peek()
        if value is not None and tok[1] != value:
            raise ValueError(f"expected {value!r}, got {tok!r}")
        self._i += 1
        return tok

    def _comparison(self):
        left = self._additive()
        kind, op = self._peek()
        if kind == "op" and op in ("=", "<>", "<", ">", "<=", ">="):
            self._take()
            right = self._additive()
            return lambda: _compare(op, left(), right())
        return left

    def _additive(self):
        node = self._term()
        while self._peek()[1] in ("+", "-") and self._peek()[0] == "op":
            op = self._take()[1]
            rhs, lhs = self._term(), node
            node = (lambda o, a, b: lambda: _arith(o, a(), b()))(op, lhs, rhs)
        return node

    def _term(self):
        node = self._unary()
        while self._peek()[1] in ("*", "/") and self._peek()[0] == "op":
            op = self._take()[1]
            rhs, lhs = self._unary(), node
            node = (lambda o, a, b: lambda: _arith(o, a(), b()))(op, lhs, rhs)
        return node

    def _unary(self):
        if self._peek() == ("op", "-"):
            self._take()
            inner = self._unary()
            return lambda: _arith("-", 0.0, inner())
        if self._peek() == ("op", "+"):
            self._take()
            return self._unary()
        return self._primary()

    def _primary(self):
        kind, text = self._take()
        sheet = self._sheet
        if kind == "num":
            return lambda: float(text)
        if kind == "str":
            return lambda: text[1:-1].replace('""', '"')
        if kind == "ref":
            if ":" in text:
                return self._range(sheet, text)
            m = _REF_RE.fullmatch(text)
            ref_sheet = m.group(1) or sheet
            return lambda: self.value(ref_sheet, m.group(2) + m.group(3))
        if kind == "op" and text == "(":
            node = self._comparison()
            self._take(")")
            return node
        if kind == "func":
            args = []
            if self._peek() != ("op", ")"):
                args.append(self._comparison())
                while self._peek() == ("op", ","):
                    self._take()
                    args.append(self._comparison())
            self._take(")")
            return self._call(text, args)
        raise ValueError(f"unexpected token {text!r}")

    def _range(self, sheet, text):
        from openpyxl.utils import column_index_from_string, get_column_letter

        a, b = text.split(":")
        ma, mb = _REF_RE.fullmatch(a), _REF_RE.fullmatch(b)
        ref_sheet = ma.group(1) or sheet
        c1, c2 = column_index_from_string(ma.group(2)), column_index_from_string(mb.group(2))
        r1, r2 = int(ma.group(3)), int(mb.group(3))
        coords = [f"{get_column_letter(c)}{r}"
                  for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
        return lambda: [self.value(ref_sheet, c) for c in coords]

    def _call(self, name, args):
        if name == "SUM":
            def _sum():
                total = 0.0
                for arg in args:
                    vals = arg()
                    for v in (vals if isinstance(vals, list) else [vals]):
                        if isinstance(v, _XlError):
                            return v
                        if isinstance(v, (int, float)) and not isinstance(v, bool):
                            total += v
                return total
            return _sum
        if name == "IF":
            def _if():
                cond = args[0]()
                if isinstance(cond, _XlError):
                    return cond
                if cond:
                    return args[1]()
                return args[2]() if len(args) > 2 else False
            return _if
        raise ValueError(f"unsupported function {name}")


def _label_rows(ws) -> dict:
    """{column-A label: row number} for a sheet (first occurrence wins)."""
    rows = {}
    for (cell,) in ws.iter_rows(min_col=1, max_col=1):
        if isinstance(cell.value, str):
            rows.setdefault(cell.value.strip(), cell.row)
    return rows


def _workbook(test, report):
    path = write_excel(report, os.path.join(_tmpdir(test), "SYNT_valuation.xlsx"))
    wb = load_workbook(path)
    return wb, XlEval(wb)


# --------------------------------------------------------------------------- #
#  Excel
# --------------------------------------------------------------------------- #
class ExcelReconciliationTests(unittest.TestCase):
    def assertClose(self, got, want, msg=None):
        self.assertIsInstance(got, float, msg)
        self.assertAlmostEqual(got, want, delta=max(1e-9 * abs(want), 1e-9), msg=msg)

    def test_dcf_bridge_subtracts_minority_interest_and_preferred(self):
        report = _report(_ClaimsProvider())
        dcf = report.dcf
        self.assertIsNotNone(dcf)
        wb, ev = _workbook(self, report)
        rows = _label_rows(wb["DCF"])
        self.assertEqual(wb["DCF"][f"B{rows['Less: minority interest']}"].value, MINORITY)
        self.assertEqual(wb["DCF"][f"B{rows['Less: preferred equity']}"].value, PREFERRED)

        def cell(label):
            return ev.value("DCF", f"B{rows[label]}")

        self.assertClose(cell("Enterprise value"), dcf.enterprise_value)
        self.assertClose(cell("Equity value"), dcf.equity_value)
        self.assertClose(cell("Implied price / share"), dcf.implied_price)
        self.assertClose(cell("Upside / (downside)"), dcf.upside)
        # The claims actually move the price, so this is not a vacuous check.
        self.assertAlmostEqual(
            dcf.enterprise_value - dcf.net_debt - dcf.equity_value, MINORITY + PREFERRED,
            delta=1.0,
        )

    def test_demo_formulas_reconcile_to_the_model(self):
        report = _report()
        wb, ev = _workbook(self, report)

        rows = _label_rows(wb["DCF"])
        self.assertClose(ev.value("DCF", f"B{rows['Equity value']}"), report.dcf.equity_value)
        self.assertClose(ev.value("DCF", f"B{rows['Implied price / share']}"),
                         report.dcf.implied_price)

        rows = _label_rows(wb["DDM_FCFE"])
        self.assertClose(ev.value("DDM_FCFE", f"B{rows['Equity value']}"),
                         report.fcfe.equity_value)
        self.assertClose(ev.value("DDM_FCFE", f"B{rows['Implied price / share']}"),
                         report.fcfe.implied_price)

        rows = _label_rows(wb["Summary"])
        cur = report.current_price
        blended = report.summary["blended_target"]
        self.assertClose(wb["Summary"][f"B{rows['Blended target (median)']}"].value, blended)
        self.assertClose(ev.value("Summary", f"C{rows['Blended target (median)']}"),
                         blended / cur - 1.0)
        self.assertClose(ev.value("Summary", f"C{rows['DCF (FCFF)']}"),
                         report.dcf.implied_price / cur - 1.0)

        errors = [(s, c, v) for s, c, v in ev.formulas() if isinstance(v, _XlError)]
        self.assertEqual(errors, [])

    def test_zero_revenue_and_zero_price_give_no_div0(self):
        report = _report()
        n = len(report.dcf.years)
        report.dcf = dataclasses.replace(
            report.dcf, revenue=[0.0] * n, ebit=[0.0] * n, nopat=[0.0] * n,
        )
        report.current_price = 0.0
        _wb, ev = _workbook(self, report)
        errors = [(s, c, v) for s, c, v in ev.formulas() if isinstance(v, _XlError)]
        self.assertEqual(errors, [])

    def test_blended_target_is_not_invented_when_engine_has_none(self):
        report = _report(run_dcf=False, run_comps=False, run_ddm=False, run_fcfe=False)
        self.assertIsNone(report.summary.get("blended_target"))
        wb, _ev = _workbook(self, report)
        rows = _label_rows(wb["Summary"])
        # No method ran, so there is nothing to ask peers for: plain n/a.
        r = rows["Blended target (median)"]
        self.assertEqual(wb["Summary"][f"B{r}"].value, "n/a")
        self.assertEqual(wb["Summary"][f"C{r}"].value, "n/a")
        self.assertEqual(wb["Summary"][f"B{rows['Verdict']}"].value, "N/A")

    def test_ddm_detail_formats_and_lists(self):
        report = _report()
        self.assertIsNotNone(report.ddm)
        wb, _ev = _workbook(self, report)
        ws = wb["DDM_FCFE"]
        rows = _label_rows(ws)
        self.assertIn("%", ws[f"B{rows['cost_of_equity']}"].number_format)
        if "high_growth_years" in rows:
            fmt = ws[f"B{rows['high_growth_years']}"].number_format
            self.assertNotIn("%", fmt)
            self.assertNotIn("$", fmt)
        for (cell,) in ws.iter_rows(min_col=2, max_col=2):
            if isinstance(cell.value, str):
                self.assertFalse(cell.value.startswith("["), f"{cell.coordinate}: {cell.value}")
        dividends = report.ddm.detail.get("dividends")
        if dividends:
            r = rows["dividends"]
            spilled = [ws.cell(row=r, column=2 + j).value for j in range(len(dividends))]
            self.assertEqual(spilled, [float(d) for d in dividends])

    def test_projection_headers_label_the_same_fiscal_years(self):
        report = _report()
        wb, _ev = _workbook(self, report)

        def header(ws, first_label):
            r = _label_rows(ws)[first_label]
            return [c.value for c in ws[r][1:] if c.value]

        dcf_hdr = header(wb["DCF"], "(values in reporting currency)")
        fcfe_hdr = header(wb["DDM_FCFE"], "(reporting currency)")
        self.assertEqual(len(dcf_hdr), len(report.dcf.years))
        self.assertEqual(dcf_hdr, fcfe_hdr[: len(dcf_hdr)])
        self.assertNotIn("FY 1", dcf_hdr)

    def test_control_characters_do_not_sink_the_workbook(self):
        report = _report()
        report.warnings.append("Comps: peer fetch failed: bad byte \x1b[0m in response")
        wb, _ev = _workbook(self, report)
        texts = [c.value for row in wb["Summary"].iter_rows() for c in row
                 if isinstance(c.value, str)]
        self.assertTrue(any("bad byte [0m in response" in t for t in texts))

    def test_one_failing_sheet_does_not_sink_the_workbook(self):
        from equity_valuation.report import excel

        report = _report()
        with patch.object(excel, "_write_comps", side_effect=RuntimeError("boom")):
            wb, _ev = _workbook(self, report)
        self.assertEqual(wb.sheetnames, ["Summary", "DCF", "Comps", "DDM_FCFE", "Sensitivity"])
        texts = [c.value for row in wb["Comps"].iter_rows() for c in row if c.value]
        self.assertTrue(any("boom" in str(t) for t in texts))


# --------------------------------------------------------------------------- #
#  HTML
# --------------------------------------------------------------------------- #
def _html(test, report, name="SYNT_valuation.html") -> str:
    path = write_html(report, os.path.join(_tmpdir(test), name))
    return Path(path).read_text(encoding="utf-8")


def _plotly_copies(text: str) -> int:
    return text.count("window.PlotlyConfig")


def _card(text: str, label: str) -> str:
    m = re.search(r"card-label'>" + re.escape(label) + r"</div><div class='card-value "
                  r"\w+'>([^<]*)<", text)
    return m.group(1) if m else None


class HtmlReportTests(unittest.TestCase):
    def test_bridge_shows_minority_interest_and_preferred(self):
        from equity_valuation.report.html import _fmt_big

        report = _report(_ClaimsProvider())
        text = _html(self, report)
        self.assertIn(f"Less: minority interest</th><td>{_fmt_big(MINORITY, '$')}<", text)
        self.assertIn(f"Less: preferred equity</th><td>{_fmt_big(PREFERRED, '$')}<", text)
        self.assertIn(f"Equity value</th><td>{_fmt_big(report.dcf.equity_value, '$')}<", text)

    def test_each_render_embeds_plotly_exactly_once_even_when_nested(self):
        # A second render starting while the first is in progress used to reset
        # the shared module state, leaving the first document with no plotly.js.
        from equity_valuation.report import html as html_mod

        report = _report()
        out = _tmpdir(self)
        original = html_mod._header_html
        calls = []

        def nested(*args, **kwargs):
            if not calls:
                calls.append(1)
                write_html(report, os.path.join(out, "inner.html"))
            return original(*args, **kwargs)

        with patch.object(html_mod, "_header_html", nested):
            write_html(report, os.path.join(out, "outer.html"))
        for name in ("outer.html", "inner.html"):
            text = Path(out, name).read_text(encoding="utf-8")
            self.assertEqual(_plotly_copies(text), 1, name)

    def test_concurrent_renders_embed_plotly_exactly_once(self):
        report = _report()
        out = _tmpdir(self)
        errors = []

        def work(i):
            try:
                for j in range(2):
                    write_html(report, os.path.join(out, f"r{i}_{j}.html"))
            except Exception as exc:  # noqa: BLE001
                errors.append(exc)

        threads = [threading.Thread(target=work, args=(i,)) for i in range(3)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        self.assertEqual(errors, [])
        for name in sorted(os.listdir(out)):
            text = Path(out, name).read_text(encoding="utf-8")
            self.assertEqual(_plotly_copies(text), 1, name)

    def test_no_blended_target_is_invented(self):
        report = _report(run_dcf=False, run_comps=False, run_ddm=False, run_fcfe=False)
        text = _html(self, report)
        self.assertEqual(_card(text, "Blended target"), "n/a")
        self.assertEqual(_card(text, "Upside / downside"), "n/a")

    def test_header_lists_every_method_behind_the_blended_target(self):
        from equity_valuation.report.html import _fmt_price

        report = _report()
        text = _html(self, report)
        self.assertEqual(_card(text, "Blended target"),
                         _fmt_price(report.summary["blended_target"], "$"))
        self.assertIn("Valuation by method", text)
        for name, price in report.summary["methods"].items():
            self.assertIn(f"{name}</th><td>{_fmt_price(price, '$')}<", text)

    def test_comps_chart_median_matches_the_trimmed_stats(self):
        from equity_valuation.report.html import _fmt_mult

        report = _report(_OutlierPeersProvider(), peers=["P0", "P1", "P2", "P3"])
        med = report.comps.stats["ev_ebitda"]["median"]
        text = _html(self, report)
        self.assertIn(f"Peer median {_fmt_mult(med)}", text)


# --------------------------------------------------------------------------- #
#  Sensitivity grids (Gordon and exit-multiple kinds)
# --------------------------------------------------------------------------- #
EXIT_MULTIPLES = [10.0, 11.0, 12.0, 13.0, 14.0]
WACC_LEVELS = [0.08, 0.087, 0.095, 0.102, 0.109]
NAN = float("nan")


def _exit_grid(centre, title="DCF implied price: WACC vs exit EV/EBITDA"):
    """A hand-built WACC x exit EV/EBITDA grid with two invalid cells.

    Cell [1][1] is None and [2][4] is NaN (invalid combinations); the centre
    cell is `centre`.
    """
    grid = [[20.0 + i + j for j in range(5)] for i in range(5)]
    grid[1][1] = None
    grid[2][4] = NAN
    grid[2][2] = centre
    return SensitivityResult(title=title, row_label="WACC", col_label="Exit EV/EBITDA",
                             row_values=list(WACC_LEVELS),
                             col_values=list(EXIT_MULTIPLES), grid=grid)


def _sens_block(ws, corner: str):
    """(header row, [body rows]) of the Sensitivity-sheet grid whose corner reads `corner`."""
    top = next(c.row for (c,) in ws.iter_rows(min_col=1, max_col=1) if c.value == corner)
    rows, r = [], top + 1
    while isinstance(ws.cell(row=r, column=1).value, (int, float, str)) \
            and ws.cell(row=r, column=1).fill.fill_type:
        rows.append(r)
        r += 1
    return top, rows


def _heatmaps(text: str) -> list:
    """[(trace, layout)] for every plotly heatmap embedded in an HTML report."""
    dec = json.JSONDecoder()
    out = []
    for m in re.finditer(r'Plotly\.newPlot\(\s*"[^"]+",\s*', text):
        data, end = dec.raw_decode(text, m.end())
        sep = re.match(r"\s*,\s*", text[end:])
        layout, _ = dec.raw_decode(text, end + sep.end())
        if data and data[0].get("type") == "heatmap":
            out.append((data[0], layout))
    return out


class SensitivityReportTests(unittest.TestCase):
    """Both grid kinds render by axis label, blank invalid cells, and mark the
    centre cell as the base case only when it is the headline DCF price."""

    def _exit_report(self, centre_is_headline=True):
        report = _report()
        headline = report.dcf.implied_price
        centre = headline if centre_is_headline else headline * 0.8
        report.sensitivities = [_exit_grid(centre)]
        return report

    def test_excel_exit_multiple_axis_shows_multiples_and_blank_invalid_cells(self):
        from equity_valuation.report.excel import MULTIPLE_FMT, PERCENT_FMT

        report = self._exit_report()
        wb, _ev = _workbook(self, report)
        ws = wb["Sensitivity"]
        top, rows = _sens_block(ws, "WACC \\ Exit EV/EBITDA")
        self.assertEqual(len(rows), 5)
        headers = [ws.cell(row=top, column=2 + j) for j in range(5)]
        self.assertEqual([c.value for c in headers], EXIT_MULTIPLES)
        self.assertEqual({c.number_format for c in headers}, {MULTIPLE_FMT})
        self.assertEqual([ws.cell(row=r, column=1).value for r in rows], WACC_LEVELS)
        self.assertEqual({ws.cell(row=r, column=1).number_format for r in rows}, {PERCENT_FMT})
        # None and NaN cells are blank, never 0 or the text "nan".
        self.assertIsNone(ws.cell(row=rows[1], column=3).value)
        self.assertIsNone(ws.cell(row=rows[2], column=6).value)

    def test_excel_marks_the_centre_only_when_it_is_the_headline(self):
        report = self._exit_report()
        wb, _ev = _workbook(self, report)
        ws = wb["Sensitivity"]
        _top, rows = _sens_block(ws, "WACC \\ Exit EV/EBITDA")
        filled = [(i, j) for i, r in enumerate(rows) for j in range(5)
                  if ws.cell(row=r, column=2 + j).fill.fill_type]
        self.assertEqual(filled, [(2, 2)])
        self.assertAlmostEqual(ws.cell(row=rows[2], column=4).value, report.dcf.implied_price)
        texts = [c.value for (c,) in ws.iter_rows(min_col=1, max_col=1) if c.value]
        self.assertIn("Highlighted cell = headline DCF (base case).", texts)

        report = self._exit_report(centre_is_headline=False)
        wb, _ev = _workbook(self, report)
        ws = wb["Sensitivity"]
        _top, rows = _sens_block(ws, "WACC \\ Exit EV/EBITDA")
        filled = [(i, j) for i, r in enumerate(rows) for j in range(5)
                  if ws.cell(row=r, column=2 + j).fill.fill_type]
        self.assertEqual(filled, [])
        texts = [c.value for (c,) in ws.iter_rows(min_col=1, max_col=1) if c.value]
        self.assertNotIn("Highlighted cell = headline DCF (base case).", texts)

    def test_excel_rate_axes_stay_percent_beyond_minus_100(self):
        # The axis format follows the label, not the magnitude: a -152% margin
        # used to render as '-1.52'.
        from equity_valuation.report.excel import PERCENT_FMT

        report = _report()
        report.sensitivities = [SensitivityResult(
            title="DCF implied price: EBIT margin vs terminal growth",
            row_label="EBIT margin", col_label="Terminal growth",
            row_values=[-1.54, -1.53, -1.52, -1.51, NAN],
            col_values=[0.015, 0.02, 0.025, 0.03, 0.035],
            grid=[[-5.0] * 5 for _ in range(5)],
        )]
        wb, _ev = _workbook(self, report)
        ws = wb["Sensitivity"]
        top, rows = _sens_block(ws, "EBIT margin \\ Terminal growth")
        self.assertEqual({ws.cell(row=r, column=1).number_format for r in rows[:4]},
                         {PERCENT_FMT})
        self.assertEqual(ws.cell(row=rows[4], column=1).value, "n/a")
        self.assertEqual({ws.cell(row=top, column=2 + j).number_format for j in range(5)},
                         {PERCENT_FMT})

    def test_gordon_grids_of_the_demo_mark_their_centre(self):
        report = _report()
        self.assertTrue(report.sensitivities)
        wb, _ev = _workbook(self, report)
        ws = wb["Sensitivity"]
        for sens in report.sensitivities:
            _top, rows = _sens_block(ws, f"{sens.row_label} \\ {sens.col_label}")
            filled = [(i, j) for i, r in enumerate(rows) for j in range(5)
                      if ws.cell(row=r, column=2 + j).fill.fill_type]
            self.assertEqual(filled, [(2, 2)], sens.title)
        for trace, layout in _heatmaps(_html(self, report)):
            self.assertTrue(all(t.endswith("%") for t in trace["x"] + trace["y"]))
            self.assertEqual(len(layout.get("shapes") or []), 1)

    def test_html_exit_multiple_axis_blank_cells_and_base_case(self):
        report = self._exit_report()
        report.sensitivities.append(_exit_grid(report.dcf.implied_price + 5.0,
                                               title="Centre is not the headline"))
        maps = _heatmaps(_html(self, report))
        self.assertEqual(len(maps), 2)
        (trace, layout), (other, other_layout) = maps
        self.assertEqual(trace["x"], ["10.0x", "11.0x", "12.0x", "13.0x", "14.0x"])
        self.assertEqual(trace["y"], ["8.00%", "8.70%", "9.50%", "10.20%", "10.90%"])
        self.assertIsNone(trace["z"][1][1])
        self.assertIsNone(trace["z"][2][4])  # NaN is written as null, not NaN
        self.assertEqual((trace["text"][1][1], trace["text"][2][4]), ("", ""))
        self.assertEqual(layout["xaxis"]["type"], "category")
        shapes = layout.get("shapes") or []
        self.assertEqual(len(shapes), 1)
        self.assertEqual((shapes[0]["x0"], shapes[0]["x1"], shapes[0]["y0"], shapes[0]["y1"]),
                         (1.5, 2.5, 1.5, 2.5))
        self.assertIn("headline DCF", layout["title"]["text"])
        # A centre that is not the headline price is not presented as the base case.
        self.assertEqual(other_layout.get("shapes") or [], [])
        self.assertNotIn("headline DCF", other_layout["title"]["text"])

    def test_blank_cells_are_explained_in_both_reports(self):
        from equity_valuation.report import excel, html

        self.assertEqual(excel.BLANK_SENSITIVITY_CELL_NOTE, html.BLANK_SENSITIVITY_CELL_NOTE)
        note = excel.BLANK_SENSITIVITY_CELL_NOTE
        report = self._exit_report()
        full = _exit_grid(report.dcf.implied_price, title="No blank cells")
        full.grid[1][1] = full.grid[2][4] = 21.0
        report.sensitivities.append(full)
        wb, _ev = _workbook(self, report)
        ws = wb["Sensitivity"]
        notes = [(c.row, c.value) for (c,) in ws.iter_rows(min_col=1, max_col=1)
                 if c.value == note]
        # Once, under the grid with blank cells (above the second grid).
        second = next(c.row for (c,) in ws.iter_rows(min_col=1, max_col=1)
                      if c.value == "No blank cells")
        self.assertEqual(len(notes), 1)
        self.assertLess(notes[0][0], second)
        text = _html(self, report)
        self.assertEqual(text.count(f"<p class='fig-note'>{note}</p>"), 1)
        # The Gordon grids of the demo have no blank cell, so no note.
        wb, _ev = _workbook(self, _report())
        self.assertNotIn(note, _label_rows(wb["Sensitivity"]))
        self.assertNotIn("fig-note'>", _html(self, _report()))

    def test_html_repeated_na_ticks_stay_distinct(self):
        # Plotly merges equal category labels, which collapsed a grid whose
        # margin axis is all n/a (pre-revenue company) into one row.
        report = _report()
        report.sensitivities = [SensitivityResult(
            title="margin", row_label="EBIT margin", col_label="Terminal growth",
            row_values=[NAN] * 5, col_values=[0.015, 0.02, 0.025, 0.03, 0.035],
            grid=[[NAN] * 5 for _ in range(5)],
        )]
        ((trace, _layout),) = _heatmaps(_html(self, report))
        self.assertEqual(len(set(trace["y"])), 5)
        self.assertTrue(all(t.startswith("n/a") for t in trace["y"]))

    def test_axis_label_rule_is_shared(self):
        from equity_valuation.report import excel, html

        for label, is_multiple in (("Exit EV/EBITDA", True), ("Exit multiple", True),
                                   ("WACC", False), ("Terminal growth", False),
                                   ("EBIT margin", False), ("EBITDA margin", False)):
            self.assertEqual(excel._is_multiple_axis(label), is_multiple, label)
            self.assertEqual(html._is_multiple_axis(label), is_multiple, label)
        self.assertEqual(html._axis_tick(12.0, "Exit EV/EBITDA"), "12.0x")
        self.assertEqual(html._axis_tick(-1.52, "EBIT margin"), "-152.00%")
        self.assertEqual(html._axis_tick(None, "WACC"), "n/a")

    def test_exit_multiple_run_shows_the_multiple_in_both_reports(self):
        from equity_valuation.report.excel import MULTIPLE_FMT
        from equity_valuation.schemas import DCFAssumptions

        report = _report(dcf_assumptions=DCFAssumptions(terminal_method="exit_multiple",
                                                        exit_ev_ebitda=12.0))
        self.assertEqual(report.dcf.assumptions["terminal_method"], "exit_multiple")
        wb, _ev = _workbook(self, report)
        ws = wb["DCF"]
        r = _label_rows(ws)["Exit EV/EBITDA"]
        self.assertEqual(ws[f"B{r}"].value, 12.0)
        self.assertEqual(ws[f"B{r}"].number_format, MULTIPLE_FMT)
        self.assertIn("Exit EV/EBITDA</th><td>12.0x<", _html(self, report))
        # Gordon runs don't list a multiple they don't use.
        wb, _ev = _workbook(self, _report())
        self.assertNotIn("Exit EV/EBITDA", _label_rows(wb["DCF"]))

    def test_exit_multiple_run_says_terminal_growth_is_not_in_the_tv(self):
        from equity_valuation.report.excel import EXIT_MULTIPLE_GROWTH_NOTE
        from equity_valuation.report.html import EXIT_MULTIPLE_GROWTH_NOTE as HTML_NOTE
        from equity_valuation.schemas import DCFAssumptions

        self.assertEqual(EXIT_MULTIPLE_GROWTH_NOTE, HTML_NOTE)
        report = _report(dcf_assumptions=DCFAssumptions(terminal_method="exit_multiple",
                                                        exit_ev_ebitda=12.0))
        wb, _ev = _workbook(self, report)
        ws = wb["DCF"]
        r = _label_rows(ws)["Terminal growth"]
        self.assertEqual(ws[f"C{r}"].value, EXIT_MULTIPLE_GROWTH_NOTE)
        text = _html(self, report)
        self.assertRegex(text, r"Terminal growth</th><td>[0-9.]+%<br><span class='aside'>"
                         + re.escape(EXIT_MULTIPLE_GROWTH_NOTE.replace("'", "&#x27;")))
        # A Gordon run's terminal growth does set its terminal value: no note.
        report = _report()
        wb, _ev = _workbook(self, report)
        ws = wb["DCF"]
        self.assertIsNone(ws[f"C{_label_rows(ws)['Terminal growth']}"].value)
        self.assertNotIn("class='aside'", _html(self, report))


# --------------------------------------------------------------------------- #
#  Methods left out of the blended target (financial institutions)
# --------------------------------------------------------------------------- #
class BlendExclusionTests(unittest.TestCase):
    """A bank's DCF and FCFE are listed for reference but marked as not in the
    blend, so the headline need not equal the median of the listed rows."""

    def setUp(self):
        self.report = _report(_BankProvider())
        self.excluded = self.report.summary["excluded_from_blend"]
        self.assertEqual(set(self.excluded), {"DCF", "FCFE"})

    def test_excel_marks_the_methods_left_out_of_the_blend(self):
        wb, _ev = _workbook(self, self.report)
        ws = wb["Summary"]
        rows = _label_rows(ws)
        table = [ws[f"A{r}"].value for r in range(rows["Method"] + 1,
                                                  rows["Blended target (median)"])]
        self.assertEqual(table, ["DCF (FCFF) (not in blend)", "Trading comps (median)",
                                 "DDM", "FCFE (not in blend)"])
        for label, key in (("DCF (FCFF)", "DCF"), ("FCFE", "FCFE")):
            r = rows[f"{label} (not in blend)"]
            self.assertEqual(ws[f"D{r}"].value, f"Not in blend: {self.excluded[key]}")
        for label in ("Trading comps (median)", "DDM"):
            self.assertIsNone(ws[f"D{rows[label]}"].value)
        # The operating company's table is unmarked.
        wb, _ev = _workbook(self, _report())
        self.assertFalse([k for k in _label_rows(wb["Summary"]) if "not in blend" in k])

    def test_html_marks_the_methods_left_out_of_the_blend(self):
        from equity_valuation.report.html import _fmt_price

        text = _html(self, self.report)
        methods = self.report.summary["methods"]
        for name, why in self.excluded.items():
            self.assertIn(f"{name} <span class='aside'>(not in blend: {why})</span></th>"
                          f"<td>{_fmt_price(methods[name], '$')}<", text)
        for name in ("Comps (median)", "DDM"):
            self.assertIn(f"{name}</th><td>{_fmt_price(methods[name], '$')}<", text)
        self.assertNotIn("not in blend", _html(self, _report()))

    def test_memo_and_deck_mark_the_methods_left_out_of_the_blend(self):
        from docx import Document
        from pptx import Presentation

        from backend import exports, valuation_service

        out = _tmpdir(self)
        with patch.object(valuation_service, "_get_provider",
                          lambda ticker, refresh=False: _BankProvider()), \
                patch.object(exports, "_ensure_out", lambda: Path(out)):
            memo = exports.export_memo(DEMO_TICKER, {"peers": ",".join(DEMO_PEERS)})
            deck = exports.export_deck(DEMO_TICKER, {"peers": ",".join(DEMO_PEERS)})
        doc = Document(memo)
        labels = [row.cells[0].text for row in doc.tables[0].rows]
        self.assertIn("DCF (not in blend)", labels)
        self.assertIn("FCFE (not in blend)", labels)
        self.assertIn("DDM", labels)
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Not in the blended target: DCF (not meaningful for a financial "
                      "institution)", text)
        cells = [cell.text for slide in Presentation(deck).slides for shape in slide.shapes
                 if shape.has_table for row in shape.table.rows for cell in row.cells]
        self.assertIn("DCF (not in blend)", cells)
        self.assertIn("FCFE (not in blend)", cells)


# --------------------------------------------------------------------------- #
#  No blended target, a withheld verdict, and reference-only methods (the
#  engine's own summaries for a lessor, a captive-finance group and a bank)
# --------------------------------------------------------------------------- #
class _KindProvider(SyntheticProvider):
    """The synthetic company flagged by its filings as a debt-funded lessor or a
    captive-finance group (``financials._financial_kind``), or presented as a
    bank by its industry. ``peer_rows`` replaces the peers' multiples."""

    def __init__(self, kind=None, industry=None, peer_rows=None):
        super().__init__()
        self.kind, self.industry, self.peer_rows = kind, industry, peer_rows

    def get_company_data(self, ticker):
        company = make_company()
        if self.kind:
            company.financials._financial_kind = self.kind
        if self.industry:
            company.market.sector = "Financial Services"
            company.market.industry = self.industry
        return company

    def suggest_peers(self, ticker):
        return []  # peers are never found automatically

    def get_peer_comp_rows(self, tickers):
        if self.peer_rows is not None:
            return list(self.peer_rows)
        return super().get_peer_comp_rows(tickers)


def _kind_report(kind=None, industry=None, peers=None, peer_rows=None, run_comps=None):
    """The engine's report for the synthetic company of that kind. Without
    peers, comps do not run unless asked (and then find no peer)."""
    return value_company(DEMO_TICKER, provider=_KindProvider(kind, industry, peer_rows),
                         peers=peers,
                         run_comps=(peers is not None) if run_comps is None else run_comps)


def _lessor_report():
    """A lessor without peers (as AER): the DCF, the FCFE and the low-payout
    DDM are reference only, so there is no target and no verdict."""
    return _kind_report("lessor")


def _captive_report():
    """A captive-finance group without peers (as GM or CAT): the same."""
    return _kind_report("captive_finance")


def _bank_report():
    """A bank without peers (as BAC): the blend is the DDM alone, so the target
    is kept but the verdict and upside are withheld."""
    return _kind_report(industry="Banks - Diversified")


def _ev_only_peers_report():
    """A lessor whose supplied peers have EV multiples only: comps ran on them
    but give a lessor no price (it uses P/E and P/B), so there is no target,
    and supplying peers is not what is missing."""
    rows = [CompRow(ticker=f"P{i}", name=f"P{i}", market_cap=5e10, enterprise_value=5.2e10,
                    ev_ebitda=18.0 + i, ev_sales=4.0) for i in range(3)]
    return _kind_report("lessor", peers=["P0", "P1", "P2"], peer_rows=rows)


def _placeholder_report():
    """Hand-built (the synthetic company always values): an ordinary company
    whose every method gave the 0.00 placeholder, so there is no target and
    no peers are asked for."""
    report = _report(run_comps=False)
    none = "no valuation (0.00)"
    report.summary = dict(
        report.summary, methods={"DCF": 0.0, "DDM": 0.0, "FCFE": 0.0},
        blended_target=None, blended_upside=None, recommendation="N/A",
        excluded_from_blend={"DCF": none, "DDM": none, "FCFE": none})
    return report


_GREEN, _RED = "000B6E0B", "00C00000"  # the workbook's upside colours


def _font_rgb(cell):
    return getattr(getattr(cell.font, "color", None), "rgb", None)


def _uncoloured(cell) -> bool:
    return _font_rgb(cell) not in (_GREEN, _RED)


def _sheet_upside(ws):
    """The DCF or FCFE sheet's 'Upside / (downside)' value cell."""
    return ws[f"B{_label_rows(ws)['Upside / (downside)']}"]


def _card_class(text: str, label: str) -> str:
    m = re.search(r"card-label'>" + re.escape(label) + r"</div><div class='card-value "
                  r"(\w+)'>", text)
    return m.group(1) if m else None


def _method_rows(text: str) -> dict:
    """{method name: (aside or None, upside cell class)} from the HTML header's
    'Valuation by method' table."""
    table = text[text.index("<table class='kv methods'>"):]
    table = table[:table.index("</table>")]
    return {m.group(1): (m.group(2), m.group(3)) for m in re.finditer(
        r"<th class='rowhead'>([^<]+?)(?: <span class='aside'>\(not in blend: ([^<]*)\)"
        r"</span>)?</th><td>[^<]*</td><td class='(\w+)'>", table)}


def _bridge_upside_class(text: str) -> str:
    return re.search(r"<th class='rowhead'>Upside</th><td class='(\w+)'>", text).group(1)


# Excel Summary label -> the summary's method name.
_EXCEL_METHODS = {"DCF (FCFF)": "DCF", "Trading comps (median)": "Comps (median)",
                  "DDM": "DDM", "FCFE": "FCFE"}


def _excel_method_rows(ws) -> dict:
    """{method name: (label, upside cell, column-D note)} on the Summary sheet."""
    rows = _label_rows(ws)
    out = {}
    for label, name in _EXCEL_METHODS.items():
        for text in (label, f"{label} (not in blend)"):
            if text in rows:
                r = rows[text]
                out[name] = (text, ws[f"C{r}"], ws[f"D{r}"].value)
    return out


class EngineKindShapeTests(unittest.TestCase):
    """The shapes the tests below render come from the engine itself."""

    def test_lessor_captive_and_bank_summaries(self):
        for build, kind in ((_lessor_report, "lessor"), (_captive_report, "captive_finance")):
            report = build()
            s = report.summary
            self.assertEqual(s["financial_kind"], kind)
            self.assertIsNone(report.comps)
            self.assertIsNone(s["blended_target"])
            self.assertIsNone(s["blended_upside"])
            self.assertEqual(s["recommendation"], "N/A")
            self.assertEqual(set(s["excluded_from_blend"]), {"DCF", "DDM", "FCFE"})
            self.assertTrue(s["excluded_from_blend"]["DDM"].startswith("dividends only"))
        self.assertIn("lessor", _lessor_report().summary["excluded_from_blend"]["DCF"])
        self.assertIn("captive finance",
                      _captive_report().summary["excluded_from_blend"]["DCF"])

        s = _bank_report().summary
        self.assertEqual(s["financial_kind"], "bank")
        self.assertEqual(set(s["excluded_from_blend"]), {"DCF", "FCFE"})
        self.assertAlmostEqual(s["blended_target"], s["methods"]["DDM"], places=12)
        self.assertIsNone(s["blended_upside"])
        self.assertEqual(s["recommendation"], "N/A")

        report = _ev_only_peers_report()
        self.assertTrue(report.comps.peers)
        self.assertNotIn("Comps (median)", report.summary["methods"])
        self.assertIsNone(report.summary["blended_target"])


class NoBlendedTargetTests(unittest.TestCase):
    def test_excel_lessor_shows_no_target_and_uncoloured_reference_rows(self):
        from equity_valuation.report.excel import NO_TARGET_SUPPLY_PEERS

        for build in (_lessor_report, _captive_report):
            report = build()
            excluded = report.summary["excluded_from_blend"]
            wb, ev = _workbook(self, report)
            ws = wb["Summary"]
            rows = _label_rows(ws)
            r = rows["Blended target (median)"]
            self.assertEqual(ws[f"B{r}"].value, NO_TARGET_SUPPLY_PEERS)
            self.assertEqual(ws[f"C{r}"].value, "n/a")
            self.assertTrue(_uncoloured(ws[f"C{r}"]))
            self.assertEqual(ws[f"B{rows['Verdict']}"].value, "N/A")
            # Every method is listed for reference, marked with the engine's
            # reason, its upside a live formula but uncoloured.
            methods = _excel_method_rows(ws)
            for name in ("DCF", "DDM", "FCFE"):
                label, up, note = methods[name]
                self.assertTrue(label.endswith(" (not in blend)"), label)
                self.assertEqual(note, f"Not in blend: {excluded[name]}")
                self.assertTrue(str(up.value).startswith("="), name)
                self.assertTrue(_uncoloured(up), name)
            self.assertEqual(methods["DDM"][0], "DDM (not in blend)")
            # The DCF and FCFE sheets' upsides are uncoloured too.
            for sheet in ("DCF", "DDM_FCFE"):
                self.assertTrue(_uncoloured(_sheet_upside(wb[sheet])), sheet)
            # Column B holds the bold, right-aligned text without clipping.
            self.assertGreaterEqual(ws.column_dimensions["B"].width,
                                    len(NO_TARGET_SUPPLY_PEERS) + 2)
            errors = [(sh, c, v) for sh, c, v in ev.formulas() if isinstance(v, _XlError)]
            self.assertEqual(errors, [])

    def test_excel_bank_keeps_its_target_but_colours_no_upside(self):
        report = _bank_report()
        wb, _ev = _workbook(self, report)
        ws = wb["Summary"]
        rows = _label_rows(ws)
        r = rows["Blended target (median)"]
        self.assertAlmostEqual(ws[f"B{r}"].value, report.summary["blended_target"], places=9)
        self.assertEqual(ws[f"C{r}"].value, "n/a")  # no live upside formula
        self.assertTrue(_uncoloured(ws[f"C{r}"]))
        self.assertEqual(ws[f"B{rows['Verdict']}"].value, "N/A")
        methods = _excel_method_rows(ws)
        # The DDM is the whole blend, yet its upside is the figure the engine
        # withholds: shown, not coloured. The reference-only rows likewise.
        self.assertEqual(methods["DDM"][0], "DDM")
        self.assertIsNone(methods["DDM"][2])
        for name in ("DCF", "DDM", "FCFE"):
            self.assertTrue(str(methods[name][1].value).startswith("="), name)
            self.assertTrue(_uncoloured(methods[name][1]), name)

        # An ordinary run colours every method's upside, the blended upside and
        # the DCF and FCFE sheets' upsides by sign.
        wb, _ev = _workbook(self, _report())
        ws = wb["Summary"]
        rows = _label_rows(ws)
        for name, (_label, up, _note) in _excel_method_rows(ws).items():
            self.assertFalse(_uncoloured(up), name)
        self.assertEqual(_font_rgb(ws[f"C{rows['Blended target (median)']}"]), _RED)
        self.assertEqual(ws[f"B{rows['Verdict']}"].value, "Overvalued")
        for sheet in ("DCF", "DDM_FCFE"):
            self.assertEqual(_font_rgb(_sheet_upside(wb[sheet])), _RED, sheet)

    def test_html_shows_no_target_and_uncoloured_reference_rows(self):
        from equity_valuation.report.excel import NO_TARGET_SUPPLY_PEERS
        from equity_valuation.report.html import _fmt_price

        report = _lessor_report()
        excluded = report.summary["excluded_from_blend"]
        text = _html(self, report)
        self.assertEqual(_card(text, "Blended target"), NO_TARGET_SUPPLY_PEERS)
        self.assertEqual(_card(text, "Upside / downside"), "n/a")
        self.assertEqual(_card_class(text, "Upside / downside"), "neutral")
        self.assertEqual(_card(text, "Verdict"), "N/A")
        self.assertEqual(_card_class(text, "Verdict"), "neutral")
        rows = _method_rows(text)
        self.assertEqual(set(rows), {"DCF", "DDM", "FCFE"})
        for name, (aside, cls) in rows.items():
            self.assertEqual(aside, excluded[name], name)
            self.assertEqual(cls, "neutral", name)
        self.assertEqual(_bridge_upside_class(text), "neutral")

        report = _bank_report()
        text = _html(self, report)
        self.assertEqual(_card(text, "Blended target"),
                         _fmt_price(report.summary["blended_target"], "$"))
        self.assertEqual(_card(text, "Upside / downside"), "n/a")
        self.assertEqual(_card_class(text, "Upside / downside"), "neutral")
        self.assertEqual(_card(text, "Verdict"), "N/A")
        rows = _method_rows(text)
        self.assertEqual(rows["DDM"], (None, "neutral"))  # in the blend, no verdict
        self.assertEqual(rows["DCF"][1], "neutral")
        self.assertEqual(_bridge_upside_class(text), "neutral")

        # An ordinary run: coloured method upsides, upside card and verdict.
        text = _html(self, _report())
        rows = _method_rows(text)
        self.assertEqual({cls for _aside, cls in rows.values()}, {"neg", "pos"})
        self.assertEqual(rows["DDM"], (None, "neg"))
        self.assertEqual(_bridge_upside_class(text), "neg")
        self.assertEqual(_card_class(text, "Upside / downside"), "neg")
        self.assertEqual(_card(text, "Verdict"), "Overvalued")

    def test_every_output_asks_for_peers_only_when_they_would_help(self):
        from backend.exports import _blended_target
        from backend.serialization import build_ai_context, report_to_dict
        from equity_valuation.report.excel import NO_TARGET_SUPPLY_PEERS, needs_peers

        supply = "none (no target; supply peers to add trading comps)"
        cases = (
            ("lessor, no peers", _lessor_report(), True),
            ("captive, no peers", _captive_report(), True),
            # Comps ran but found no usable peer: the engine asks for peers too.
            ("lessor, no usable peer", _kind_report("lessor", run_comps=True), True),
            ("lessor, peers without P/E or P/B", _ev_only_peers_report(), False),
            ("ordinary, no valuation", _placeholder_report(), False),
        )
        for label, report, want in cases:
            with self.subTest(label):
                self.assertIsNone(report.summary["blended_target"])
                self.assertEqual(needs_peers(report.summary, report.comps), want)
                text = NO_TARGET_SUPPLY_PEERS if want else "n/a"
                wb, _ev = _workbook(self, report)
                ws = wb["Summary"]
                self.assertEqual(ws[f"B{_label_rows(ws)['Blended target (median)']}"].value,
                                 text)
                self.assertEqual(_card(_html(self, report), "Blended target"), text)
                self.assertEqual(_blended_target(report.summary, "USD", report.comps), text)
                ctx = build_ai_context(report_to_dict(report))
                self.assertIn("blended target: " + (supply if want else "n/a"), ctx)

    def _office(self, build):
        """(memo paragraphs, memo table rows, deck title runs, deck table rows)."""
        from docx import Document
        from pptx import Presentation

        from backend import exports, valuation_service

        real = exports.run_valuation_report

        def engine_built(ticker, payload):
            _report_, echo = real(ticker, payload)
            return build(), echo

        out = _tmpdir(self)
        with patch.object(valuation_service, "_get_provider",
                          lambda ticker, refresh=False: SyntheticProvider()), \
                patch.object(exports, "run_valuation_report", engine_built), \
                patch.object(exports, "_ensure_out", lambda: Path(out)):
            memo = exports.export_memo(DEMO_TICKER, {})
            deck = exports.export_deck(DEMO_TICKER, {})
        doc = Document(memo)
        paragraphs = [p.text for p in doc.paragraphs]
        memo_rows = [[c.text for c in row.cells] for row in doc.tables[0].rows]
        slides = list(Presentation(deck).slides)
        title_runs = [r for shape in slides[0].shapes if shape.has_text_frame
                      for p in shape.text_frame.paragraphs for r in p.runs]
        deck_rows = [[c.text for c in row.cells] for shape in slides[1].shapes
                     if shape.has_table for row in shape.table.rows]
        return paragraphs, memo_rows, title_runs, deck_rows

    def test_memo_and_deck_show_no_target_uncoloured(self):
        from pptx.dml.color import RGBColor

        excluded = _lessor_report().summary["excluded_from_blend"]
        paragraphs, memo_rows, title_runs, deck_rows = self._office(_lessor_report)
        text = "\n".join(paragraphs)
        self.assertIn("blended fair value: No target (supply peers) · model verdict: N/A", text)
        self.assertIn(["Blended target", "No target (supply peers)", "n/a"], memo_rows)
        for rows in (memo_rows, deck_rows):
            self.assertEqual([row[0] for row in rows[1:-1]],
                             ["DCF (not in blend)", "DDM (not in blend)", "FCFE (not in blend)"])
        self.assertIn("Not in the blended target: "
                      + "; ".join(f"{k} ({excluded[k]})" for k in ("DCF", "DDM", "FCFE")) + ".",
                      text)
        headline = [r for r in title_runs if "blended fair value" in r.text]
        self.assertEqual(len(headline), 1)
        self.assertIn("blended fair value: No target (supply peers) · N/A", headline[0].text)
        # Dim grey, not the green / rose / amber of a verdict.
        self.assertEqual(headline[0].font.color.rgb, RGBColor(0x5C, 0x6B, 0x84))
        self.assertIn(["Blended target", "No target (supply peers)", "n/a"], deck_rows)

    def test_memo_and_deck_keep_a_target_whose_upside_is_withheld(self):
        from backend.exports import _money

        paragraphs, memo_rows, title_runs, deck_rows = self._office(_bank_report)
        target = _money(_bank_report().summary["blended_target"])
        self.assertIn(f"blended fair value {target} (upside n/a) · model verdict: N/A",
                      "\n".join(paragraphs))
        self.assertIn(["Blended target", target, "n/a"], memo_rows)
        self.assertIn(["Blended target", target, "n/a"], deck_rows)
        self.assertIn("DDM", [row[0] for row in memo_rows])
        headline = [r for r in title_runs if "blended fair value" in r.text][0]
        self.assertIn(f"blended fair value {target} (upside n/a) · N/A", headline.text)

    def test_reference_only_upsides_are_uncoloured_beside_a_verdict(self):
        # With peers, the bank's and the lessor's blends rest on comps (and the
        # bank's DDM) and give a verdict; their reference-only rows stay
        # uncoloured while the methods in the blend keep their colour.
        for report in (_report(_BankProvider()), _kind_report("lessor", peers=DEMO_PEERS)):
            s = report.summary
            excluded = s["excluded_from_blend"]
            with self.subTest(s["financial_kind"]):
                self.assertNotEqual(s["recommendation"], "N/A")
                self.assertIn("Comps (median)", s["methods"])
                self.assertNotIn("Comps (median)", excluded)
                wb, _ev = _workbook(self, report)
                for name, (_label, up, _note) in _excel_method_rows(wb["Summary"]).items():
                    self.assertEqual(_uncoloured(up), name in excluded, name)
                self.assertTrue(_uncoloured(_sheet_upside(wb["DCF"])))
                self.assertTrue(_uncoloured(_sheet_upside(wb["DDM_FCFE"])))
                text = _html(self, report)
                for name, (_aside, cls) in _method_rows(text).items():
                    self.assertEqual(cls == "neutral", name in excluded, name)
                self.assertEqual(_bridge_upside_class(text), "neutral")

    def test_cli_summary_prints_no_upside_for_no_target(self):
        from equity_valuation import engine
        from equity_valuation.report.excel import NO_TARGET_SUPPLY_PEERS

        report = _lessor_report()
        out = _tmpdir(self)
        with patch.object(engine, "value_company", lambda *a, **k: report):
            code, stdout, err = _run_cli(["LSR", "--excel", "--out", out])
        self.assertEqual(code, 0, err)
        self.assertIn("Recommendation: N/A", stdout)
        line = next(ln for ln in stdout.splitlines() if "Blended target" in ln)
        # The upside column reads n/a; the target column 'n/a' or, once the
        # CLI shares the reports' wording, 'No target (supply peers)'.
        self.assertEqual(line.split()[-1], "n/a")
        target = line.split("Blended target", 1)[1].rsplit("n/a", 1)[0].strip()
        self.assertIn(target, ("n/a", NO_TARGET_SUPPLY_PEERS))
        for name in ("DCF", "DDM", "FCFE"):
            self.assertIn(f"{name} (excluded)", stdout)


class SharedTextTests(unittest.TestCase):
    """The dashboard's lib/format.ts keeps a copy of the rules the Python
    outputs share; pin its constants to the Python ones."""

    FORMAT_TS = Path(__file__).resolve().parent.parent / "frontend" / "lib" / "format.ts"

    def test_dashboard_copies_the_no_target_text_and_the_blume_weights(self):
        from equity_valuation.data import market
        from equity_valuation.report.excel import BLUME_NOTE, NO_TARGET_SUPPLY_PEERS

        self.assertEqual(BLUME_NOTE, f"Blume-adjusted toward 1 ({market._BLUME_RAW_WEIGHT:g} "
                                     f"x raw + {market._BLUME_MARKET_WEIGHT:g})")
        if not self.FORMAT_TS.is_file():
            self.skipTest("no frontend in this tree")
        ts = self.FORMAT_TS.read_text(encoding="utf-8")
        self.assertIn("export const NO_TARGET_SUPPLY_PEERS = "
                      f"{json.dumps(NO_TARGET_SUPPLY_PEERS)};", ts)
        consts = dict(re.findall(r"export const (BLUME_\w+_WEIGHT) = ([\d.]+);", ts))
        self.assertEqual(float(consts["BLUME_RAW_WEIGHT"]), market._BLUME_RAW_WEIGHT)
        self.assertEqual(float(consts["BLUME_MARKET_WEIGHT"]), market._BLUME_MARKET_WEIGHT)


# --------------------------------------------------------------------------- #
#  Adjusted beta
# --------------------------------------------------------------------------- #
class _AdjustedBetaProvider(SyntheticProvider):
    """The synthetic company with a Blume-adjusted beta and Yahoo's raw beta."""

    def get_company_data(self, ticker):
        company = make_company()
        company.market.raw_beta = 2.217
        company.market.beta = round(0.67 * 2.217 + 0.33, 6)
        return company


class AdjustedBetaTests(unittest.TestCase):
    def test_excel_and_html_label_the_adjusted_beta_with_its_raw_value(self):
        report = _report(_AdjustedBetaProvider())
        beta = report.dcf.wacc.beta
        self.assertAlmostEqual(beta, 1.81539, places=6)
        wb, _ev = _workbook(self, report)
        ws = wb["DCF"]
        rows = _label_rows(ws)
        self.assertNotIn("Beta", rows)
        r = rows["Beta (adj.)"]
        self.assertAlmostEqual(ws[f"B{r}"].value, beta, places=9)
        self.assertEqual(ws[f"C{r}"].value,
                         "Raw beta 2.217, Blume-adjusted toward 1 (0.67 x raw + 0.33)")

        text = _html(self, report)
        self.assertIn("<th class='rowhead'>Beta (adj.)</th><td>1.82<br><span class='aside'>"
                      "Raw beta 2.217, Blume-adjusted toward 1 (0.67 x raw + 0.33)</span>",
                      text)

    def test_an_unadjusted_beta_keeps_its_plain_label(self):
        report = _report()  # the synthetic beta is not adjusted (raw_beta None)
        self.assertIsNone(report.company.market.raw_beta)
        wb, _ev = _workbook(self, report)
        rows = _label_rows(wb["DCF"])
        self.assertIn("Beta", rows)
        self.assertIsNone(wb["DCF"][f"C{rows['Beta']}"].value)
        self.assertFalse([k for k in rows if "(adj.)" in k])
        text = _html(self, report)
        self.assertIn("<th class='rowhead'>Beta</th><td>1.10</td>", text)
        self.assertNotIn("Beta (adj.)", text)


# --------------------------------------------------------------------------- #
#  CLI
# --------------------------------------------------------------------------- #
def _run_cli(argv):
    """(exit code, stdout, stderr) of cli.main, with SystemExit mapped to its code."""
    from equity_valuation.cli import main

    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        try:
            code = main(argv)
        except SystemExit as exc:
            code = exc.code
    return code, out.getvalue(), err.getvalue()


class CliTests(unittest.TestCase):
    def test_terminal_growth_reaches_ddm_and_fcfe(self):
        from equity_valuation import engine

        seen = {}
        real = engine.value_company

        def spy(*args, **kwargs):
            seen.update(kwargs)
            return real(*args, **kwargs)

        with patch.object(engine, "value_company", spy):
            code, _out, err = _run_cli(["--demo", "--terminal-growth", "0.04", "--quiet",
                                        "--excel", "--out", _tmpdir(self)])
        self.assertEqual(code, 0, err)
        self.assertEqual(seen["dcf_assumptions"].terminal_growth, 0.04)
        self.assertEqual(seen["ddm_assumptions"].terminal_growth, 0.04)

    def test_bad_numeric_flags_are_rejected(self):
        cases = [
            (["--rf", "4.3"], "0.043"),
            (["--tax", "21"], "0.21"),
            (["--target-ebit-margin", "25"], "decimals"),
            (["--rf", "nan"], "finite"),
            (["--erp", "inf"], "finite"),
            (["--terminal-growth", "nan"], "finite"),
            (["--cost-of-debt", "-0.01"], ">= 0"),
            (["--forecast-years", "0"], "between 1 and"),
            (["--forecast-years", "-3"], "between 1 and"),
            (["--exit-ev-ebitda", "-5", "--terminal-method", "exit_multiple"], "> 0"),
            (["--terminal-method", "exit_multiple"], "requires --exit-ev-ebitda"),
        ]
        out = _tmpdir(self)
        for extra, hint in cases:
            with self.subTest(args=extra):
                code, _out, err = _run_cli(["--demo", "--quiet", "--out", out] + extra)
                self.assertEqual(code, 2)
                self.assertIn(hint, err)
                self.assertNotIn("Traceback", err)
        self.assertEqual(os.listdir(out), [])

    def test_valid_decimal_flags_still_run(self):
        code, _out, err = _run_cli([
            "--demo", "--quiet", "--excel", "--out", _tmpdir(self), "--rf", "-0.005",
            "--tax", "0.21", "--terminal-method", "exit_multiple", "--exit-ev-ebitda", "12",
        ])
        self.assertEqual(code, 0, err)

    def test_export_failure_exits_nonzero(self):
        out = _tmpdir(self)
        os.mkdir(os.path.join(out, "SYNT_valuation.xlsx"))
        os.mkdir(os.path.join(out, "SYNT_valuation.html"))
        code, _out, err = _run_cli(["--demo", "--quiet", "--out", out])
        self.assertEqual(code, 1)
        self.assertIn("Excel export failed", err)
        self.assertIn("HTML export failed", err)

    def test_out_pointing_at_a_file_is_a_clean_error(self):
        path = os.path.join(_tmpdir(self), "not_a_dir")
        Path(path).write_text("x", encoding="utf-8")
        code, _out, err = _run_cli(["--demo", "--quiet", "--out", path])
        self.assertEqual(code, 2)
        self.assertIn("not a directory", err)
        self.assertNotIn("Traceback", err)

    def test_ticker_is_sanitized_in_output_filenames(self):
        from equity_valuation import engine

        report = _report()
        report.company = dataclasses.replace(report.company, ticker="BRK/B")
        out = _tmpdir(self)
        with patch.object(engine, "value_company", lambda *a, **k: report):
            code, _stdout, err = _run_cli(["BRK/B", "--quiet", "--excel", "--out", out])
        self.assertEqual(code, 0, err)
        self.assertEqual(os.listdir(out), ["BRK_B_valuation.xlsx"])

    def test_demo_rejects_real_peer_tickers(self):
        code, _out, err = _run_cli(["--demo", "--peers", "MSFT,GOOGL", "--quiet",
                                    "--out", _tmpdir(self)])
        self.assertEqual(code, 2)
        self.assertIn("synthetic peers", err)

        code, _out, err = _run_cli(["--demo", "--peers", "peer1,PEER3", "--quiet", "--excel",
                                    "--out", _tmpdir(self)])
        self.assertEqual(code, 0, err)



class ExternalDataExportTests(unittest.TestCase):
    def test_external_text_cannot_become_an_excel_formula(self):
        report = value_company("SYNT", provider=SyntheticProvider(), run_sensitivity=False)
        hostile = '=HYPERLINK("https://invalid.example/","Name")'
        report.comps.peers[0].name = hostile
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "report.xlsx"
            write_excel(report, str(path))
            wb = load_workbook(path)
            peers = wb["Comps"]
            name_cells = [cell for row in peers for cell in row if cell.value == hostile]
            self.assertEqual(len(name_cells), 1)
            self.assertEqual(name_cells[0].data_type, "s")
            formulas = [cell for row in wb["DCF"] for cell in row if cell.data_type == "f"]
            self.assertTrue(any("SUM(" in cell.value for cell in formulas))
            self.assertTrue(any("/" in cell.value for cell in formulas))
            wb.close()

    def test_peer_money_displays_its_own_currency(self):
        report = value_company("SYNT", provider=SyntheticProvider(), run_sensitivity=False)
        report.comps.peers[0] = CompRow("GBPPEER", "Peer", market_cap=1e9,
                                      enterprise_value=2e9, pe=20, currency="GBP")
        markup = _comps_table(report, "$")
        row = markup.split("GBPPEER", 1)[1].split("</tr>", 1)[0]
        self.assertIn("GBP", row)
        self.assertIn("£", row)
        self.assertNotIn("$", row)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "report.xlsx"
            write_excel(report, str(path))
            wb = load_workbook(path)
            rows = list(wb["Comps"].values)
            peer_row = next(r for r in rows if r[0] == "GBPPEER")
            self.assertIn("GBP", peer_row)
            wb.close()

class UnavailableValuationReportTests(unittest.TestCase):
    @staticmethod
    def _report(*, missing_shares=False, missing_margin=False, zero_value=False):
        company = make_company()
        if missing_shares:
            company.market.shares_outstanding = 0.0
            company.financials.diluted_shares = [0.0] * 5
        if missing_margin:
            company.financials.net_income = [None] * 5
        if zero_value:
            company.balance_sheet.total_debt = 0.0
            company.balance_sheet.cash_and_investments = 0.0
            for name in ("ebit", "net_income", "pretax_income", "interest_expense", "change_in_nwc"):
                setattr(company.financials, name, [0.0] * 5)
            company.financials.capex = company.financials.dep_amort[:]
        provider = SyntheticProvider()
        provider.get_company_data = lambda ticker: company
        return value_company("SYNT", provider=provider, run_ddm=False,
                             run_comps=False, run_sensitivity=False)

    def test_missing_shares_show_no_price_or_downside_in_exports_and_cli(self):
        report = self._report(missing_shares=True)
        wb, _ = _workbook(self, report)
        summary_rows = _label_rows(wb["Summary"])
        for name in ("DCF (FCFF) (not in blend)", "FCFE (not in blend)"):
            row = summary_rows[name]
            self.assertIsNone(wb["Summary"].cell(row, 2).value)
            self.assertEqual(wb["Summary"].cell(row, 3).value, "n/a")
        for sheet in ("DCF", "DDM_FCFE"):
            rows = _label_rows(wb[sheet])
            self.assertIsNone(wb[sheet].cell(rows["Implied price / share"], 2).value)
            self.assertIsNone(wb[sheet].cell(rows["Upside / (downside)"], 2).value)
        from equity_valuation.report.html import _dcf_table, _methods_table
        for markup in (_dcf_table(report, "$"), _methods_table(report, "$")):
            self.assertNotIn("-100.0%", markup)
            self.assertNotIn("$0.00", markup)
        with contextlib.redirect_stdout(io.StringIO()) as output:
            from equity_valuation.cli import _print_summary
            _print_summary(report)
        self.assertNotIn("-100.0%", output.getvalue())

    def test_missing_fcfe_margin_is_unavailable_even_with_valid_shares(self):
        report = self._report(missing_margin=True)
        wb, _ = _workbook(self, report)
        rows = _label_rows(wb["DDM_FCFE"])
        self.assertIsNone(wb["DDM_FCFE"].cell(rows["Implied price / share"], 2).value)
        self.assertIsNone(wb["DDM_FCFE"].cell(rows["Upside / (downside)"], 2).value)

    def test_completed_zero_values_keep_the_real_total_loss(self):
        report = self._report(zero_value=True)
        wb, evaluator = _workbook(self, report)
        for sheet in ("DCF", "DDM_FCFE"):
            rows = _label_rows(wb[sheet])
            self.assertEqual(evaluator.value(sheet, f"B{rows['Implied price / share']}"), 0.0)
            self.assertEqual(evaluator.value(sheet, f"B{rows['Upside / (downside)']}"), -1.0)
        from equity_valuation.report.html import _methods_table
        self.assertIn("-100.0%", _methods_table(report, "$"))

    def test_cli_returns_failure_with_diagnostic_exports_when_every_model_unavailable(self):
        from equity_valuation import engine
        out = _tmpdir(self)
        with patch.object(engine, "value_company", return_value=self._report(missing_shares=True)):
            code, _stdout, error = _run_cli(["SYNT", "--quiet", "--excel", "--out", out])
        self.assertEqual(code, 1)
        self.assertIn("no requested model produced an available valuation", error)
        self.assertTrue((Path(out) / "SYNT_valuation.xlsx").exists())

    def test_calculation_time_is_preserved_separately_from_export_time(self):
        report = self._report()
        self.assertTrue(report.valued_at.endswith("+00:00"))
        report.valued_at = "2026-01-02T03:04:05+00:00"
        wb, _ = _workbook(self, report)
        rows = _label_rows(wb["Summary"])
        self.assertEqual(wb["Summary"].cell(rows["Valuation calculated (UTC)"], 2).value,
                         report.valued_at)
        self.assertNotEqual(wb["Summary"].cell(rows["Report generated (UTC)"], 2).value,
                            report.valued_at)
        from equity_valuation.report.html import _header_html
        self.assertIn("Valuation calculated (UTC): " + report.valued_at, _header_html(report, "$"))


# --------------------------------------------------------------------------- #
#  Workbook layout and cached formula results
# --------------------------------------------------------------------------- #
def _excel_text(value, fmt):
    """The text Excel shows for a number in one of the workbook's masks, or None
    for the General format (which Excel shortens to fit)."""
    v = float(value)
    m = re.fullmatch(r'(?:"([^"]*)")?#,##0(\.0+)?', fmt)
    if m:  # money, or a count with thousands separators
        dec = len(m.group(2) or ".") - 1
        return ("-" if round(v, dec) < 0 else "") + (m.group(1) or "") + f"{abs(v):,.{dec}f}"
    m = re.fullmatch(r"0(\.0+)?(%|\"x\")?", fmt)
    if m:
        dec = len(m.group(1) or ".") - 1
        suffix = {"%": "%", '"x"': "x"}.get(m.group(2) or "", "")
        return f"{v * (100 if suffix == '%' else 1):.{dec}f}{suffix}"
    return None


def _col_widths(ws) -> dict:
    """{column index: width} for the columns given a width."""
    widths = {}
    for dim in ws.column_dimensions.values():
        if dim.width:
            for c in range(dim.min, (dim.max or dim.min) + 1):
                widths[c] = dim.width
    return widths


def _scaled_report(factor=40.0):
    """The demo with every DCF and FCFE amount scaled up (a $14T enterprise
    value) and net cash, i.e. a negative net debt."""
    report = _report()
    d = report.dcf
    up = lambda xs: [x * factor for x in xs]  # noqa: E731
    report.dcf = dataclasses.replace(
        d, revenue=up(d.revenue), ebit=up(d.ebit), nopat=up(d.nopat), fcff=up(d.fcff),
        pv_fcff=up(d.pv_fcff), terminal_value=d.terminal_value * factor,
        pv_terminal=d.pv_terminal * factor, net_debt=-abs(d.enterprise_value) * 2)
    f = report.fcfe
    report.fcfe = dataclasses.replace(
        f, fcfe=up(f.fcfe), pv_fcfe=up(f.pv_fcfe), terminal_value=f.terminal_value * factor,
        pv_terminal=f.pv_terminal * factor)
    return report


class ExcelLayoutAndCacheTests(unittest.TestCase):
    def _write(self, report):
        path = write_excel(report, os.path.join(_tmpdir(self), "SYNT_valuation.xlsx"))
        return load_workbook(path), load_workbook(path, data_only=True)

    def test_every_number_fits_its_column(self):
        # Excel shows '####' for a formatted number wider than its column (a
        # $345B enterprise value with cents did not fit the old width of 16).
        for label, report in (("demo", _report()), ("scaled", _scaled_report())):
            wb, values = self._write(report)
            checked = 0
            for ws in wb.worksheets:
                widths = _col_widths(ws)
                for row in ws.iter_rows():
                    for cell in row:
                        v = values[ws.title][cell.coordinate].value
                        if isinstance(v, bool) or not isinstance(v, (int, float)):
                            continue
                        text = _excel_text(v, cell.number_format)
                        if text is None:
                            continue
                        checked += 1
                        width = widths.get(cell.column, 8.43)
                        self.assertLessEqual(len(text) + 1, width,
                                             f"{label} {ws.title}!{cell.coordinate} {text}")
            self.assertGreater(checked, 100, label)
        # The demo's large amounts really are wider than the old width.
        wb, values = self._write(_report())
        rows = _label_rows(wb["DCF"])
        ev = values["DCF"][f"B{rows['Enterprise value']}"].value
        self.assertGreater(len(_excel_text(ev, wb["DCF"][f"B{rows['Enterprise value']}"]
                                           .number_format)) + 1, 16)

    def test_formulas_are_saved_with_their_results(self):
        # A viewer that does not calculate (file previews, pandas) reads the
        # cached values; they must be what Excel computes, and the formulas stay.
        report = _report()
        wb, values = self._write(report)
        evaluated = list(XlEval(wb).formulas())
        self.assertGreater(len(evaluated), 20)
        for sheet, coord, want in evaluated:
            got = values[sheet][coord].value
            self.assertIsInstance(want, float, f"{sheet}!{coord}")
            self.assertIsInstance(got, float, f"{sheet}!{coord}")
            self.assertAlmostEqual(got, want, delta=1e-9 * max(1.0, abs(want)),
                                   msg=f"{sheet}!{coord}")
            self.assertEqual(wb[sheet][coord].data_type, "f")
        dcf_rows = _label_rows(values["DCF"])
        self.assertAlmostEqual(values["DCF"][f"B{dcf_rows['Implied price / share']}"].value,
                               report.dcf.implied_price, delta=1e-9)
        self.assertAlmostEqual(values["DCF"][f"B{dcf_rows['Enterprise value']}"].value,
                               report.dcf.enterprise_value, delta=1e-3)
        fcfe_rows = _label_rows(values["DDM_FCFE"])
        self.assertAlmostEqual(values["DDM_FCFE"][f"B{fcfe_rows['Implied price / share']}"]
                               .value, report.fcfe.implied_price, delta=1e-9)
        rows = _label_rows(values["Summary"])
        self.assertAlmostEqual(values["Summary"][f"C{rows['Blended target (median)']}"].value,
                               report.summary["blended_upside"], delta=1e-9)
        self.assertAlmostEqual(values["Summary"][f"C{rows['DCF (FCFF)']}"].value,
                               report.dcf.upside, delta=1e-9)

    def test_an_empty_string_result_is_cached_as_text(self):
        report = _report()
        n = len(report.dcf.years)
        report.dcf = dataclasses.replace(report.dcf, revenue=[0.0] * n, ebit=[0.0] * n,
                                         nopat=[0.0] * n)
        wb, values = self._write(report)
        r = _label_rows(wb["DCF"])["growth %"]
        cell = values["DCF"][f"C{r}"]
        self.assertEqual(cell.data_type, "str")  # IF(...=0,"",...) gave ""
        self.assertIn(cell.value, (None, ""))
        self.assertEqual(wb["DCF"][f"C{r}"].data_type, "f")

    def test_unmapped_currency_masks_carry_the_iso_code(self):
        from equity_valuation.report.excel import _money_fmt

        report = _report()
        for code, mask in (("USD", '"$"#,##0.00'), ("CAD", '"C$"#,##0.00'),
                           ("SEK", '"SEK "#,##0.00'), (None, "#,##0.00")):
            report.company.market.currency = code
            self.assertEqual(_money_fmt(report), mask, code)
        # The mask's literal is counted when the columns are fitted.
        report.company.market.currency = "SEK"
        wb, values = self._write(report)
        widths = _col_widths(wb["DCF"])
        r = _label_rows(wb["DCF"])["Enterprise value"]
        text = _excel_text(values["DCF"][f"B{r}"].value, wb["DCF"][f"B{r}"].number_format)
        self.assertTrue(text.startswith("SEK "))
        self.assertLessEqual(len(text) + 1, widths[2])

    def test_a_malformed_currency_code_leaves_a_plain_number_mask(self):
        from equity_valuation.report.excel import _money_fmt

        report = _report()
        # A quote in provider text would end the mask's literal early, and
        # openpyxl then reads the mask back as a date format.
        for code in ('U"S', "<B>X</B>", "US", "USDX", "usd", "\x1b"):
            report.company.market.currency = code
            self.assertEqual(_money_fmt(report), "#,##0.00", repr(code))
        report.company.market.currency = 'U"S'
        wb, _values = self._write(report)
        cell = wb["DCF"][f"B{_label_rows(wb['DCF'])['Enterprise value']}"]
        self.assertEqual(cell.number_format, "#,##0.00")
        self.assertFalse(cell.is_date)

    def test_the_workbook_is_written_like_any_other_file(self):
        out = _tmpdir(self)
        plain = os.path.join(out, "plain.txt")
        with open(plain, "w", encoding="utf-8") as fh:
            fh.write("x")
        path = write_excel(_report(), os.path.join(out, "SYNT_valuation.xlsx"))
        # The usual permissions (a temporary file renamed into place was 0600),
        # and no temporary file left beside it.
        self.assertEqual(os.stat(path).st_mode & 0o777, os.stat(plain).st_mode & 0o777)
        self.assertEqual(sorted(os.listdir(out)), ["SYNT_valuation.xlsx", "plain.txt"])
        # An existing file keeps its mode when it is written again.
        os.chmod(path, 0o640)
        write_excel(_report(), path)
        self.assertEqual(os.stat(path).st_mode & 0o777, 0o640)

    def test_concurrent_writes_to_one_path_do_not_fail(self):
        # The saved workbook is never read back from disk, so two runs writing
        # into one output folder cannot trip over each other's half-written file.
        report = _report()
        path = os.path.join(_tmpdir(self), "SYNT_valuation.xlsx")
        errors = []

        def run():
            for _ in range(3):
                try:
                    write_excel(report, path)
                except Exception as exc:  # noqa: BLE001 - collected for the assertion
                    errors.append(exc)

        threads = [threading.Thread(target=run) for _ in range(4)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        self.assertEqual(errors, [])


# --------------------------------------------------------------------------- #
#  HTML formatting, the comps chart, and the DDM / FCFE section
# --------------------------------------------------------------------------- #
class _CaptureFigs:
    """An embedder that keeps each figure instead of serializing it."""

    def __init__(self):
        self.figs = []

    def to_html(self, fig):
        self.figs.append(fig)
        return "<div class='fig'></div>"


def _without_scripts(text: str) -> str:
    return re.sub(r"<script.*?</script>", "", text, flags=re.S)


class HtmlFormattingTests(unittest.TestCase):
    def test_negative_prices_put_the_minus_before_the_symbol(self):
        from equity_valuation.report.html import _fmt_price

        self.assertEqual(_fmt_price(-6.52, "$"), "-$6.52")
        self.assertEqual(_fmt_price(-1234.5, "€"), "-€1,234.50")
        self.assertEqual(_fmt_price(-0.004, "$"), "$0.00")  # rounds to zero: no sign
        self.assertEqual(_fmt_price(40.84, "$"), "$40.84")

        report = _report()
        report.dcf.implied_price = -6.52
        report.summary["methods"]["DCF"] = -6.52
        dcf_row = next(r for r in report.football_field if r.method == "DCF")
        dcf_row.low = -10.0
        text = _html(self, report)
        self.assertIn("DCF</th><td>-$6.52<", text)
        self.assertIn("Implied price</th><td>-$6.52<", text)
        self.assertNotIn("$-", _without_scripts(text))
        # The tooltips too (plotly figure JSON): preformatted, sign first.
        self.assertIn('"-$10.00"', text)
        self.assertNotIn('"$-', text)

    def test_currency_markers_and_axis_titles(self):
        report = _report()
        for code, price in (("CAD", "C$40.84"), ("SEK", "SEK 40.84")):
            report.company.market.currency = code
            text = _html(self, report)
            self.assertEqual(_card(text, "Current price"), price, code)
            self.assertIn(f'"text":"Implied price per share ({code})"', text)
            self.assertIn(f'"text":"FCFF ({code})"', text)
        report.company.market.currency = None
        text = _html(self, report)
        self.assertEqual(_card(text, "Current price"), "40.84")
        self.assertIn('"text":"FCFF"', text)
        self.assertNotIn('"text":"FCFF ()"', text)
        self.assertNotIn("per share ()", text)

    def test_every_mapped_currency_keeps_its_symbol(self):
        from equity_valuation import config
        from equity_valuation.report.html import _currency_symbol

        for code, symbol in config.CURRENCY_SYMBOLS.items():
            self.assertEqual(_currency_symbol(code), symbol, code)
        self.assertEqual(_currency_symbol("SEK"), "SEK ")
        self.assertEqual(_currency_symbol(None), "")

    def test_a_malformed_currency_code_never_reaches_the_markup(self):
        from equity_valuation.report.html import _currency_symbol

        for code in ('U"S', "<B>X</B>", "US", "USDX", "usd"):
            self.assertEqual(_currency_symbol(code), "", repr(code))
        report = _report()
        report.company.market.currency = "<B>X</B>"
        report.comps.peers[0].currency = "<SVG ONLOAD=1>"
        text = _html(self, report)
        body = _without_scripts(text)
        for tag in ("<B>", "</B>", "<SVG"):
            self.assertNotIn(tag, body)
        # Shown as text where the report names the currency, never as a prefix.
        self.assertIn("&lt;B&gt;X&lt;/B&gt;", body)
        self.assertIn("&lt;SVG ONLOAD=1&gt;", body)
        self.assertEqual(_card(text, "Current price"), "40.84")
        # Nor in the charts: the axis titles carry no code.
        self.assertNotIn("<B>X</B>", text)
        self.assertIn('"text":"FCFF"', text)


class CompsChartTests(unittest.TestCase):
    @staticmethod
    def _chart(report):
        from equity_valuation.report.html import _comps_fig

        cap = _CaptureFigs()
        html = _comps_fig(report, cap)
        return (cap.figs[0] if cap.figs else None), html

    def test_an_operating_company_charts_ev_ebitda(self):
        report = _report()
        fig, _ = self._chart(report)
        self.assertIn("EV/EBITDA", fig.layout.title.text)
        bars = dict(zip(fig.data[0].x, fig.data[0].y))
        self.assertEqual(bars[DEMO_TICKER], report.comps.target.ev_ebitda)

    def test_a_bank_charts_the_multiple_that_gave_its_price(self):
        report = _report(_BankProvider())
        implied = report.comps.implied
        self.assertIsNone(implied["ev_ebitda"])
        self.assertIsNotNone(implied["pe"])
        fig, _ = self._chart(report)
        self.assertEqual(fig.layout.title.text, "Comps — P/E: peers vs target")
        self.assertEqual(fig.layout.yaxis.title.text, "P / E (x)")
        bars = dict(zip(fig.data[0].x, fig.data[0].y))
        for peer in report.comps.peers:
            self.assertEqual(bars[peer.ticker], peer.pe)
        self.assertIn("P/E", fig.data[0].hovertemplate)

    def test_no_chart_without_a_priced_multiple_or_a_peer_that_has_it(self):
        # Comps ran but found no peer: nothing to set the target against.
        report = _kind_report("lessor", run_comps=True)
        self.assertEqual(report.comps.peers, [])
        self.assertEqual(self._chart(report), (None, ""))
        # Peers with EV multiples only: a lessor's comps give no price at all.
        self.assertEqual(self._chart(_ev_only_peers_report()), (None, ""))
        # A priced multiple that no peer has: the target is never charted alone.
        report = _report(_BankProvider())
        report.comps.peers = [dataclasses.replace(p, pe=None) for p in report.comps.peers]
        report.comps.implied = {"pe": 50.0}
        self.assertEqual(self._chart(report), (None, ""))


class HtmlDdmFcfeSectionTests(unittest.TestCase):
    def test_a_bank_shows_the_ddm_inputs_behind_its_target(self):
        from equity_valuation.report.html import _fmt_price

        report = _bank_report()
        detail = report.ddm.detail
        self.assertEqual(report.summary["blended_target"], report.ddm.implied_price)
        text = _without_scripts(_html(self, report))
        self.assertIn("<h2>Dividend discount &amp; FCFE</h2>", text)
        section = text[text.index("<h2>Dividend discount &amp; FCFE</h2>"):]
        section = section[:section.index("</section>")]
        self.assertIn(f"Latest dividend per share (D0)</th><td>{_fmt_price(detail['D0'], '$')}<",
                      section)
        self.assertIn(f"High-growth rate</th><td>{detail['high_growth'] * 100:.2f}%<", section)
        self.assertIn(f"High-growth years</th><td>{detail['high_growth_years']}<", section)
        self.assertIn(f"PV of terminal value</th><td>{_fmt_price(detail['pv_terminal'], '$')}<",
                      section)
        self.assertIn(f"Implied price</th><td>{_fmt_price(report.ddm.implied_price, '$')}<",
                      section)
        self.assertIn(f"Dividend per share</th><td>{_fmt_price(detail['dividends'][0], '$')}<",
                      section)
        # The FCFE is reference only for a bank, and the DCF section says so too.
        why = report.summary["excluded_from_blend"]["FCFE"]
        fcfe = section[section.index("Free cash flow to equity (FCFE)"):]
        self.assertIn(f"Not in the blended target: {why}.", fcfe)
        self.assertNotIn("Not in the blended target", section[:section.index("FCFE")])
        dcf = text[text.index("<h2>Discounted cash flow (DCF)</h2>"):]
        self.assertIn("Not in the blended target: "
                      f"{report.summary['excluded_from_blend']['DCF']}.",
                      dcf[:dcf.index("</section>")])
        # The DDM variant by its name, not its identifier.
        self.assertEqual(report.ddm.method, "two_stage")
        self.assertIn("Method</th><td>Two-stage<", section)
        self.assertNotIn("two_stage", text)
        # No verdict for a DDM-only blend: the upsides are uncoloured.
        self.assertNotIn("<td class='pos'>", section)
        self.assertNotIn("<td class='neg'>", section)

    def test_a_failed_method_is_not_called_reference_only(self):
        # The engine also leaves out a method that gave no price; the line
        # gives its reason without suggesting it was set aside on purpose.
        report = _report()
        report.summary.setdefault("excluded_from_blend", {})["FCFE"] = "non-finite implied price"
        text = _without_scripts(_html(self, report))
        self.assertIn("Not in the blended target: non-finite implied price.", text)
        self.assertNotIn("(reference only)", text)

    def test_ddm_variants_are_named(self):
        from equity_valuation.report.excel import ddm_method_label

        self.assertEqual([ddm_method_label(m) for m in ("two_stage", "h_model", "gordon")],
                         ["Two-stage", "H-model", "Gordon growth"])
        self.assertEqual(ddm_method_label("custom"), "custom")
        self.assertEqual(ddm_method_label(None), "n/a")
        report = _report()
        report.ddm = dataclasses.replace(report.ddm, method="h_model")
        self.assertIn("Method</th><td>H-model<", _html(self, report))
        wb = load_workbook(write_excel(report, os.path.join(_tmpdir(self), "d.xlsx")))
        ws = wb["DDM_FCFE"]
        self.assertEqual(ws[f"B{_label_rows(ws)['Method']}"].value, "H-model")

    def test_section_follows_the_models_that_ran(self):
        report = _report()
        text = _without_scripts(_html(self, report))
        section = text[text.index("<h2>Dividend discount &amp; FCFE</h2>"):]
        section = section[:section.index("</section>")]
        self.assertNotIn("Not in the blended target", text)
        self.assertIn(f"Equity value</th><td>${report.fcfe.equity_value / 1e9:,.2f}B<", section)
        self.assertIn("<td class='neg'>", section)  # in the blend, with a verdict
        text = _html(self, _report(run_ddm=False, run_fcfe=False))
        self.assertNotIn("Dividend discount &amp; FCFE", text)


# --------------------------------------------------------------------------- #
#  Reverse DCF, terminal reinvestment check, comps basis and peer counts
# --------------------------------------------------------------------------- #
SHEETS = ["Summary", "DCF", "Comps", "DDM_FCFE", "Sensitivity"]
# The trailing-basis fixture: Yahoo's figures to this quarter end, read on
# TTM_TODAY (the synthetic fiscal year 2024 ended on 2024-12-31).
TTM_AS_OF, TTM_TODAY = "2025-03-31", datetime.date(2025, 6, 1)


def _dcf_section(text: str) -> str:
    text = _without_scripts(text)
    section = text[text.index("<h2>Discounted cash flow (DCF)</h2>"):]
    return section[:section.index("</section>")]


def _ttm_report(ebitda=True):
    """The synthetic company valued on the opt-in trailing comps basis, its
    trailing figures 10% above fiscal 2024's (no trailing EBITDA when
    ``ebitda`` is False, so EV/EBITDA falls back to the fiscal year)."""
    from equity_valuation.models import comps as comps_mod

    company = make_company()
    fin = company.financials
    company.market = dataclasses.replace(
        company.market, ttm_as_of=TTM_AS_OF, fiscal_year_end="2024-12-31",
        ttm_currency="USD", ttm_revenue=fin.revenue[-1] * 1.1,
        ttm_ebitda=fin.ebitda[-1] * 1.1 if ebitda else None,
        ttm_net_income=fin.net_income[-1] * 1.1)

    class _Ttm(SyntheticProvider):
        def get_company_data(self, ticker):
            return company

    with patch.object(comps_mod, "_today", return_value=TTM_TODAY):
        return _report(_Ttm(), comps_basis="ttm", run_sensitivity=False)


class ReverseDcfReportTests(unittest.TestCase):
    """The workbook and the HTML report show the engine's reverse DCF (what
    the price implies for year-1 growth, the target EBIT margin and the
    discount rate, one at a time), the terminal year's reinvestment check, the
    comps' target-metric basis and the peer count per multiple."""

    @classmethod
    def setUpClass(cls):
        cls.report = _report()

    def assertClose(self, got, want, msg=None):
        self.assertIsInstance(got, float, msg)
        self.assertAlmostEqual(got, want, delta=1e-9, msg=msg)

    def test_dcf_sheet_shows_the_reverse_dcf_as_values(self):
        from equity_valuation.report.excel import PERCENT_FMT, REVERSE_DCF_TITLE

        reverse = self.report.reverse_dcf
        wb, ev = _workbook(self, self.report)
        self.assertEqual(wb.sheetnames, SHEETS)  # no new sheet
        ws = wb["DCF"]
        rows = _label_rows(ws)
        self.assertIn(REVERSE_DCF_TITLE, rows)
        self.assertGreater(rows[REVERSE_DCF_TITLE], rows["Upside / (downside)"])
        self.assertEqual([ws.cell(row=rows["Driver"], column=c).value for c in (1, 2, 3)],
                         ["Driver", "Model", "Market-implied"])
        expected = {
            "Year-1 revenue growth": (reverse["current_assumption_y1"],
                                      reverse["implied_growth_y1"]),
            "Target EBIT margin": (reverse["margin"]["current_target_ebit_margin"],
                                   reverse["margin"]["implied_target_ebit_margin"]),
            "Discount rate (WACC)": (reverse["wacc"]["current_wacc"],
                                     reverse["wacc"]["implied_wacc"]),
        }
        for label, (model, implied) in expected.items():
            r = rows[label]
            self.assertClose(ws[f"B{r}"].value, model, label)
            self.assertClose(ws[f"C{r}"].value, implied, label)
            for col in "BC":  # solved by the engine: values, not formulas
                self.assertEqual(ws[f"{col}{r}"].data_type, "n", label)
                self.assertEqual(ws[f"{col}{r}"].number_format, PERCENT_FMT, label)
        # The demo's figures, as the CLI prints them.
        self.assertEqual([f"{ws[f'C{rows[k]}'].value:.1%}" for k in expected],
                         ["17.5%", "30.8%", "8.2%"])
        notes = [c.value for row in ws.iter_rows(min_row=rows[REVERSE_DCF_TITLE])
                 for c in row if isinstance(c.value, str)]
        self.assertIn("Solved by the model, one driver at a time; values, not live formulas.",
                      notes)
        # The bridge's formulas still reconcile to the model.
        self.assertClose(ev.value("DCF", f"B{rows['Implied price / share']}"),
                         self.report.dcf.implied_price)

        summary = wb["Summary"]
        srows = _label_rows(summary)
        r = srows["Market-implied year-1 growth"]
        self.assertClose(summary[f"B{r}"].value, reverse["implied_growth_y1"])
        self.assertEqual(summary[f"B{r}"].number_format, PERCENT_FMT)
        self.assertEqual(summary[f"C{r}"].value, "Model 8.0%; reverse DCF on the DCF sheet")
        self.assertEqual(r, srows["Verdict"] + 1)

    def test_terminal_check_rows_equal_the_model(self):
        from equity_valuation.report.excel import PERCENT_FMT, TERMINAL_DIAGNOSTICS

        assumptions = self.report.dcf.assumptions
        wb, _ev = _workbook(self, self.report)
        ws = wb["DCF"]
        rows = _label_rows(ws)
        labels = [label for label, _key in TERMINAL_DIAGNOSTICS]
        self.assertEqual(labels, ["Terminal net reinvestment (% of NOPAT)",
                                  "Implied return on new invested capital", "Latest ROIC"])
        # Under the PV of the terminal value, indented, before the enterprise value.
        self.assertEqual([rows[k] for k in labels],
                         [rows["PV of terminal value"] + i for i in (1, 2, 3)])
        self.assertEqual(rows["Enterprise value"], rows["PV of terminal value"] + 4)
        for label, key in TERMINAL_DIAGNOSTICS:
            cell = ws[f"B{rows[label]}"]
            self.assertClose(cell.value, assumptions[key], label)
            self.assertEqual(cell.number_format, PERCENT_FMT, label)
            self.assertTrue(ws[f"A{rows[label]}"].value.startswith("  "), label)
        self.assertAlmostEqual(assumptions["implied_terminal_ronic"], 0.3714450, places=6)
        # An exit multiple has no stable year to check: n/a, never a number.
        report = _report(dcf_assumptions=DCFAssumptions(terminal_method="exit_multiple",
                                                        exit_ev_ebitda=12.0),
                         run_sensitivity=False)
        wb, ev = _workbook(self, report)
        ws = wb["DCF"]
        rows = _label_rows(ws)
        self.assertEqual([ws[f"B{rows[k]}"].value for k in labels], ["n/a"] * 3)
        self.assertClose(ev.value("DCF", f"B{rows['Implied price / share']}"),
                         report.dcf.implied_price)

        text = _dcf_section(_html(self, self.report))
        for label, key in TERMINAL_DIAGNOSTICS:
            self.assertIn(f"<th class='rowhead sub'>{label}</th>"
                          f"<td>{assumptions[key] * 100:.1f}%</td>", text)
        self.assertIn("Implied return on new invested capital</th><td>37.1%<", text)

    def test_html_shows_the_reverse_dcf_and_the_peer_counts(self):
        from equity_valuation.report.excel import REVERSE_DCF_TITLE

        text = _html(self, self.report)
        section = _dcf_section(text)
        self.assertIn(f"<h3 class='subhead'>{REVERSE_DCF_TITLE}</h3>", section)
        for label, model, implied in (("Year-1 revenue growth", "8.0%", "17.5%"),
                                      ("Target EBIT margin", "25.0%", "30.8%"),
                                      ("Discount rate (WACC)", "9.5%", "8.2%")):
            self.assertIn(f"{label}</th><td>{model}</td><td>{implied}</td>", section)
        self.assertIn("Target metrics: fiscal year 2024. Peers per multiple: EV/EBITDA n=3 · "
                      "EV/Sales n=3 · P/E n=3 · P/B n=3 · PEG n=3</caption>", text)

    def test_comps_sheet_names_the_basis_and_the_peer_count(self):
        report = _report(_OutlierPeersProvider(), peers=["P0", "P1", "P2", "P3"],
                         run_sensitivity=False)
        wb, _ev = _workbook(self, report)
        ws = wb["Comps"]
        rows = _label_rows(ws)
        self.assertIn("Target metrics: fiscal year 2024", rows)
        self.assertEqual(rows["Target metrics: fiscal year 2024"], rows["Multiples table"] + 1)
        header = rows["Multiple"]  # the stats header comes first
        self.assertEqual(ws.cell(row=header, column=8).value, "N")
        counts = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=8).value
                  for r in range(header + 1, header + 6)}
        for key, label in (("ev_ebitda", "EV/EBITDA"), ("ev_sales", "EV/Sales"),
                           ("pe", "P/E"), ("pb", "P/B"), ("peg", "PEG")):
            self.assertEqual(counts[label], report.comps.stats[key]["n"], label)
        # The outlying EV/EBITDA is trimmed: three values behind it, four elsewhere.
        self.assertEqual((counts["EV/EBITDA"], counts["P/E"]), (3, 4))
        self.assertIn("EV/EBITDA n=3 · EV/Sales n=4", _html(self, report))

    def test_trailing_basis_is_named_in_both_reports(self):
        from equity_valuation.report.excel import comps_basis_text

        report = _ttm_report()
        self.assertEqual(comps_basis_text(report), f"trailing twelve months to {TTM_AS_OF}")
        wb, _ev = _workbook(self, report)
        self.assertIn(f"Target metrics: trailing twelve months to {TTM_AS_OF}",
                      _label_rows(wb["Comps"]))
        self.assertIn(f"<caption>Target metrics: trailing twelve months to {TTM_AS_OF}. ",
                      _html(self, report))
        # A metric on the fiscal year (no trailing EBITDA) is pointed out.
        self.assertEqual(comps_basis_text(_ttm_report(ebitda=False)),
                         f"trailing twelve months to {TTM_AS_OF} (some metrics on the fiscal "
                         "year; see the notes)")
        # Asked for but unusable (no trailing figures): the fiscal year.
        self.assertEqual(comps_basis_text(_report(comps_basis="ttm", run_sensitivity=False)),
                         "fiscal year 2024")

    def test_a_value_driver_run_names_its_ronic(self):
        report = _report(dcf_assumptions=DCFAssumptions(terminal_method="value_driver",
                                                        terminal_ronic=0.15),
                         run_sensitivity=False)
        self.assertEqual(report.dcf.assumptions["terminal_ronic_used"], 0.15)
        wb, _ev = _workbook(self, report)
        ws = wb["DCF"]
        r = _label_rows(ws)["Terminal method"]
        self.assertEqual(ws[f"B{r}"].value, "Value driver (RONIC 15.0%)")
        self.assertGreaterEqual(ws.column_dimensions["B"].width,
                                len("Value driver (RONIC 15.0%)") + 2)
        self.assertIn("Terminal method</th><td>Value driver (RONIC 15.0%)<",
                      _html(self, report))
        # Gordon and exit-multiple runs keep the method's name.
        wb, _ev = _workbook(self, self.report)
        ws = wb["DCF"]
        self.assertEqual(ws[f"B{_label_rows(ws)['Terminal method']}"].value, "gordon")

    def test_a_value_driver_check_shows_the_reinvestment_the_tv_uses(self):
        from equity_valuation.report.excel import (PERCENT_FMT, TERMINAL_CHECK_NOTE,
                                                   TERMINAL_DIAGNOSTICS,
                                                   VALUE_DRIVER_REINVESTMENT_LABEL)

        report = _report(dcf_assumptions=DCFAssumptions(terminal_method="value_driver",
                                                        terminal_ronic=0.15),
                         run_sensitivity=False)
        a = report.dcf.assumptions
        # The TV reinvests g / RONIC of NOPAT (FCFF = NOPAT x (1 - g / RONIC)),
        # not the drivers' 6.7%, whose implied RONIC (37.1%) is not the 15% used.
        uses = a["terminal_growth_used"] / a["terminal_ronic_used"]
        self.assertAlmostEqual(uses, 1 - a["terminal_fcff"] / a["terminal_nopat"], places=12)
        self.assertEqual(f"{uses:.1%}", "16.7%")
        self.assertEqual(f"{a['terminal_reinvestment_rate']:.1%}", "6.7%")
        self.assertEqual(f"{a['implied_terminal_ronic']:.1%}", "37.1%")
        note = ("From the capex, D&A and working-capital drivers, for comparison; the "
                "value-driver TV reinvests g / RONIC = 16.7% of NOPAT at RONIC 15.0%")

        wb, ev = _workbook(self, report)
        ws = wb["DCF"]
        rows = _label_rows(ws)
        labels = [VALUE_DRIVER_REINVESTMENT_LABEL] + [k for k, _key in TERMINAL_DIAGNOSTICS]
        # Its own row first under the PV of the TV, then the drivers' check.
        self.assertEqual([rows[k] for k in labels],
                         [rows["PV of terminal value"] + i for i in (1, 2, 3, 4)])
        self.assertEqual(rows["Enterprise value"], rows["PV of terminal value"] + 5)
        r = rows[VALUE_DRIVER_REINVESTMENT_LABEL]
        self.assertAlmostEqual(ws[f"B{r}"].value, uses, places=12)
        self.assertEqual(ws[f"B{r}"].number_format, PERCENT_FMT)
        self.assertTrue(ws[f"A{r}"].value.startswith("  "))
        first = rows[TERMINAL_DIAGNOSTICS[0][0]]
        self.assertEqual(ws[f"C{first}"].value, note)
        self.assertAlmostEqual(ws[f"B{first}"].value, a["terminal_reinvestment_rate"], places=12)
        # The bridge formulas skip the check rows and still reconcile.
        self.assertAlmostEqual(ev.value("DCF", f"B{rows['Implied price / share']}"),
                               report.dcf.implied_price, delta=1e-9)

        text = _dcf_section(_html(self, report))
        self.assertIn(f"<th class='rowhead sub'>{VALUE_DRIVER_REINVESTMENT_LABEL}</th>"
                      "<td>16.7%</td>", text)
        self.assertIn("Terminal net reinvestment (% of NOPAT)</th><td>6.7%<br>"
                      f"<span class='aside'>{note.replace('&', '&amp;')}</span></td>", text)

        # Negative growth: the TV releases capital instead.
        report = _report(dcf_assumptions=DCFAssumptions(terminal_method="value_driver",
                                                        terminal_ronic=0.15,
                                                        terminal_growth=-0.02),
                         run_sensitivity=False)
        wb, _ev = _workbook(self, report)
        ws = wb["DCF"]
        rows = _label_rows(ws)
        self.assertAlmostEqual(ws[f"B{rows[VALUE_DRIVER_REINVESTMENT_LABEL]}"].value,
                               -0.02 / 0.15, places=12)
        self.assertEqual(ws[f"C{rows[TERMINAL_DIAGNOSTICS[0][0]]}"].value,
                         "From the capex, D&A and working-capital drivers, for comparison; "
                         "the value-driver TV releases -g / RONIC = 13.3% of NOPAT a year at "
                         "RONIC 15.0%")

        # A Gordon run: no such row, and the note describes the check itself.
        wb, _ev = _workbook(self, self.report)
        ws = wb["DCF"]
        rows = _label_rows(ws)
        self.assertNotIn(VALUE_DRIVER_REINVESTMENT_LABEL, rows)
        self.assertEqual(ws[f"C{rows[TERMINAL_DIAGNOSTICS[0][0]]}"].value, TERMINAL_CHECK_NOTE)
        text = _dcf_section(_html(self, self.report))
        self.assertNotIn(VALUE_DRIVER_REINVESTMENT_LABEL, text)
        self.assertNotIn("for comparison", text)

    def test_a_flagged_issuer_says_why_there_is_no_reverse_dcf(self):
        for build, whom in ((_bank_report, "a financial institution"),
                            (_lessor_report, "a debt-funded lessor")):
            report = build()
            note = f"Not meaningful for {whom}: the FCFF DCF is reference-only."
            self.assertEqual(report.reverse_dcf["note"], note)
            with self.subTest(whom):
                wb, _ev = _workbook(self, report)
                ws = wb["DCF"]
                rows = _label_rows(ws)
                self.assertIn(note, rows)
                self.assertNotIn("Driver", rows)  # no figures at all
                self.assertNotIn("Year-1 revenue growth", rows)
                summary = wb["Summary"]
                r = _label_rows(summary)["Market-implied year-1 growth"]
                self.assertEqual((summary[f"B{r}"].value, summary[f"C{r}"].value),
                                 ("n/a", note))
                section = _dcf_section(_html(self, report))
                self.assertIn(f"<p class='rdcf-note'>{note}</p>", section)
                self.assertNotIn("Market-implied</th>", section)

    def test_an_unsolved_driver_shows_na_and_the_solver_note(self):
        why = ("Market price is outside the solvable discount-rate range (2.51% to 25.00%) "
               "with the current assumptions.")
        report = dataclasses.replace(self.report, reverse_dcf=dict(self.report.reverse_dcf, wacc={
            "converged": False, "implied_wacc": None, "current_wacc": 0.095, "note": why}))
        wb, _ev = _workbook(self, report)
        ws = wb["DCF"]
        rows = _label_rows(ws)
        r = rows["Discount rate (WACC)"]
        self.assertEqual((ws[f"B{r}"].value, ws[f"C{r}"].value), (0.095, "n/a"))
        self.assertIn(f"Discount rate (WACC): {why}", rows)
        section = _dcf_section(_html(self, report))
        self.assertIn("Discount rate (WACC)</th><td>9.5%</td><td>n/a</td>", section)
        self.assertIn(f"<li>Discount rate (WACC): {why}</li>", section)
        # No reverse DCF at all (skipped, or no DCF valuation to invert).
        report = dataclasses.replace(self.report, reverse_dcf=None)
        wb, _ev = _workbook(self, report)
        self.assertIn("Not available for this run.", _label_rows(wb["DCF"]))
        self.assertIn("<p class='rdcf-note'>Not available for this run.</p>",
                      _dcf_section(_html(self, report)))


if __name__ == "__main__":
    unittest.main()
