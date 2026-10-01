"""Excel exporter: a formatted, multi-sheet .xlsx valuation workbook.

The workbook mirrors the `ValuationReport` produced by the engine. Where it is
practical we emit *live* Excel formulas (PV = FCFF * discount-factor, EV = SUM(PVs)
+ PV_terminal, implied = equity / shares, upside = implied / current - 1) that
reference real cells, so an analyst can tweak an input and let Excel recalculate.
Static values are used where a live formula would be impractical or fragile.

Workbook conventions:
  * Money is in absolute units; we format with the "#,##0.00" currency mask (and a
    currency-symbol prefix where we know the reporting currency).
  * Rates / margins / upside are decimals; we format them with the "0.0%" mask so
    0.082 renders as "8.2%".
  * Every `report.dcf / comps / ddm / fcfe` may be None -- we guard each section and
    write a human-readable "not available" note instead of crashing.
  * No network calls, no print(); we only build and save a workbook.

"""

from __future__ import annotations

import datetime as _dt
from typing import Optional

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .. import config
from ..data.market import _BLUME_MARKET_WEIGHT, _BLUME_RAW_WEIGHT
from ..schemas import ValuationReport
from ..utils import is_num, median

# --------------------------------------------------------------------------- #
#  Styling constants
# --------------------------------------------------------------------------- #
CURRENCY_FMT = "#,##0.00"        # absolute money / per-share values
PERCENT_FMT = "0.0%"             # decimals -> percent (0.082 -> 8.2%)
MULTIPLE_FMT = "0.0\"x\""        # trading multiples, e.g. 12.3x
FACTOR_FMT = "0.0000"            # discount factors

_TITLE_FONT = Font(bold=True, size=14, color="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_SUBHEADER_FONT = Font(bold=True, color="1F3864")
_LABEL_FONT = Font(bold=True)
_NOTE_FONT = Font(italic=True, color="808080")
_GREEN_FONT = Font(color="0B6E0B")   # positive upside
_RED_FONT = Font(color="C00000")     # negative upside
_BASE_CASE_FILL = PatternFill("solid", fgColor="FFF2CC")  # sensitivity cell = headline DCF
_RIGHT = Alignment(horizontal="right")
_LEFT = Alignment(horizontal="left")
_CENTER = Alignment(horizontal="center")

# --------------------------------------------------------------------------- #
#  Rules shared by every output (this workbook, the HTML report, the memo and
#  deck, and the AI context import them from here; the dashboard's
#  lib/format.ts keeps a copy that a test pins to these values)
# --------------------------------------------------------------------------- #
# The Blume adjustment the market data applies to Yahoo's raw beta.
BLUME_FORMULA = f"{_BLUME_RAW_WEIGHT:g} x raw + {_BLUME_MARKET_WEIGHT:g}"
BLUME_NOTE = f"Blume-adjusted toward 1 ({BLUME_FORMULA})"

# In place of a blended target when the engine gives none because the
# company's kind sets its methods aside (a bank's, a captive-finance group's or
# a lessor's DCF and FCFE, and a low-payout or lone DDM) and no usable peers
# were supplied: trading comps would give a target.
NO_TARGET_SUPPLY_PEERS = "No target (supply peers)"


def _has_peers(comps: object) -> bool:
    """True when trading comps ran on at least one usable peer. ``comps`` is a
    CompsResult, its JSON dict, or None (comps not run)."""
    if comps is None:
        return False
    peers = comps.get("peers") if isinstance(comps, dict) else getattr(comps, "peers", None)
    return bool(peers)


def needs_peers(summary: Optional[dict], comps: object = None) -> bool:
    """True when there is no blended target because the company is flagged
    (``financial_kind``) and every method it has is reference only, and no
    usable peers were supplied. Not when peers were supplied but comps gave no
    price, and not when an ordinary company's methods gave no valuation (0.00);
    those read 'n/a'. The engine asks for peers in the same cases."""
    summary = summary or {}
    if is_num(summary.get("blended_target")):
        return False
    methods = summary.get("methods") or {}
    if not methods or any(str(name).startswith("Comps") for name in methods):
        return False
    if not (summary.get("financial_kind") or summary.get("financial_institution")):
        return False
    return not _has_peers(comps)


def no_target_text(summary: Optional[dict], comps: object = None) -> str:
    """What to show for a missing blended target: 'No target (supply peers)'
    when ``needs_peers``, else 'n/a'."""
    return NO_TARGET_SUPPLY_PEERS if needs_peers(summary, comps) else "n/a"


def method_upside_toned(summary: Optional[dict], name: str) -> bool:
    """Whether a method's upside is coloured by its sign. Not for a method left
    out of the blended target (``excluded_from_blend``: shown for reference
    only, e.g. a lessor's DCF), and not for any method when the engine gives no
    verdict ('N/A': no target, or a verdict withheld for a blend that rests on
    the DDM alone), so no output shows a direction the engine does not give.
    ``name`` is the summary's method name ("DCF", "Comps (median)", "DDM",
    "FCFE")."""
    summary = summary or {}
    if name in (summary.get("excluded_from_blend") or {}):
        return False
    return summary.get("recommendation") != "N/A"


# Beside the terminal growth of an exit-multiple DCF, whose terminal value is
# EBITDA_N x the exit multiple (the HTML report's Key assumptions say the same).
EXIT_MULTIPLE_GROWTH_NOTE = (
    "Not used in this DCF's terminal value (the exit multiple sets it); still the "
    "rate revenue growth fades to, and the terminal growth of DDM and FCFE."
)


# --------------------------------------------------------------------------- #
#  Small cell helpers (all None-safe)
# --------------------------------------------------------------------------- #
def _num(value: object) -> Optional[float]:
    """Return a finite float for writing, else None (so the cell stays blank)."""
    return float(value) if is_num(value) else None


def _set(ws: Worksheet, row: int, col: int, value: object,
         *, fmt: Optional[str] = None, font: Optional[Font] = None,
         align: Optional[Alignment] = None, formula: bool = False) -> "Cell":  # type: ignore[name-defined]
    """Write a value into (row, col) and apply optional number format / style."""
    if isinstance(value, str):
        # Control characters (e.g. from exception text in a warning) are illegal
        # in XLSX XML and would make openpyxl refuse to write the whole file.
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell = ws.cell(row=row, column=col, value=value)
    # Company names, tickers and model notes come from external providers.
    # Preserve them as text even when they start with '='; only formulas built
    # here opt into Excel evaluation.
    if isinstance(value, str) and not formula:
        cell.data_type = "s"
    if fmt is not None:
        cell.number_format = fmt
    if font is not None:
        cell.font = font
    if align is not None:
        cell.alignment = align
    return cell


def _formula(ws: Worksheet, row: int, col: int, value: str, **style):
    """Write an internally constructed formula, never provider-supplied text."""
    return _set(ws, row, col, value, formula=True, **style)


def _title(ws: Worksheet, text: str) -> None:
    """Write the per-sheet title row (row 1)."""
    _set(ws, 1, 1, text, font=_TITLE_FONT)


def _header_row(ws: Worksheet, row: int, labels: list[str], start_col: int = 1) -> None:
    """Write a styled (bold, filled) header row."""
    for j, label in enumerate(labels):
        cell = _set(ws, row, start_col + j, label, font=_HEADER_FONT, align=_CENTER)
        cell.fill = _HEADER_FILL


def _note(ws: Worksheet, row: int, text: str, col: int = 1) -> None:
    """Write a greyed-out italic note line."""
    _set(ws, row, col, text, font=_NOTE_FONT)


def _set_widths(ws: Worksheet, widths: dict[int, float]) -> None:
    """Set column widths from a {col_index: width} map."""
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _money_fmt(report: ValuationReport) -> str:
    """Currency mask, prefixed with the reporting-currency symbol when known."""
    cur = getattr(getattr(report.company, "market", None), "currency", None)
    symbol = config.CURRENCY_SYMBOLS.get(cur) if cur else None
    if symbol:
        # Quote the symbol so non-ASCII (€, £, ¥) is treated literally by Excel.
        return f'"{symbol}"{CURRENCY_FMT}'
    return CURRENCY_FMT


# Token fragments that, when present in a detail key name, mark it as a rate
# (rendered as a percent). Everything else -- d0, dps, *_pv, *_value, price,
# stage PVs -- is money.
_RATE_KEY_TOKENS = (
    "growth", "rate", "ke", "coe", "roe", "retention", "yield", "wacc", "cost_of",
)


def _is_rate_key(key: object) -> bool:
    """True if a detail key's NAME indicates a rate (-> percent format)."""
    name = str(key).lower()
    if name == "g" or name.endswith("_g"):
        return True
    return any(tok in name for tok in _RATE_KEY_TOKENS)


def _is_years_key(key: object) -> bool:
    """True if a detail key's NAME is a horizon in years (e.g. high_growth_years)."""
    name = str(key).lower()
    return name.endswith("years") or name == "h_half"


def _detail_fmt(key: object, value: float, money_fmt: str) -> str:
    """Number format for one numeric model-detail entry, chosen by key name."""
    if _is_years_key(key):
        return "0" if float(value).is_integer() else "0.0"
    return PERCENT_FMT if _is_rate_key(key) else money_fmt


def _year_labels(report: ValuationReport, years: list) -> list[str]:
    """Column headers for a projection table.

    Calendar years render as "FY 2025". Relative indices (the DCF numbers its
    forecast 1..n from the latest fiscal year) are mapped onto that fiscal year
    when it is known, so the DCF and FCFE tables label the same periods alike;
    otherwise they render as "Year 1".
    """
    fin = getattr(report.company, "financials", None)
    fys = [int(y) for y in (getattr(fin, "fiscal_years", None) or []) if is_num(y)]
    last_fy = fys[-1] if fys else None
    labels = []
    for y in years:
        if is_num(y) and y >= 1000:
            labels.append(f"FY {int(y)}")
        elif is_num(y) and last_fy is not None:
            labels.append(f"FY {last_fy + int(y)}")
        else:
            labels.append(f"Year {y}")
    return labels


# --------------------------------------------------------------------------- #
#  Sheet: Summary
# --------------------------------------------------------------------------- #
def _write_summary(ws: Worksheet, report: ValuationReport, money_fmt: str) -> None:
    company = report.company
    market = getattr(company, "market", None)
    name = getattr(company, "name", None) or getattr(company, "ticker", "") or ""
    ticker = getattr(company, "ticker", "") or ""

    _title(ws, f"Valuation Summary — {name} ({ticker})")

    row = 3
    _set(ws, row, 1, "Company", font=_LABEL_FONT)
    _set(ws, row, 2, name)
    row += 1
    _set(ws, row, 1, "Ticker", font=_LABEL_FONT)
    _set(ws, row, 2, ticker)
    row += 1
    _set(ws, row, 1, "Sector", font=_LABEL_FONT)
    _set(ws, row, 2, getattr(market, "sector", None) or "n/a")
    row += 1
    _set(ws, row, 1, "Currency", font=_LABEL_FONT)
    _set(ws, row, 2, getattr(market, "currency", None) or "n/a")
    row += 1
    _set(ws, row, 1, "Current price", font=_LABEL_FONT)
    cur_price = _num(report.current_price)
    _set(ws, row, 2, cur_price, fmt=money_fmt, align=_RIGHT)
    current_price_cell = f"B{row}"
    row += 1
    _set(ws, row, 1, "Report date", font=_LABEL_FONT)
    _set(ws, row, 2, _dt.date.today().isoformat())
    row += 2

    # --- Method valuation table ------------------------------------------- #
    _set(ws, row, 1, "Valuation by method", font=_SUBHEADER_FONT)
    row += 1
    _header_row(ws, row, ["Method", "Implied price", "Upside vs. current"])
    row += 1

    # Methods the engine shows for reference but leaves out of the blended
    # target (e.g. a bank's DCF and FCFE), keyed by its summary method name.
    summary = getattr(report, "summary", None) or {}
    excluded = summary.get("excluded_from_blend") or {}

    def _method_line(label: str, implied: object, key: str) -> None:
        nonlocal row
        imp = _num(implied)
        if key in excluded:
            label = f"{label} (not in blend)"
            _note(ws, row, f"Not in blend: {excluded[key]}", col=4)
        _set(ws, row, 1, label)
        _set(ws, row, 2, imp, fmt=money_fmt, align=_RIGHT)
        if imp is not None and cur_price:
            # Live upside formula referencing the implied-price cell and current price.
            up_cell = _formula(ws, row, 3, f"=B{row}/{current_price_cell}-1",
                           fmt=PERCENT_FMT, align=_RIGHT)
            # Best-effort sign coloring (Excel won't recolor on edit; this is the
            # value as-computed now -- a "plus", per the contract). None for a
            # method left out of the blend or when the engine gives no verdict.
            up_val = imp / cur_price - 1.0
            if method_upside_toned(summary, key):
                up_cell.font = _GREEN_FONT if up_val >= 0 else _RED_FONT
        else:
            _set(ws, row, 3, "n/a", align=_RIGHT)
        row += 1

    dcf = report.dcf
    comps = report.comps
    ddm = report.ddm
    fcfe = report.fcfe

    _method_line("DCF (FCFF)", getattr(dcf, "implied_price", None) if dcf else None, "DCF")
    comps_med = None
    if comps is not None:
        comps_med = (comps.implied_price_summary or {}).get("median")
    _method_line("Trading comps (median)", comps_med, "Comps (median)")
    _method_line("DDM", getattr(ddm, "implied_price", None) if ddm else None, "DDM")
    _method_line("FCFE", getattr(fcfe, "implied_price", None) if fcfe else None, "FCFE")

    # Blended target: prefer the engine-computed value on report.summary so the
    # Excel and HTML headline targets always agree (a None there means no
    # target, written as 'n/a' or 'No target (supply peers)', never a number);
    # fall back to the local median-of-methods only when that key is absent.
    if "blended_target" in summary:
        blended = _num(summary.get("blended_target"))
    else:
        method_prices = [
            _num(getattr(dcf, "implied_price", None) if dcf else None),
            _num(comps_med),
            _num(getattr(ddm, "implied_price", None) if ddm else None),
            _num(getattr(fcfe, "implied_price", None) if fcfe else None),
        ]
        blended = median([p for p in method_prices if p is not None])
    _set(ws, row, 1, "Blended target (median)", font=_LABEL_FONT)
    blended = _num(blended)
    if blended is None:
        # No target: say so in words (never a number), with no upside.
        _set(ws, row, 2, no_target_text(summary, comps), font=_LABEL_FONT, align=_RIGHT)
        _set(ws, row, 3, "n/a", align=_RIGHT)
    else:
        _set(ws, row, 2, blended, fmt=money_fmt, font=_LABEL_FONT, align=_RIGHT)
    # The engine withholds the upside (None) when it gives no verdict, e.g. for
    # a blend that rests on the DDM alone: the target stays, the upside is n/a.
    withheld = "blended_upside" in summary and _num(summary.get("blended_upside")) is None
    if blended is not None and cur_price and not withheld:
        up_cell = _formula(ws, row, 3, f"=B{row}/{current_price_cell}-1",
                       fmt=PERCENT_FMT, align=_RIGHT)
        # Prefer the engine-computed blended_upside for sign-coloring (keeps the
        # Excel/HTML headline consistent); fall back to the local computation.
        if "blended_upside" in summary:
            up_val = _num(summary.get("blended_upside"))
        else:
            up_val = blended / cur_price - 1.0
        if up_val is not None:
            up_cell.font = _GREEN_FONT if up_val >= 0 else _RED_FONT
    elif blended is not None:
        _set(ws, row, 3, "n/a", align=_RIGHT)
    row += 1
    _set(ws, row, 1, "Verdict", font=_LABEL_FONT)
    _set(ws, row, 2, summary.get("recommendation") or "n/a", align=_RIGHT)
    row += 2

    # --- Football-field ranges -------------------------------------------- #
    _set(ws, row, 1, "Valuation ranges (football field)", font=_SUBHEADER_FONT)
    row += 1
    ff = report.football_field or []
    if ff:
        _header_row(ws, row, ["Method", "Low", "Base", "High"])
        row += 1
        for r in ff:
            _set(ws, row, 1, getattr(r, "method", "") or "")
            _set(ws, row, 2, _num(getattr(r, "low", None)), fmt=money_fmt, align=_RIGHT)
            _set(ws, row, 3, _num(getattr(r, "base", None)), fmt=money_fmt, align=_RIGHT)
            _set(ws, row, 4, _num(getattr(r, "high", None)), fmt=money_fmt, align=_RIGHT)
            row += 1
    else:
        _note(ws, row, "Football-field ranges not available.")
        row += 1
    row += 1

    # --- Warnings ---------------------------------------------------------- #
    warnings = list(report.warnings or [])
    if warnings:
        _set(ws, row, 1, "Warnings", font=_SUBHEADER_FONT)
        row += 1
        for w in warnings:
            _note(ws, row, f"• {w}")
            row += 1

    # Column B is wide enough for the bold, right-aligned 'No target (supply
    # peers)', which cannot overflow into column A's label.
    _set_widths(ws, {1: 28, 2: 26, 3: 18, 4: 18})


# --------------------------------------------------------------------------- #
#  Sheet: DCF
# --------------------------------------------------------------------------- #
def _beta_label(report: ValuationReport) -> tuple[str, Optional[str]]:
    """(row label, note) for the DCF's beta.

    The market data Blume-adjusts Yahoo's beta toward 1 (``BLUME_FORMULA``) and
    keeps the raw figure (``MarketData.raw_beta``; the WACC records it as
    ``detail['beta_raw']``), so an adjusted beta reads 'Beta (adj.)' with the
    raw value in a note. The HTML report's Key assumptions say the same.
    """
    wacc = getattr(getattr(report, "dcf", None), "wacc", None)
    beta = getattr(wacc, "beta", None)
    detail = getattr(wacc, "detail", None) or {}
    if detail.get("beta_source") == "DEFAULT_BETA":
        return "Beta (default)", "No usable market beta."
    raw = detail.get("beta_raw")
    if not is_num(raw):
        raw = getattr(getattr(report.company, "market", None), "raw_beta", None)
    if is_num(beta) and is_num(raw) and abs(float(raw) - float(beta)) > 1e-9:
        return "Beta (adj.)", f"Raw beta {float(raw):.3f}, {BLUME_NOTE}"
    if "adjusted" in str(detail.get("beta_source") or ""):
        return "Beta (adj.)", "Adjusted by the market data (see the notes)."
    return "Beta", None


def _write_dcf(ws: Worksheet, report: ValuationReport, money_fmt: str) -> None:
    _title(ws, "Discounted Cash Flow (Unlevered FCFF)")
    dcf = report.dcf
    if dcf is None:
        _note(ws, 3, "DCF model not available for this company.")
        _set_widths(ws, {1: 30})
        return

    assumptions = dict(getattr(dcf, "assumptions", {}) or {})
    wacc_res = getattr(dcf, "wacc", None)

    # --- Assumptions block ------------------------------------------------- #
    row = 3
    _set(ws, row, 1, "Key assumptions", font=_SUBHEADER_FONT)
    row += 1

    # WACC and its components from the WACCResult (with graceful fallbacks).
    wacc_val = _num(getattr(wacc_res, "wacc", None)) if wacc_res else None
    _set(ws, row, 1, "WACC", font=_LABEL_FONT)
    wacc_cell_ref = f"B{row}"
    _set(ws, row, 2, wacc_val, fmt=PERCENT_FMT, align=_RIGHT)
    row += 1
    if wacc_res is not None:
        beta_label, beta_note = _beta_label(report)
        for label, attr, fmt in (
            ("Cost of equity", "cost_of_equity", PERCENT_FMT),
            ("After-tax cost of debt", "after_tax_cost_of_debt", PERCENT_FMT),
            ("Weight equity", "weight_equity", PERCENT_FMT),
            ("Weight debt", "weight_debt", PERCENT_FMT),
            (beta_label, "beta", "0.00"),
        ):
            _set(ws, row, 1, label)
            _set(ws, row, 2, _num(getattr(wacc_res, attr, None)), fmt=fmt, align=_RIGHT)
            if attr == "beta" and beta_note:
                _note(ws, row, beta_note, col=3)
            row += 1

    # Selected assumption-dict entries (terminal method/growth, tax, mid-year).
    term_method = assumptions.get("terminal_method")
    # The growth actually used (the model clamps it below WACC when needed).
    term_growth = assumptions.get("terminal_growth_used", assumptions.get("terminal_growth"))
    tax_rate = assumptions.get("tax_rate")
    mid_year = assumptions.get("mid_year_convention")
    _set(ws, row, 1, "Terminal method")
    _set(ws, row, 2, str(term_method) if term_method is not None else "n/a", align=_RIGHT)
    row += 1
    # The multiple that sets an exit-multiple terminal value (and centres the
    # exit EV/EBITDA sensitivity axis).
    exit_mult = assumptions.get("exit_ev_ebitda")
    if term_method == "exit_multiple" and is_num(exit_mult):
        _set(ws, row, 1, "Exit EV/EBITDA")
        _set(ws, row, 2, _num(exit_mult), fmt=MULTIPLE_FMT, align=_RIGHT)
        row += 1
    _set(ws, row, 1, "Terminal growth")
    term_growth_cell_ref = f"B{row}"
    _set(ws, row, 2, _num(term_growth), fmt=PERCENT_FMT, align=_RIGHT)
    if term_method == "exit_multiple":
        # The multiple alone sets this DCF's terminal value; the growth rate
        # still ends the revenue-growth fade and drives DDM and FCFE.
        _note(ws, row, EXIT_MULTIPLE_GROWTH_NOTE, col=3)
    row += 1
    _set(ws, row, 1, "Tax rate")
    _set(ws, row, 2, _num(tax_rate), fmt=PERCENT_FMT, align=_RIGHT)
    row += 1
    _set(ws, row, 1, "Mid-year convention")
    _set(ws, row, 2, ("Yes" if mid_year else "No") if mid_year is not None else "n/a",
         align=_RIGHT)
    row += 2

    # --- Projection table (metrics as ROWS, forecast years as COLUMNS) ----- #
    years = list(getattr(dcf, "years", []) or [])
    n = len(years)
    revenue = list(getattr(dcf, "revenue", []) or [])
    ebit = list(getattr(dcf, "ebit", []) or [])
    nopat = list(getattr(dcf, "nopat", []) or [])
    fcff = list(getattr(dcf, "fcff", []) or [])
    dfs = list(getattr(dcf, "discount_factors", []) or [])
    pv_fcff = list(getattr(dcf, "pv_fcff", []) or [])

    _set(ws, row, 1, "FCFF projection", font=_SUBHEADER_FONT)
    row += 1
    table_top = row  # header row of the projection table

    # Header: metric label column + one column per forecast year.
    _header_row(ws, table_top, ["(values in reporting currency)"]
                + _year_labels(report, years))
    # Column index of the first data year (column 2 = "B").
    first_year_col = 2

    def _series_row(r: int, label: str, series: list, fmt: str) -> None:
        _set(ws, r, 1, label, font=_LABEL_FONT)
        for j in range(n):
            val = _num(series[j]) if j < len(series) else None
            _set(ws, r, first_year_col + j, val, fmt=fmt, align=_RIGHT)

    body = table_top + 1
    r_rev = body
    _series_row(r_rev, "Revenue", revenue, money_fmt)
    # Revenue growth (live formula vs. prior year; first year derived from
    # historical base which we don't store here -> static blank / value).
    r_growth = body + 1
    _set(ws, r_growth, 1, "  growth %", font=_LABEL_FONT)
    for j in range(n):
        col = first_year_col + j
        if j == 0:
            _set(ws, r_growth, col, None, fmt=PERCENT_FMT, align=_RIGHT)
        else:
            prev = get_column_letter(col - 1)
            cur = get_column_letter(col)
            # Guarded: the model projects zero revenue when it has no base.
            _formula(ws, r_growth, col,
                 f'=IF({prev}{r_rev}=0,"",{cur}{r_rev}/{prev}{r_rev}-1)',
                 fmt=PERCENT_FMT, align=_RIGHT)

    r_ebit = body + 2
    _series_row(r_ebit, "EBIT", ebit, money_fmt)
    # EBIT margin = EBIT / Revenue (live formula).
    r_margin = body + 3
    _set(ws, r_margin, 1, "  EBIT margin %", font=_LABEL_FONT)
    for j in range(n):
        col = get_column_letter(first_year_col + j)
        _formula(ws, r_margin, first_year_col + j,
             f'=IF({col}{r_rev}=0,"",{col}{r_ebit}/{col}{r_rev})',
             fmt=PERCENT_FMT, align=_RIGHT)

    r_nopat = body + 4
    _series_row(r_nopat, "NOPAT", nopat, money_fmt)

    # D&A, Capex, ΔNWC: not stored on DCFResult, so back them out where possible.
    # FCFF = NOPAT + D&A - Capex - ΔNWC. We display the FCFF directly (the engine
    # already computed it) and leave the individual add-backs blank with a note,
    # rather than fabricating numbers that don't reconcile.
    r_da = body + 5
    _set(ws, r_da, 1, "D&A", font=_LABEL_FONT)
    r_capex = body + 6
    _set(ws, r_capex, 1, "Capex", font=_LABEL_FONT)
    r_nwc = body + 7
    _set(ws, r_nwc, 1, "Δ NWC", font=_LABEL_FONT)
    for j in range(n):
        for rr in (r_da, r_capex, r_nwc):
            _set(ws, rr, first_year_col + j, None, fmt=money_fmt, align=_RIGHT)

    r_fcff = body + 8
    _series_row(r_fcff, "FCFF", fcff, money_fmt)

    r_df = body + 9
    _series_row(r_df, "Discount factor", dfs, FACTOR_FMT)

    # PV of FCFF -- LIVE formula PV = FCFF * discount factor (recalcs on edit).
    r_pv = body + 10
    _set(ws, r_pv, 1, "PV of FCFF", font=_LABEL_FONT)
    for j in range(n):
        col = get_column_letter(first_year_col + j)
        # If we have both inputs in-sheet, use a formula; else fall back to value.
        if j < len(fcff) and j < len(dfs) and _num(fcff[j]) is not None \
                and _num(dfs[j]) is not None:
            _formula(ws, r_pv, first_year_col + j, f"={col}{r_fcff}*{col}{r_df}",
                 fmt=money_fmt, align=_RIGHT)
        else:
            val = _num(pv_fcff[j]) if j < len(pv_fcff) else None
            _set(ws, r_pv, first_year_col + j, val, fmt=money_fmt, align=_RIGHT)

    _note(ws, r_pv + 1,
          "D&A / Capex / ΔNWC components are summarized within FCFF "
          "(not broken out on the result object).")

    # --- Valuation bridge -------------------------------------------------- #
    row = r_pv + 3
    _set(ws, row, 1, "Valuation bridge", font=_SUBHEADER_FONT)
    row += 1

    # PV(explicit FCFF) = SUM of the PV row across the forecast columns.
    last_year_col = get_column_letter(first_year_col + n - 1) if n else "B"
    first_year_col_letter = get_column_letter(first_year_col)
    _set(ws, row, 1, "Σ PV of explicit FCFF", font=_LABEL_FONT)
    if n:
        sum_pv_formula = f"=SUM({first_year_col_letter}{r_pv}:{last_year_col}{r_pv})"
        _formula(ws, row, 2, sum_pv_formula, fmt=money_fmt, align=_RIGHT)
    else:
        _set(ws, row, 2, None, fmt=money_fmt, align=_RIGHT)
    sum_pv_cell = f"B{row}"
    row += 1

    _set(ws, row, 1, "Terminal value (undiscounted)", font=_LABEL_FONT)
    _set(ws, row, 2, _num(getattr(dcf, "terminal_value", None)),
         fmt=money_fmt, align=_RIGHT)
    row += 1

    _set(ws, row, 1, "PV of terminal value", font=_LABEL_FONT)
    _set(ws, row, 2, _num(getattr(dcf, "pv_terminal", None)),
         fmt=money_fmt, align=_RIGHT)
    pv_terminal_cell = f"B{row}"
    row += 1

    # Enterprise value = Σ PV explicit + PV terminal  (LIVE formula).
    _set(ws, row, 1, "Enterprise value", font=_LABEL_FONT)
    _formula(ws, row, 2, f"={sum_pv_cell}+{pv_terminal_cell}", fmt=money_fmt, align=_RIGHT)
    ev_cell = f"B{row}"
    row += 1

    _set(ws, row, 1, "Less: net debt", font=_LABEL_FONT)
    _set(ws, row, 2, _num(getattr(dcf, "net_debt", None)), fmt=money_fmt, align=_RIGHT)
    net_debt_cell = f"B{row}"
    row += 1

    # The model bridges EV to common equity through every senior claim, not just
    # net debt (models/dcf.py): minority interest and preferred equity too.
    bs = getattr(report.company, "balance_sheet", None)
    claim_cells = []
    for label, attr in (("Less: minority interest", "minority_interest"),
                        ("Less: preferred equity", "preferred_equity")):
        claim = _num(getattr(bs, attr, None)) if bs is not None else None
        _set(ws, row, 1, label, font=_LABEL_FONT)
        _set(ws, row, 2, claim if claim is not None else 0.0, fmt=money_fmt, align=_RIGHT)
        claim_cells.append(f"B{row}")
        row += 1

    # Equity value = EV - net debt - minority - preferred  (LIVE formula).
    _set(ws, row, 1, "Equity value", font=_LABEL_FONT)
    _formula(ws, row, 2, f"={ev_cell}-{net_debt_cell}-" + "-".join(claim_cells),
         fmt=money_fmt, align=_RIGHT)
    equity_cell = f"B{row}"
    row += 1

    _set(ws, row, 1, "Shares outstanding", font=_LABEL_FONT)
    _set(ws, row, 2, _num(getattr(dcf, "shares", None)), fmt="#,##0", align=_RIGHT)
    shares_cell = f"B{row}"
    row += 1

    # Implied price = equity value / shares  (LIVE formula).
    _set(ws, row, 1, "Implied price / share", font=_LABEL_FONT)
    shares_val = _num(getattr(dcf, "shares", None))
    if shares_val:
        _formula(ws, row, 2, f"={equity_cell}/{shares_cell}", fmt=money_fmt, align=_RIGHT)
    else:
        _set(ws, row, 2, _num(getattr(dcf, "implied_price", None)),
             fmt=money_fmt, align=_RIGHT)
    implied_cell = f"B{row}"
    row += 1

    _set(ws, row, 1, "Current price", font=_LABEL_FONT)
    _set(ws, row, 2, _num(getattr(dcf, "current_price", None)),
         fmt=money_fmt, align=_RIGHT)
    current_cell = f"B{row}"
    row += 1

    # Upside = implied / current - 1  (LIVE formula, sign-colored).
    _set(ws, row, 1, "Upside / (downside)", font=_LABEL_FONT)
    up_val = _num(getattr(dcf, "upside", None))
    cur_p = _num(getattr(dcf, "current_price", None))
    if cur_p:
        up_cell = _formula(ws, row, 2, f"={implied_cell}/{current_cell}-1",
                       fmt=PERCENT_FMT, align=_RIGHT)
    else:
        up_cell = _set(ws, row, 2, up_val, fmt=PERCENT_FMT, align=_RIGHT)
    # Uncoloured for a reference-only DCF or when the engine gives no verdict.
    if up_val is not None and method_upside_toned(report.summary, "DCF"):
        up_cell.font = _GREEN_FONT if up_val >= 0 else _RED_FONT

    _set_widths(ws, {1: 26, **{c: 16 for c in range(2, max(3, n + 2))}})


# --------------------------------------------------------------------------- #
#  Sheet: Comps
# --------------------------------------------------------------------------- #
_COMP_COLS = [
    ("ticker", "Ticker", None),
    ("name", "Name", None),
    ("market_cap", "Market cap", "#,##0"),
    ("enterprise_value", "EV", "#,##0"),
    ("ev_ebitda", "EV/EBITDA", MULTIPLE_FMT),
    ("ev_sales", "EV/Sales", MULTIPLE_FMT),
    ("pe", "P/E", MULTIPLE_FMT),
    ("pb", "P/B", MULTIPLE_FMT),
    ("peg", "PEG", "0.00"),
    ("currency", "Currency", None),
]
# Multiples that participate in the stats / implied tables.
_STAT_MULTIPLES = ["ev_ebitda", "ev_sales", "pe", "pb", "peg"]
_STAT_LABELS = {
    "ev_ebitda": "EV/EBITDA", "ev_sales": "EV/Sales",
    "pe": "P/E", "pb": "P/B", "peg": "PEG",
}


def _write_comps(ws: Worksheet, report: ValuationReport, money_fmt: str) -> None:
    _title(ws, "Trading Comparables")
    comps = report.comps
    if comps is None:
        _note(ws, 3, "Trading-comps analysis not available for this company.")
        _set_widths(ws, {1: 30})
        return

    row = 3

    # --- Target + peer multiples table ------------------------------------ #
    _set(ws, row, 1, "Multiples table", font=_SUBHEADER_FONT)
    row += 1
    _header_row(ws, row, [label for _, label, _ in _COMP_COLS])
    row += 1

    def _write_comp_row(r: int, comp_row, *, bold: bool = False) -> None:
        for j, (attr, _label, fmt) in enumerate(_COMP_COLS):
            val = getattr(comp_row, attr, None)
            if attr in ("ticker", "name", "currency"):
                cell = _set(ws, r, j + 1, val or "")
            else:
                cell = _set(ws, r, j + 1, _num(val), fmt=fmt, align=_RIGHT)
            if bold:
                cell.font = _LABEL_FONT

    target = getattr(comps, "target", None)
    if target is not None:
        _write_comp_row(row, target, bold=True)
        row += 1

    for peer in (getattr(comps, "peers", None) or []):
        _write_comp_row(row, peer)
        row += 1
    row += 1

    # --- Stats table (median / mean / min / max / p25 / p75) -------------- #
    stats = getattr(comps, "stats", None) or {}
    _set(ws, row, 1, "Peer statistics", font=_SUBHEADER_FONT)
    row += 1
    stat_keys = ["median", "mean", "min", "max", "p25", "p75"]
    _header_row(ws, row, ["Multiple"] + [k.upper() for k in stat_keys])
    row += 1
    for mult in _STAT_MULTIPLES:
        s = stats.get(mult) or {}
        _set(ws, row, 1, _STAT_LABELS[mult], font=_LABEL_FONT)
        fmt = "0.00" if mult == "peg" else MULTIPLE_FMT
        for j, k in enumerate(stat_keys):
            _set(ws, row, 2 + j, _num(s.get(k)), fmt=fmt, align=_RIGHT)
        row += 1
    row += 1

    # --- Implied prices per multiple -------------------------------------- #
    implied = getattr(comps, "implied", None) or {}
    _set(ws, row, 1, "Implied price by multiple", font=_SUBHEADER_FONT)
    row += 1
    _header_row(ws, row, ["Multiple", "Implied price"])
    row += 1
    for mult in _STAT_MULTIPLES:
        _set(ws, row, 1, _STAT_LABELS[mult])
        _set(ws, row, 2, _num(implied.get(mult)), fmt=money_fmt, align=_RIGHT)
        row += 1

    summary = getattr(comps, "implied_price_summary", None) or {}
    for label, key in (("Low", "low"), ("Median", "median"), ("High", "high")):
        _set(ws, row, 1, f"Summary — {label}", font=_LABEL_FONT)
        _set(ws, row, 2, _num(summary.get(key)), fmt=money_fmt, align=_RIGHT)
        row += 1
    row += 1

    # --- Notes ------------------------------------------------------------- #
    notes = getattr(comps, "notes", None) or []
    if notes:
        _set(ws, row, 1, "Notes", font=_SUBHEADER_FONT)
        row += 1
        for nline in notes:
            _note(ws, row, f"• {nline}")
            row += 1

    _set_widths(ws, {1: 22, 2: 26, 3: 16, 4: 16, 5: 12, 6: 12,
                     7: 12, 8: 12, 9: 12, 10: 12})


# --------------------------------------------------------------------------- #
#  Sheet: DDM_FCFE
# --------------------------------------------------------------------------- #
def _write_ddm_fcfe(ws: Worksheet, report: ValuationReport, money_fmt: str) -> None:
    _title(ws, "Dividend Discount & FCFE Models")
    ddm = report.ddm
    fcfe = report.fcfe

    row = 3

    # --- DDM block --------------------------------------------------------- #
    _set(ws, row, 1, "Dividend Discount Model", font=_SUBHEADER_FONT)
    row += 1
    if ddm is None:
        _note(ws, row, "DDM not available (company pays no dividend or data missing).")
        row += 2
    else:
        _set(ws, row, 1, "Method", font=_LABEL_FONT)
        _set(ws, row, 2, getattr(ddm, "method", None) or "n/a", align=_RIGHT)
        row += 1
        _set(ws, row, 1, "Cost of equity", font=_LABEL_FONT)
        _set(ws, row, 2, _num(getattr(ddm, "cost_of_equity", None)),
             fmt=PERCENT_FMT, align=_RIGHT)
        row += 1
        _set(ws, row, 1, "Implied price", font=_LABEL_FONT)
        _set(ws, row, 2, _num(getattr(ddm, "implied_price", None)),
             fmt=money_fmt, align=_RIGHT)
        row += 1
        # Spill any scalar detail entries (growth inputs, D0, stage PVs, ...).
        detail = getattr(ddm, "detail", None) or {}
        if detail:
            _set(ws, row, 1, "Detail", font=_LABEL_FONT)
            row += 1
            for key, val in detail.items():
                _set(ws, row, 1, f"  {key}")
                if is_num(val):
                    # Classify by KEY NAME, not magnitude: a $0.96 dividend or an
                    # $0.85 per-share PV must not render as "96.0%"/"85.0%". Only
                    # keys whose name signals a rate get the percent mask.
                    _set(ws, row, 2, float(val), fmt=_detail_fmt(key, val, money_fmt),
                         align=_RIGHT)
                elif isinstance(val, (list, tuple)) and val and all(is_num(v) for v in val):
                    # Per-year series (dividends, stage PVs): one value per column.
                    for j, v in enumerate(val):
                        _set(ws, row, 2 + j, float(v),
                             fmt=_detail_fmt(key, v, money_fmt), align=_RIGHT)
                elif isinstance(val, (list, tuple)):
                    # Text lists (notes): readable text, not a Python repr.
                    text = "; ".join(str(v) for v in val) if val else "none"
                    _set(ws, row, 2, text, align=_RIGHT)
                else:
                    _set(ws, row, 2, str(val), align=_RIGHT)
                row += 1
        row += 1

    # --- FCFE block -------------------------------------------------------- #
    _set(ws, row, 1, "Free Cash Flow to Equity (FCFE)", font=_SUBHEADER_FONT)
    row += 1
    if fcfe is None:
        _note(ws, row, "FCFE model not available for this company.")
        _set_widths(ws, {1: 26, 2: 18, 3: 16, 4: 16, 5: 16, 6: 16, 7: 16})
        return

    _set(ws, row, 1, "Cost of equity", font=_LABEL_FONT)
    _set(ws, row, 2, _num(getattr(fcfe, "cost_of_equity", None)),
         fmt=PERCENT_FMT, align=_RIGHT)
    row += 2

    years = list(getattr(fcfe, "years", []) or [])
    n = len(years)
    fcfe_series = list(getattr(fcfe, "fcfe", []) or [])
    pv_series = list(getattr(fcfe, "pv_fcfe", []) or [])

    # Projection table: metrics as rows, forecast years as columns.
    _header_row(ws, row, ["(reporting currency)"] + _year_labels(report, years))
    table_top = row
    first_col = 2
    row += 1

    r_fcfe = row
    _set(ws, r_fcfe, 1, "FCFE", font=_LABEL_FONT)
    for j in range(n):
        _set(ws, r_fcfe, first_col + j,
             _num(fcfe_series[j]) if j < len(fcfe_series) else None,
             fmt=money_fmt, align=_RIGHT)
    row += 1

    r_pv = row
    _set(ws, r_pv, 1, "PV of FCFE", font=_LABEL_FONT)
    for j in range(n):
        _set(ws, r_pv, first_col + j,
             _num(pv_series[j]) if j < len(pv_series) else None,
             fmt=money_fmt, align=_RIGHT)
    row += 2

    # Valuation bridge with a live Σ-PV formula.
    last_col = get_column_letter(first_col + n - 1) if n else "B"
    first_col_letter = get_column_letter(first_col)
    _set(ws, row, 1, "Σ PV of explicit FCFE", font=_LABEL_FONT)
    if n:
        _formula(ws, row, 2, f"=SUM({first_col_letter}{r_pv}:{last_col}{r_pv})",
             fmt=money_fmt, align=_RIGHT)
    else:
        _set(ws, row, 2, None, fmt=money_fmt, align=_RIGHT)
    sum_pv_cell = f"B{row}"
    row += 1

    _set(ws, row, 1, "Terminal value (undiscounted)", font=_LABEL_FONT)
    _set(ws, row, 2, _num(getattr(fcfe, "terminal_value", None)),
         fmt=money_fmt, align=_RIGHT)
    row += 1

    _set(ws, row, 1, "PV of terminal value", font=_LABEL_FONT)
    _set(ws, row, 2, _num(getattr(fcfe, "pv_terminal", None)),
         fmt=money_fmt, align=_RIGHT)
    pv_term_cell = f"B{row}"
    row += 1

    # Equity value = Σ PV + PV terminal  (LIVE formula).
    _set(ws, row, 1, "Equity value", font=_LABEL_FONT)
    _formula(ws, row, 2, f"={sum_pv_cell}+{pv_term_cell}", fmt=money_fmt, align=_RIGHT)
    equity_cell = f"B{row}"
    row += 1

    _set(ws, row, 1, "Shares outstanding", font=_LABEL_FONT)
    _set(ws, row, 2, _num(getattr(fcfe, "shares", None)), fmt="#,##0", align=_RIGHT)
    shares_cell = f"B{row}"
    row += 1

    # Implied price = equity / shares  (LIVE formula).
    _set(ws, row, 1, "Implied price / share", font=_LABEL_FONT)
    if _num(getattr(fcfe, "shares", None)):
        _formula(ws, row, 2, f"={equity_cell}/{shares_cell}", fmt=money_fmt, align=_RIGHT)
    else:
        _set(ws, row, 2, _num(getattr(fcfe, "implied_price", None)),
             fmt=money_fmt, align=_RIGHT)
    implied_cell = f"B{row}"
    row += 1

    _set(ws, row, 1, "Current price", font=_LABEL_FONT)
    _set(ws, row, 2, _num(getattr(fcfe, "current_price", None)),
         fmt=money_fmt, align=_RIGHT)
    current_cell = f"B{row}"
    row += 1

    # Upside = implied / current - 1  (LIVE formula, sign-colored).
    _set(ws, row, 1, "Upside / (downside)", font=_LABEL_FONT)
    cur_p = _num(getattr(fcfe, "current_price", None))
    imp = _num(getattr(fcfe, "implied_price", None))
    if cur_p:
        up_cell = _formula(ws, row, 2, f"={implied_cell}/{current_cell}-1",
                       fmt=PERCENT_FMT, align=_RIGHT)
        # Uncoloured for a reference-only FCFE or when the engine gives no verdict.
        if imp is not None and method_upside_toned(report.summary, "FCFE"):
            up_cell.font = _GREEN_FONT if (imp / cur_p - 1.0) >= 0 else _RED_FONT
    else:
        _set(ws, row, 2, None, fmt=PERCENT_FMT, align=_RIGHT)

    _set_widths(ws, {1: 28, **{c: 16 for c in range(2, max(3, n + 2))}})


# --------------------------------------------------------------------------- #
#  Sheet: Sensitivity
# --------------------------------------------------------------------------- #
_RATE_AXIS_TOKENS = ("wacc", "growth", "margin", "rate", "discount")

# Why a sensitivity cell can be blank (models/sensitivity.py stores NaN for a
# cell it cannot price at its own row/column inputs). The HTML report shows the
# same text under its heatmaps.
BLANK_SENSITIVITY_CELL_NOTE = (
    "Blank cell = no valid DCF price at those inputs: terminal growth too close "
    "to WACC, a non-positive WACC the model would replace with its fallback rate, "
    "an exit multiple of zero or less, or no base EBIT margin to vary."
)


def _is_multiple_axis(label: object) -> bool:
    """True for a sensitivity axis that varies a valuation multiple.

    The shared contract: labels containing 'EV/EBITDA' or 'multiple' (e.g. the
    'Exit EV/EBITDA' axis of an exit-multiple DCF grid) hold multiples such as
    12.0; every other axis the engine builds holds a decimal rate.
    """
    low = str(label or "").lower()
    return "ev/ebitda" in low or "multiple" in low


def _axis_fmt(label: object, value: object) -> str:
    """Number format for one sensitivity axis level, chosen by the axis label.

    Multiple axes -> 12.0x; rate axes (WACC, growth, margin) -> percent whatever
    the magnitude, so a -152% margin does not show as -1.52. Only an axis with
    an unrecognised label falls back on magnitude (|x| < 1 -> percent).
    """
    if _is_multiple_axis(label):
        return MULTIPLE_FMT
    low = str(label or "").lower()
    if any(tok in low for tok in _RATE_AXIS_TOKENS):
        return PERCENT_FMT
    return PERCENT_FMT if (is_num(value) and abs(float(value)) < 1) else "0.00"


def _axis_value(value: object) -> object:
    """An axis level for writing: the number, or 'n/a' when it is missing."""
    num = _num(value)
    return "n/a" if num is None else num


def _base_case_cell(sens: object, report: ValuationReport) -> Optional[tuple[int, int]]:
    """(row, col) of a grid's centre cell when it is the headline DCF case.

    The centre of an odd-sized grid holds the unshifted inputs, so it should
    reproduce the headline DCF implied price. When it does not (e.g. a Gordon
    grid shown beside an exit-multiple headline, or a missing centre cell) it
    is not the base case, and None is returned so nothing is highlighted.
    """
    dcf = getattr(report, "dcf", None)
    headline = getattr(dcf, "implied_price", None) if dcf is not None else None
    rows = list(getattr(sens, "row_values", None) or [])
    cols = list(getattr(sens, "col_values", None) or [])
    grid = list(getattr(sens, "grid", None) or [])
    if not is_num(headline) or len(rows) % 2 == 0 or len(cols) % 2 == 0:
        return None
    i, j = len(rows) // 2, len(cols) // 2
    grid_row = grid[i] if i < len(grid) else None
    cell = grid_row[j] if grid_row and j < len(grid_row) else None
    if not is_num(cell):
        return None
    headline = float(headline)
    if abs(float(cell) - headline) > max(1e-6, 1e-6 * abs(headline)):
        return None
    return i, j


def _write_sensitivity(ws: Worksheet, report: ValuationReport, money_fmt: str) -> None:
    _title(ws, "Sensitivity Analysis")
    sensitivities = report.sensitivities or []
    if not sensitivities:
        _note(ws, 3, "No sensitivity grids available.")
        _set_widths(ws, {1: 30})
        return

    row = 3
    max_cols = 1  # track widest grid for column-width sizing
    for sens in sensitivities:
        title = getattr(sens, "title", None) or "Sensitivity grid"
        row_label = getattr(sens, "row_label", None) or "Rows"
        col_label = getattr(sens, "col_label", None) or "Cols"
        row_values = list(getattr(sens, "row_values", []) or [])
        col_values = list(getattr(sens, "col_values", []) or [])
        grid = list(getattr(sens, "grid", []) or [])

        _set(ws, row, 1, title, font=_SUBHEADER_FONT)
        row += 1
        # Axis legend line.
        _set(ws, row, 1, f"rows = {row_label}   |   columns = {col_label}",
             font=_NOTE_FONT)
        row += 1

        ncols = len(col_values)
        max_cols = max(max_cols, ncols + 1)

        # Header row: corner cell shows the column-axis label, then column values.
        corner = _set(ws, row, 1, f"{row_label} \\ {col_label}",
                      font=_HEADER_FONT, align=_CENTER)
        corner.fill = _HEADER_FILL
        # Column headers are the axis *levels*, formatted by the axis label
        # (rates -> percent, EV/EBITDA multiples -> 12.0x).
        for j, cval in enumerate(col_values):
            cell = _set(ws, row, 2 + j, _axis_value(cval),
                        fmt=_axis_fmt(col_label, cval),
                        font=_HEADER_FONT, align=_CENTER)
            cell.fill = _HEADER_FILL
        row += 1

        # Body rows: row-axis level in column 1, then implied prices. The
        # centre cell is marked as the base case only when it reproduces the
        # headline DCF price.
        base_cell = _base_case_cell(sens, report)
        has_blank = False
        for i, rval in enumerate(row_values):
            rcell = _set(ws, row, 1, _axis_value(rval), fmt=_axis_fmt(row_label, rval),
                         font=_LABEL_FONT, align=_RIGHT)
            rcell.fill = PatternFill("solid", fgColor="D9E1F2")
            grid_row = grid[i] if i < len(grid) else []
            for j in range(ncols):
                val = grid_row[j] if j < len(grid_row) else None
                # NaN/None (invalid combination) -> blank, not the literal "nan".
                cell = _set(ws, row, 2 + j, _num(val), fmt=money_fmt, align=_RIGHT)
                has_blank = has_blank or cell.value is None
                if base_cell == (i, j):
                    cell.font = _LABEL_FONT
                    cell.fill = _BASE_CASE_FILL
            row += 1
        if base_cell is not None:
            _note(ws, row, "Highlighted cell = headline DCF (base case).")
            row += 1
        if has_blank:
            _note(ws, row, BLANK_SENSITIVITY_CELL_NOTE)
            row += 1
        row += 2  # spacer between grids

    widths = {1: 18}
    for c in range(2, max_cols + 1):
        widths[c] = 14
    _set_widths(ws, widths)


# --------------------------------------------------------------------------- #
#  Public entry point
# --------------------------------------------------------------------------- #
def write_excel(report: ValuationReport, path: str) -> str:
    """Write a formatted multi-sheet valuation workbook and return `path`.

    Sheets: Summary, DCF, Comps, DDM_FCFE, Sensitivity. Every model section is
    guarded against being None and degrades to a human-readable "not available"
    note. Live Excel formulas are used where practical (PV = FCFF*DF,
    EV = SUM(PVs)+PV_TV, equity = EV - net debt - minority interest - preferred
    equity, implied = equity/shares, upside = implied/current - 1) so the
    workbook recalculates on user edits.
    """
    money_fmt = _money_fmt(report)

    wb = Workbook()
    # Re-purpose the default first sheet as Summary, then add the rest in order.
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_dcf = wb.create_sheet("DCF")
    ws_comps = wb.create_sheet("Comps")
    ws_ddm = wb.create_sheet("DDM_FCFE")
    ws_sens = wb.create_sheet("Sensitivity")

    # Each writer is independently guarded; one bad section must not sink the file.
    for writer, ws in (
        (_write_summary, ws_summary),
        (_write_dcf, ws_dcf),
        (_write_comps, ws_comps),
        (_write_ddm_fcfe, ws_ddm),
        (_write_sensitivity, ws_sens),
    ):
        try:
            writer(ws, report, money_fmt)
        except Exception as exc:  # noqa: BLE001 - degrade to a note on that sheet
            _note(ws, ws.max_row + 2, f"This section could not be written: {exc}")

    wb.save(path)
    return path
